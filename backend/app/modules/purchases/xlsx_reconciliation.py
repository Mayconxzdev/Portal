from __future__ import annotations

import csv
import hashlib
import os
import re
import tempfile
import unicodedata
import uuid
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable

import openpyxl
from sqlalchemy import desc, func, or_
from sqlalchemy.orm import Session, joinedload

from app.models.master_data import Person, PersonContact, ProductFamily, ProductItem, Supplier
from app.models.purchase import (
    PurchaseCatalogSearchIndex,
    PurchasePriceEvidence,
    PurchasePriceHistory,
    PurchasePriceReference,
    PurchaseSupplierPriceOffer,
    PurchaseXlsxCatalogRow,
    PurchaseXlsxImportRun,
)


DEFAULT_XLSX_PATH = Path("storage/imports/compras-demo.xlsx")
IGNORED_SHEET_NAMES = {"indice", "índice", "planilha1"}
SMART_CATALOG_PARSER_VERSION = "smart_catalog_v2.3"
TECHNICAL_KEYWORDS = {
    "diametro": "diametro",
    "diâmetro": "diametro",
    "cv": "cv",
    "polo": "polo",
    "polos": "polo",
    "pólo": "polo",
    "pas": "pas",
    "pás": "pas",
    "vazao": "vazao",
    "vazão": "vazao",
    "pressao": "pressao",
    "pressão": "pressao",
    "pintura": "pintura",
    "fabricante": "fabricante",
}


@dataclass
class ParsedPurchaseRow:
    row_key: str
    sheet: str
    row_number: int
    parser_type: str
    category: str
    family: str
    subfamily: str | None
    variation: str
    display_name: str
    original_description: str
    code: str | None = None
    manufacturer_code: str | None = None
    manufacturer: str | None = None
    supplier_name: str | None = None
    company_name: str | None = None
    email: str | None = None
    phone: str | None = None
    raw_price: Decimal | None = None
    ipi: Decimal | None = None
    adjustment: Decimal | None = None
    final_value: Decimal | None = None
    price_used: Decimal | None = None
    price_type: str = "nao_identificado"
    unit: str = "un"
    weight: str | None = None
    updated_at: str | None = None
    attributes: dict[str, Any] = field(default_factory=dict)
    issues: list[str] = field(default_factory=list)
    portal_item_id: uuid.UUID | None = None
    portal_item_name: str | None = None
    portal_supplier_id: uuid.UUID | None = None
    portal_supplier_name: str | None = None
    portal_price: Decimal | None = None
    difference_amount: Decimal | None = None
    difference_percent: Decimal | None = None
    situation: str = "Pronto"


def _text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\n", " ")).strip()


def _norm(value: Any) -> str:
    text = _text(value).lower()
    trans = str.maketrans("áàâãäéèêëíìîïóòôõöúùûüç", "aaaaaeeeeiiiiooooouuuuc")
    return text.translate(trans)


def _money(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            return Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        except InvalidOperation:
            return None
    text = _text(value).replace("R$", "").replace("%", "").strip()
    if not text or text in {"-", "."}:
        return None
    text = re.sub(r"[^0-9,.\-]", "", text)
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".") if text.rfind(",") > text.rfind(".") else text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return Decimal(text).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError):
        return None


def _email(value: Any) -> str | None:
    text = _text(value).lower()
    match = re.search(r"[\w.\-+]+@[\w.\-]+\.\w+", text)
    return match.group(0) if match else None


def _phone(value: Any) -> str | None:
    text = _text(value)
    digits = re.sub(r"\D", "", text)
    if len(digits) < 8:
        return None
    return digits[:15]


def _clean_supplier(value: Any) -> str | None:
    text = _text(value)
    if not text:
        return None
    text = re.sub(r"\s+", " ", text).strip(" -")
    return text[:255] or None


def _row_key(sheet: str, row_number: int, values: Iterable[Any]) -> str:
    raw = "|".join(_text(v) for v in values)
    digest = hashlib.sha1(f"{sheet}:{row_number}:{raw}".encode("utf-8", errors="ignore")).hexdigest()[:16]
    return f"{sheet}:{row_number}:{digest}"


def _looks_like_title(values: list[Any]) -> bool:
    texts = [_text(v) for v in values if _text(v)]
    nums = [_money(v) for v in values]
    if not texts:
        return False
    if any(n is not None and n > 0 for n in nums):
        return False
    return len(texts) <= 3 and len(" ".join(texts)) <= 80


def _is_measure_title(value: Any) -> bool:
    text = _norm(value).strip()
    if not text:
        return False
    return bool(re.fullmatch(r"\d+(?:[,.]\d+)?(?:\s*(?:mm|m|kg|pol|\"))?", text))


def _is_family_title(value: Any, sheet: str) -> bool:
    text = _text(value).strip()
    norm = _norm(text)
    if not text or _is_measure_title(text):
        return False
    if norm in {"arame", "tela", "valor m2", "valor mt2", "preco por m2", "preco p kg", "email", "telefone"}:
        return False
    if len(norm) <= 3:
        return False
    strong_terms = [
        "arame btc",
        "arames inox",
        "tela quadrada",
        "tela",
        "chapa",
        "cantoneira",
        "barra",
        "tubo",
        "tarugo",
        "perfil",
        "helice",
        "pvc",
    ]
    if any(term in norm for term in strong_terms):
        return True
    return len(norm.split()) >= 2 and not any(term in _norm(sheet) for term in ["helice"])


def _is_root_family_title(value: Any) -> bool:
    norm = _norm(value)
    if not norm:
        return False
    root_terms = [
        "chapa",
        "tela quadrada",
        "arame btc",
        "arame mig",
        "arame solda",
        "arames inox",
        "cantoneira",
        "barra chata",
        "barra redonda",
        "tubo",
        "tarugo",
        "perfil",
        "pvc",
    ]
    if " arame btc cl " in f" {norm} ":
        return False
    if norm.startswith("ch ") or norm.startswith("ch."):
        return False
    return any(term in norm for term in root_terms)


def _find_header_map(rows: list[tuple[int, list[Any]]]) -> tuple[int, dict[str, int]] | None:
    aliases = {
        "codigo": ["codigo", "código", "cod"],
        "description": ["material", "descricao", "descrição", "produto", "item"],
        "raw_price": ["preco", "preço", "preco unit", "preço unit", "preco p", "preço p"],
        "ipi": ["ipi"],
        "adjustment": ["reajuste"],
        "final_value": ["valor final", "preco final", "preço final"],
        "weight": ["peso", "kg/m"],
        "supplier": ["fornecedor"],
        "updated_at": ["atualizado", "ultima atualizacao", "última atualização"],
        "company": ["empresa"],
        "email": ["email", "e-mail"],
        "phone": ["telefone", "fone"],
        "manufacturer": ["fabricante"],
        "manufacturer_code": ["codigo do fabricante", "código do fabricante", "cod fabricante"],
        "diametro": ["diametro", "diâmetro"],
        "cv": ["cv"],
        "polo": ["polo", "pólo"],
        "pas": ["pas", "pás"],
        "vazao": ["vazao", "vazão"],
        "pressao": ["pressao", "pressão"],
        "pintura": ["pintura"],
    }
    best: tuple[int, dict[str, int], int] | None = None
    for pos, (row_num, values) in enumerate(rows[:30]):
        current = [_norm(v) for v in values]
        has_code = any("codigo" in cell or "código" in cell for cell in current)
        has_description = any(any(term in cell for term in ("material", "descricao", "descrição", "produto", "item")) for cell in current)
        has_price = any("preco" in cell or "preço" in cell or "valor final" in cell for cell in current)
        has_technical = any(any(term in cell for term in ("diametro", "diâmetro", "cv", "polo", "pólo", "pas", "pás", "vazao", "vazão", "pressao", "pressão")) for cell in current)
        if not (has_code and (has_description or (has_price and has_technical))):
            continue
        next_values = rows[pos + 1][1] if pos + 1 < len(rows) else []
        max_len = max(len(values), len(next_values))
        header = [
            _norm(values[idx] if idx < len(values) else "") + " " + _norm(next_values[idx] if idx < len(next_values) else "")
            for idx in range(max_len)
        ]
        mapping: dict[str, int] = {}
        for key, terms in aliases.items():
            for idx, cell in enumerate(header):
                if any(term in cell for term in terms):
                    mapping[key] = idx
                    break
        score = int("description" in mapping) + int("raw_price" in mapping) + int("final_value" in mapping) + int("supplier" in mapping)
        if score >= 2 and (best is None or score > best[2]):
            best = (row_num, mapping, score)
    if not best:
        return None
    return best[0], best[1]


def _sheet_parser_type(sheet_name: str, header_map: dict[str, int] | None) -> str:
    name = _norm(sheet_name)
    if "helice" in name or "hélice" in sheet_name.lower():
        return "helices"
    if any(term in name for term in ["chapa", "cantoneira", "barra", "tubo", "tarugo", "arame", "tela", "perfil"]):
        return "metais_blocos"
    if any(term in name for term in ["climatizador", "conex"]):
        return "climatizadores_conexoes"
    return "materiais_padrao" if header_map else "fallback_revisao"


def _extract_attributes(text: str, row: dict[str, Any] | None = None) -> dict[str, Any]:
    attrs: dict[str, Any] = {}
    source = f"{text} " + " ".join(f"{k} {v}" for k, v in (row or {}).items() if v)
    norm = _norm(source)
    thickness = re.search(r"(\d+(?:[,.]\d+)?)\s*mm", norm)
    if thickness:
        attrs["medida_milimetro"] = thickness.group(0).replace(",", ".")
    inch = re.findall(r"\d+(?:[./]\d+)?\s*\"", source)
    if inch:
        attrs["medida_polegada"] = " x ".join(inch[:3])
    dimension = re.search(r"(\d+(?:[,.]\d+)?)\s*x\s*(\d+(?:[,.]\d+)?)(?:\s*x\s*(\d+(?:[,.]\d+)?))?\s*(m|mm)?", norm)
    if dimension:
        attrs["dimensao"] = dimension.group(0).strip()
    kgm = re.search(r"(\d+(?:[,.]\d+)?)\s*kg\s*/?\s*m", norm)
    if kgm:
        attrs["kg_m"] = kgm.group(0).replace(",", ".")
    for label, canonical in TECHNICAL_KEYWORDS.items():
        match = re.search(rf"{label}\s*[:\-]?\s*([a-z0-9,./\"º°\s]+)", norm)
        if match:
            attrs[canonical] = match.group(1).strip()[:60]
    return attrs


def _family_from_context(sheet: str, context: list[str], description: str) -> tuple[str, str | None]:
    invalid = {"nenhum nome faltando", "preco p/ mt", "preço p/ mt", "email", "telefone"}
    candidates = [c for c in context if c and _norm(c) not in invalid]
    family = candidates[0] if candidates else sheet
    subfamily = candidates[-1] if len(candidates) > 1 and candidates[-1] != family else None
    if _norm(sheet).startswith("helice fm"):
        family = "HÉLICE FIBRAMETAL"
    elif _norm(sheet).startswith("helice mw"):
        family = "HÉLICE MULTI WING"
    elif not family or len(family) < 3:
        family = description.split("|")[0].strip()[:80] or sheet
    return family.upper(), subfamily


def _price_type(sheet: str, description: str, unit: str, header_map: dict[str, int]) -> str:
    text = _norm(f"{sheet} {description} {unit}")
    if "kg" in text:
        return "kg"
    if "metro" in text or " mt" in text or "/m" in text:
        return "metro"
    if any(term in text for term in ["barra", "cantoneira", "tarugo"]):
        return "barra"
    if "chapa" in text:
        return "chapa"
    if "rolo" in text:
        return "rolo"
    if "caixa" in text:
        return "caixa"
    if "fabricante" in header_map:
        return "fabricante"
    return "unidade"


def _decimal_to_json(value: Decimal | None) -> float | None:
    return float(value) if value is not None else None


SEARCH_ALIASES = {
    "bct": "btc",
    "btc": "bct",
    "ch": "chapa",
    "galv": "galvanizada",
    "galvaniz": "galvanizada",
    "diam": "diametro",
    "ø": "diametro",
    "kgm": "kg/m",
}


def normalize_search_text(value: Any) -> str:
    text = _text(value).lower().replace("ø", " diametro ")
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = text.replace("kg/m", " kgm ")
    text = re.sub(r"[^a-z0-9/.,\" ]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    tokens = []
    for token in text.split():
        tokens.append(token)
        alias = SEARCH_ALIASES.get(token)
        if alias:
            tokens.append(alias)
    return " ".join(dict.fromkeys(tokens))


def search_matches(search_text: str, term: str | None) -> bool:
    normalized_term = normalize_search_text(term or "")
    if not normalized_term:
        return True
    haystack = normalize_search_text(search_text)
    return all(part in haystack for part in normalized_term.split())


def _canonical_key(*parts: Any) -> str:
    return normalize_search_text(" ".join(_text(part) for part in parts if _text(part)))[:500]


def _status_from_row(row: ParsedPurchaseRow, *, has_email: bool, ambiguous: bool = False) -> tuple[str, list[str]]:
    ignored = {"Linha complementar da planilha"}
    issues = [issue for issue in row.issues if issue not in ignored]
    if not row.supplier_name and not row.company_name and "Sem fornecedor" not in issues:
        issues.append("Sem fornecedor")
    if not has_email and "Fornecedor sem e-mail" not in issues:
        issues.append("Fornecedor sem e-mail")
    if ambiguous and "Variação ambígua" not in issues:
        issues.append("Variação ambígua")
    if "Sem fornecedor" in issues:
        return "Sem fornecedor", list(dict.fromkeys(issues))
    if "Fornecedor sem e-mail" in issues:
        return "Sem e-mail", list(dict.fromkeys(issues))
    if "Variação ambígua" in issues:
        return "Variação ambígua", list(dict.fromkeys(issues))
    if issues:
        return "Precisa revisão", list(dict.fromkeys(issues))
    return "Pronto", []


def _latest_import_run(db: Session) -> PurchaseXlsxImportRun | None:
    return db.query(PurchaseXlsxImportRun).order_by(desc(PurchaseXlsxImportRun.started_at)).first()


def _file_hash(path: Path) -> str | None:
    if not path.exists():
        return None
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _parse_observed_sort(value: str | None) -> datetime:
    text = _text(value)
    if not text:
        return datetime.min
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(text[:10], fmt)
        except ValueError:
            continue
    return datetime.min


def _supplier_display_name(row: ParsedPurchaseRow) -> str | None:
    return _clean_supplier(row.supplier_name) or _clean_supplier(row.company_name)


def _price_for_offer(row: ParsedPurchaseRow) -> Decimal | None:
    return row.raw_price or row.price_used or row.final_value


def _serialize_catalog_family(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "family_id": row.get("family_id"),
        "family_name": row.get("family_name"),
        "description": row.get("description"),
        "category": row.get("category"),
        "variations_count": row.get("variations_count", 0),
        "suppliers_count": row.get("suppliers_count", 0),
        "min_price": _decimal_to_json(row.get("min_price")),
        "max_price": _decimal_to_json(row.get("max_price")),
        "last_updated_at": row.get("last_updated_at"),
        "variations": row.get("variations", []),
    }


def _serialize_offer(offer: PurchaseSupplierPriceOffer) -> dict[str, Any]:
    supplier_name = None
    if offer.supplier and offer.supplier.person:
        supplier_name = offer.supplier.person.name
    return {
        "id": offer.id,
        "product_item_id": offer.product_item_id,
        "supplier_id": offer.supplier_id,
        "supplier_name": supplier_name or offer.supplier_name_snapshot,
        "email": offer.email,
        "phone": offer.phone,
        "raw_price": _decimal_to_json(offer.raw_price),
        "final_value": _decimal_to_json(offer.final_value),
        "current_price": _decimal_to_json(offer.final_value or offer.raw_price),
        "currency": offer.currency,
        "unit": offer.unit,
        "observed_at": offer.observed_at,
        "is_current_supplier": offer.is_current_supplier,
        "is_consolidated": offer.is_consolidated,
        "status": offer.status,
        "source_row_key": offer.source_row_key,
        "updated_at": offer.updated_at,
    }


def _serialize(row: ParsedPurchaseRow, *, include_technical: bool = False) -> dict[str, Any]:
    data = {
        "row_key": row.row_key,
        "sheet": row.sheet,
        "row_number": row.row_number,
        "category": row.category,
        "family": row.family,
        "subfamily": row.subfamily,
        "variation": row.variation,
        "display_name": row.display_name,
        "supplier_name": row.supplier_name,
        "company_name": row.company_name,
        "email": row.email,
        "phone": row.phone,
        "spreadsheet_price": _decimal_to_json(row.price_used),
        "raw_price": _decimal_to_json(row.raw_price),
        "final_value": _decimal_to_json(row.final_value),
        "portal_price": _decimal_to_json(row.portal_price),
        "difference_amount": _decimal_to_json(row.difference_amount),
        "difference_percent": _decimal_to_json(row.difference_percent),
        "price_type": row.price_type,
        "unit": row.unit,
        "updated_at": row.updated_at,
        "situation": row.situation,
        "issues": row.issues,
        "portal_item_id": str(row.portal_item_id) if row.portal_item_id else None,
        "portal_supplier_id": str(row.portal_supplier_id) if row.portal_supplier_id else None,
        "portal_item_name": row.portal_item_name,
        "portal_supplier_name": row.portal_supplier_name,
        "ready_for_quote": bool(row.email and row.supplier_name and not row.issues),
    }
    if include_technical:
        data["technical_details"] = {
            "parser_type": row.parser_type,
            "codigo_interno": row.code,
            "codigo_fabricante": row.manufacturer_code,
            "fabricante": row.manufacturer,
            "descricao_original_planilha": row.original_description,
            "atributos": row.attributes,
        }
    return data


def _serialize_catalog_family(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "family_id": row.get("family_id"),
        "family_name": row.get("family_name"),
        "description": row.get("description"),
        "category": row.get("category"),
        "variations_count": row.get("variations_count", 0),
        "suppliers_count": row.get("suppliers_count", 0),
        "min_price": _decimal_to_json(row.get("min_price")),
        "max_price": _decimal_to_json(row.get("max_price")),
        "last_updated_at": row.get("last_updated_at"),
        "variations": row.get("variations", []),
    }


def _serialize_offer(offer: PurchaseSupplierPriceOffer) -> dict[str, Any]:
    supplier_name = None
    if offer.supplier and offer.supplier.person:
        supplier_name = offer.supplier.person.name
    return {
        "id": offer.id,
        "product_item_id": offer.product_item_id,
        "supplier_id": offer.supplier_id,
        "supplier_name": supplier_name or offer.supplier_name_snapshot,
        "email": offer.email,
        "phone": offer.phone,
        "raw_price": _decimal_to_json(offer.raw_price),
        "final_value": _decimal_to_json(offer.final_value),
        "current_price": _decimal_to_json(offer.final_value or offer.raw_price),
        "currency": offer.currency,
        "unit": offer.unit,
        "observed_at": offer.observed_at,
        "is_current_supplier": offer.is_current_supplier,
        "is_consolidated": offer.is_consolidated,
        "status": offer.status,
        "source_row_key": offer.source_row_key,
        "updated_at": offer.updated_at,
    }


def _serialize(row: ParsedPurchaseRow, *, include_technical: bool = False) -> dict[str, Any]:
    data = {
        "row_key": row.row_key,
        "sheet": row.sheet,
        "row_number": row.row_number,
        "category": row.category,
        "family": row.family,
        "subfamily": row.subfamily,
        "variation": row.variation,
        "display_name": row.display_name,
        "supplier_name": row.supplier_name,
        "company_name": row.company_name,
        "email": row.email,
        "phone": row.phone,
        "spreadsheet_price": _decimal_to_json(row.price_used),
        "raw_price": _decimal_to_json(row.raw_price),
        "final_value": _decimal_to_json(row.final_value),
        "portal_price": _decimal_to_json(row.portal_price),
        "difference_amount": _decimal_to_json(row.difference_amount),
        "difference_percent": _decimal_to_json(row.difference_percent),
        "price_type": row.price_type,
        "unit": row.unit,
        "updated_at": row.updated_at,
        "situation": row.situation,
        "issues": row.issues,
        "portal_item_id": str(row.portal_item_id) if row.portal_item_id else None,
        "portal_supplier_id": str(row.portal_supplier_id) if row.portal_supplier_id else None,
        "portal_item_name": row.portal_item_name,
        "portal_supplier_name": row.portal_supplier_name,
        "ready_for_quote": bool(row.email and row.supplier_name and not row.issues),
    }
    if include_technical:
        data["technical_details"] = {
            "parser_type": row.parser_type,
            "codigo_interno": row.code,
            "codigo_fabricante": row.manufacturer_code,
            "fabricante": row.manufacturer,
            "descricao_original_planilha": row.original_description,
            "atributos": row.attributes,
        }
    return data


class PurchasesXlsxReconciliation:
    @staticmethod
    def source_path() -> Path:
        return Path(os.getenv("PURCHASES_XLSX_PATH") or DEFAULT_XLSX_PATH)

    @staticmethod
    def parse_workbook(path: Path | None = None, *, max_rows_per_sheet: int | None = None) -> tuple[list[ParsedPurchaseRow], dict[str, Any]]:
        source = path or PurchasesXlsxReconciliation.source_path()
        if not source.exists():
            return [], {
                "source_path": str(source),
                "error": "Planilha de compras não encontrada.",
                "sheets_read": [],
                "sheets_ignored": [],
                "layouts": {},
            }
        wb = openpyxl.load_workbook(source, data_only=True)
        parsed: list[ParsedPurchaseRow] = []
        ignored: list[str] = []
        layouts: dict[str, int] = Counter()
        sheets_read: list[str] = []
        contacts_registry: dict[str, dict[str, Any]] = {}
        def _get_cell_bg_color(cell: Any) -> str | None:
            if not cell or not cell.fill or cell.fill.fill_type is None:
                return None
            try:
                color = cell.fill.start_color
                if color and color.type == "rgb" and isinstance(color.rgb, str):
                    return color.rgb.upper()
            except Exception:
                pass
            return None
        try:
            for ws in wb.worksheets:
                sheet_norm = _norm(ws.title)
                if sheet_norm in IGNORED_SHEET_NAMES or sheet_norm.replace(" ", "") in IGNORED_SHEET_NAMES:
                    continue
                cell_rows = list(ws.iter_rows())
                if not cell_rows:
                    continue
                value_rows = [(idx, [cell.value for cell in row]) for idx, row in enumerate(cell_rows, start=1)]
                header = _find_header_map(value_rows)
                header_row = header[0] if header else 0
                mapping = header[1] if header else {}
                if not mapping:
                    continue
                for row_idx, row_cells in enumerate(cell_rows, start=1):
                    if row_idx <= header_row:
                        continue
                    row_data = {key: row_cells[idx].value if idx < len(row_cells) else None for key, idx in mapping.items()}
                    fornecedor = _clean_supplier(row_data.get("supplier"))
                    empresa = _clean_supplier(row_data.get("company"))
                    email = _email(row_data.get("email")) or _email(row_data.get("phone")) or _email(empresa)
                    phone = _phone(row_data.get("phone")) or _phone(empresa)
                    if not email and not phone:
                        continue
                    if empresa:
                        emp_norm = _norm(empresa)
                        if emp_norm:
                            entry = contacts_registry.setdefault(emp_norm, {"name": empresa, "email": None, "phone": None})
                            if email: entry["email"] = email
                            if phone: entry["phone"] = phone
                    if fornecedor:
                        forn_norm = _norm(fornecedor)
                        if forn_norm:
                            if empresa and forn_norm == _norm(empresa):
                                entry = contacts_registry.setdefault(forn_norm, {"name": fornecedor, "email": None, "phone": None})
                                if email: entry["email"] = email
                                if phone: entry["phone"] = phone
            for ws in wb.worksheets:
                sheet_norm = _norm(ws.title)
                if sheet_norm in IGNORED_SHEET_NAMES or sheet_norm.replace(" ", "") in IGNORED_SHEET_NAMES:
                    ignored.append(ws.title)
                    continue
                cell_rows = list(ws.iter_rows())
                value_rows = [(idx, [cell.value for cell in row]) for idx, row in enumerate(cell_rows, start=1)]
                if not any(any(_text(cell) for cell in values) for _, values in value_rows):
                    ignored.append(ws.title)
                    continue
                header = _find_header_map(value_rows)
                header_row = header[0] if header else 0
                mapping = header[1] if header else {}
                parser_type = _sheet_parser_type(ws.title, mapping)
                layouts[parser_type] += 1
                sheets_read.append(ws.title)
                has_fills = False
                for row_cells in cell_rows:
                    if len(row_cells) > 1:
                        bg = _get_cell_bg_color(row_cells[1])
                        if bg in ("FFFFC000", "FFFC6E04", "FFE6EDEE"):
                            has_fills = True
                            break
                context: list[str] = []
                current_family_title: str | None = None
                current_subfamily_title: str | None = None
                current_measure_title: str | None = None
                last_weight: str | None = None
                last_description: str | None = None
                last_variation_description: str | None = None
                data_rows = cell_rows[header_row:] if header_row else cell_rows
                if max_rows_per_sheet:
                    data_rows = data_rows[:max_rows_per_sheet]
                for row_cells in data_rows:
                    row_idx = row_cells[0].row if row_cells else 0
                    row_data = {key: row_cells[idx].value if idx < len(row_cells) else None for key, idx in mapping.items()}
                    source_desc = _text(row_data.get("description"))
                    raw_price = _money(row_data.get("raw_price"))
                    final_value = _money(row_data.get("final_value"))
                    supplier = _clean_supplier(row_data.get("supplier"))
                    company = _clean_supplier(row_data.get("company"))
                    cell_b = row_cells[1] if len(row_cells) > 1 else None
                    val_b = cell_b.value if cell_b else None
                    val_b_str = _text(val_b)
                    bg_color_b = _get_cell_bg_color(cell_b) if cell_b else None
                    is_bold_b = cell_b.font.bold if cell_b and cell_b.font else False
                    if has_fills:
                        if bg_color_b in ("FFFFC000", "FFFC6E04"):
                            if val_b_str:
                                current_family_title = val_b_str
                                current_subfamily_title = None
                                current_measure_title = None
                                context = [val_b_str]
                            continue
                        if is_bold_b and bg_color_b not in ("FFE6EDEE", "FFFFC000", "FFFC6E04") and val_b_str:
                            if _is_measure_title(val_b_str):
                                current_measure_title = val_b_str
                            else:
                                current_subfamily_title = val_b_str
                            context = (
                                ([current_family_title] if current_family_title else [])
                                + ([current_subfamily_title] if current_subfamily_title else [])
                                + ([current_measure_title] if current_measure_title else [])
                            )
                            continue
                        cell_f = row_cells[mapping["final_value"]] if "final_value" in mapping and mapping["final_value"] < len(row_cells) else None
                        cell_h = row_cells[mapping["supplier"]] if "supplier" in mapping and mapping["supplier"] < len(row_cells) else None
                        bg_f = _get_cell_bg_color(cell_f) if cell_f else None
                        bg_h = _get_cell_bg_color(cell_h) if cell_h else None
                        if bg_f == "FFFFD966" or bg_h == "FFC7D0DB":
                            continue
                    else:
                        if source_desc and raw_price is None and final_value is None:
                            if _is_family_title(source_desc, ws.title):
                                if not current_family_title or _is_root_family_title(source_desc):
                                    current_family_title = source_desc
                                    current_subfamily_title = None
                                    current_measure_title = None
                                    context = [source_desc]
                                else:
                                    current_subfamily_title = source_desc
                                    context = [current_family_title, source_desc]
                            elif _is_measure_title(source_desc):
                                current_measure_title = source_desc
                                context = (
                                    ([current_family_title] if current_family_title else [])
                                    + ([current_subfamily_title] if current_subfamily_title else [])
                                    + [source_desc]
                                )
                            continue
                        if not source_desc and raw_price is None:
                            continue
                    desc = source_desc or last_variation_description or last_description or ""
                    if parser_type == "helices" and _text(row_data.get("codigo")):
                        desc = _text(row_data.get("codigo"))
                    if not desc:
                        continue
                    weight_text = _text(row_data.get("weight"))
                    if weight_text and "kg" in _norm(weight_text):
                        last_weight = weight_text
                    price = raw_price if raw_price is not None else final_value
                    if price is None or price <= 0:
                        continue
                    if source_desc:
                        last_description = source_desc
                        last_variation_description = source_desc
                    family_context = []
                    row_family_override = None
                    if source_desc and _is_root_family_title(source_desc):
                        current_norm = normalize_search_text(current_family_title or "")
                        desc_norm = normalize_search_text(source_desc)
                        if not current_norm or not desc_norm.startswith(current_norm):
                            row_family_override = source_desc
                            current_family_title = source_desc
                            current_subfamily_title = None
                            current_measure_title = None
                    if row_family_override:
                        family_context.append(row_family_override)
                    elif current_family_title:
                        family_context.append(current_family_title)
                    if current_subfamily_title:
                        family_context.append(current_subfamily_title)
                    if current_measure_title:
                        family_context.append(current_measure_title)
                    context_for_row = family_context or context
                    family, subfamily = _family_from_context(ws.title, context_for_row, desc)
                    code = _text(row_data.get("codigo")) or None
                    manufacturer = _clean_supplier(row_data.get("manufacturer"))
                    manufacturer_code = _text(row_data.get("manufacturer_code")) or None
                    attrs = _extract_attributes(" ".join([*context_for_row, desc, weight_text or last_weight or ""]), row_data)
                    for tech_key in ("diametro", "cv", "polo", "pas", "vazao", "pressao", "pintura"):
                        tech_value = _text(row_data.get(tech_key))
                        if tech_value:
                            attrs[tech_key] = tech_value
                    if current_measure_title:
                        attrs.setdefault("medida_bloco", current_measure_title)
                    if weight_text or last_weight:
                        attrs["peso"] = weight_text or last_weight
                    if manufacturer:
                        attrs["fabricante"] = manufacturer
                    if manufacturer_code:
                        attrs["codigo_fabricante"] = manufacturer_code
                    if final_value is not None:
                        attrs["valor_final_original"] = float(final_value)
                    unit = "un"
                    if "kg" in _norm(desc) or "kg_m" in attrs:
                        unit = "kg"
                    elif "barra" in _norm(ws.title + " " + desc):
                        unit = "barra"
                    elif "chapa" in _norm(ws.title + " " + desc):
                        unit = "chapa"
                    elif "m2" in _norm(desc) or "malha" in _norm(desc):
                        unit = "m2"
                    issues: list[str] = []
                    if not source_desc and last_variation_description:
                        issues.append("Oferta de fornecedor da mesma variacao")
                    if not supplier and not company:
                        issues.append("Sem fornecedor")
                    email = None
                    phone = None
                    if supplier and company:
                        forn_norm = _norm(supplier)
                        emp_norm = _norm(company)
                        if forn_norm == emp_norm:
                            entry = contacts_registry.get(forn_norm)
                            if entry:
                                email = entry.get("email")
                                phone = entry.get("phone")
                        else:
                            issues.append(f"Revisar contato do fornecedor: e-mail na linha pertence à empresa {company}")
                    else:
                        name_to_lookup = _norm(supplier or company or "")
                        entry = contacts_registry.get(name_to_lookup)
                        if entry:
                            email = entry.get("email")
                            phone = entry.get("phone")
                    if not email:
                        email = _email(row_data.get("email")) or _email(row_data.get("phone"))
                    if not phone:
                        phone = _phone(row_data.get("phone"))
                    if not email:
                        issues.append("Fornecedor sem e-mail")
                    variation_base = desc
                    if current_measure_title and current_measure_title not in variation_base:
                        variation_base = f"{current_measure_title} | {variation_base}"
                    variation_parts = [variation_base]
                    for key in ("dimensao", "medida_polegada", "medida_milimetro", "kg_m", "diametro", "cv", "polo", "pas", "vazao", "pressao", "pintura"):
                        if attrs.get(key) and str(attrs[key]) not in variation_base:
                            variation_parts.append(str(attrs[key]))
                    variation = " | ".join(dict.fromkeys(p for p in variation_parts if p))
                    clean_fam = family.upper().strip()
                    clean_var = variation.strip()
                    var_norm = _norm(clean_var)
                    fam_norm = _norm(clean_fam)
                    if var_norm.startswith(fam_norm) or fam_norm in var_norm:
                        display_name = clean_var
                    else:
                        display_name = f"{clean_fam} | {clean_var}"
                    parsed.append(ParsedPurchaseRow(
                        row_key=_row_key(ws.title, row_idx, [c.value for c in row_cells]),
                        sheet=ws.title,
                        row_number=row_idx,
                        parser_type=parser_type,
                        category=ws.title,
                        family=family,
                        subfamily=subfamily,
                        variation=variation[:255],
                        display_name=display_name[:255],
                        original_description=desc,
                        code=code,
                        manufacturer_code=manufacturer_code,
                        manufacturer=manufacturer,
                        supplier_name=supplier or company,
                        company_name=company,
                        email=email,
                        phone=phone,
                        raw_price=raw_price,
                        ipi=_money(row_data.get("ipi")),
                        adjustment=_money(row_data.get("adjustment")),
                        final_value=final_value,
                        price_used=price,
                        price_type=_price_type(ws.title, desc, unit, mapping),
                        unit=unit,
                        weight=weight_text or last_weight,
                        updated_at=_text(row_data.get("updated_at")) or None,
                        attributes=attrs,
                        issues=issues,
                    ))
        finally:
            wb.close()
        meta = {
            "source_path": str(source),
            "sheets_read": sheets_read,
            "sheets_ignored": ignored,
            "layouts": dict(layouts),
            "file_updated_at": datetime.fromtimestamp(source.stat().st_mtime).isoformat() if source.exists() else None,
            "contacts_registry": contacts_registry,
        }
        return parsed, meta

    @staticmethod
    def compare_with_portal(db: Session, rows: list[ParsedPurchaseRow]) -> list[ParsedPurchaseRow]:
        people = db.query(Person).outerjoin(Supplier).filter(Supplier.id.isnot(None)).all()
        suppliers_by_name = {_norm(person.name): person.supplier for person in people if person.supplier}
        products = db.query(ProductItem).options(joinedload(ProductItem.family)).all()
        products_by_key = {_norm(product.canonical_key or product.sku or product.name): product for product in products}
        products_by_name = {_norm(product.name): product for product in products}

        for row in rows:
            supplier_key = _norm(row.supplier_name or row.company_name or "")
            supplier = suppliers_by_name.get(supplier_key)
            if not supplier and supplier_key:
                supplier = next((sup for name, sup in suppliers_by_name.items() if supplier_key in name or name in supplier_key), None)
            if supplier:
                row.portal_supplier_id = supplier.id
                row.portal_supplier_name = supplier.person.name if supplier.person else None
            else:
                row.issues.append("Fornecedor faltando no Portal")

            canonical = _norm(" ".join([row.family, row.variation, row.unit, row.manufacturer or ""]))
            product = products_by_key.get(canonical) or products_by_name.get(_norm(row.variation))
            if not product:
                product = next((item for name, item in products_by_name.items() if _norm(row.original_description) in name or name in _norm(row.original_description)), None)
            if product:
                row.portal_item_id = product.id
                row.portal_item_name = product.name
                if not product.family_id:
                    row.issues.append("Item sem família no Portal")
                ref_query = db.query(PurchasePriceReference).filter(
                    PurchasePriceReference.product_item_id == product.id,
                    PurchasePriceReference.is_active == True,
                )
                if row.portal_supplier_id:
                    ref_query = ref_query.filter(or_(
                        PurchasePriceReference.supplier_id == row.portal_supplier_id,
                        PurchasePriceReference.supplier_id == None,
                    ))
                ref = ref_query.order_by(PurchasePriceReference.supplier_id.desc(), PurchasePriceReference.updated_at.desc()).first()
                if ref:
                    row.portal_price = Decimal(str(ref.current_unit_price)).quantize(Decimal("0.01"))
                    if row.price_used is not None:
                        row.difference_amount = (row.price_used - row.portal_price).quantize(Decimal("0.01"))
                        if row.portal_price > 0:
                            row.difference_percent = ((row.difference_amount / row.portal_price) * Decimal("100")).quantize(Decimal("0.01"))
                        if abs(row.difference_amount) > Decimal("0.01"):
                            row.issues.append("Preço diferente")
                else:
                    row.issues.append("Sem preço atual no Portal")
            else:
                row.issues.append("Produto faltando no Portal")

            if len([r for r in rows if _norm(r.family) == _norm(row.family) and _norm(r.variation) == _norm(row.variation)]) > 1:
                row.issues.append("Variação ambígua")
            row.issues = list(dict.fromkeys(row.issues))
            if "Preço diferente" in row.issues:
                row.situation = "Preço diferente"
            elif "Produto faltando no Portal" in row.issues:
                row.situation = "Faltando no Portal"
            elif "Fornecedor sem e-mail" in row.issues:
                row.situation = "Sem e-mail"
            elif "Variação ambígua" in row.issues:
                row.situation = "Variação ambígua"
            elif row.issues:
                row.situation = "Precisa revisão"
            else:
                row.situation = "Pronto"
        return rows

    @staticmethod
    def payload(
        db: Session,
        *,
        search: str | None = None,
        status_filter: str = "all",
        limit: int = 80,
        offset: int = 0,
        include_technical: bool = False,
    ) -> dict[str, Any]:
        rows, meta = PurchasesXlsxReconciliation.parse_workbook()
        rows = PurchasesXlsxReconciliation.compare_with_portal(db, rows)
        term = _norm(search or "")
        if term:
            rows = [
                row for row in rows
                if term in _norm(" ".join([
                    row.sheet, row.family, row.variation, row.original_description,
                    row.supplier_name or "", row.company_name or "", row.email or "",
                    row.code or "", row.manufacturer_code or "",
                ]))
            ]
        if status_filter and status_filter != "all":
            mapping = {
                "divergent": "Preço diferente",
                "missing": "Faltando no Portal",
                "no_supplier": "Sem fornecedor",
                "no_email": "Sem e-mail",
                "ambiguous": "Variação ambígua",
                "price_diff": "Preço diferente",
                "ready": "Pronto",
            }
            expected = mapping.get(status_filter, status_filter)
            rows = [row for row in rows if row.situation == expected or expected in row.issues]
        total = len(rows)
        limit = max(1, min(limit, 200))
        offset = max(0, offset)
        page = rows[offset:offset + limit]
        suppliers = {(row.supplier_name or row.company_name or "").strip() for row in rows if row.supplier_name or row.company_name}
        emails = {row.email for row in rows if row.email}
        phones = {row.phone for row in rows if row.phone}
        families = {_norm(row.family) for row in rows if row.family}
        prices = [row for row in rows if row.price_used is not None]
        using_final = [row for row in rows if row.final_value is not None and row.price_used == row.final_value]
        matched_prices = [row for row in rows if row.portal_price is not None and row.difference_amount in (Decimal("0.00"), None)]
        divergent = [row for row in rows if row.situation == "Preço diferente"]
        summary = {
            "sheets_read": len(meta.get("sheets_read") or []),
            "suppliers_found": len(suppliers),
            "suppliers_with_email": len(emails),
            "suppliers_with_phone": len(phones),
            "suppliers_in_portal": len({row.portal_supplier_id for row in rows if row.portal_supplier_id}),
            "suppliers_missing": len({(row.supplier_name or row.company_name) for row in rows if (row.supplier_name or row.company_name) and not row.portal_supplier_id}),
            "products_found": len(rows),
            "families_found": len(families),
            "variations_found": len({_norm(row.variation) for row in rows}),
            "ambiguous_items": len([row for row in rows if "Variação ambígua" in row.issues]),
            "prices_detected": len(prices),
            "prices_using_final_value": len(using_final),
            "prices_matching_portal": len(matched_prices),
            "prices_divergent": len(divergent),
            "needs_review": len([row for row in rows if row.situation not in {"Pronto", "Preço diferente"}]),
            "suppliers_without_email": len({(row.supplier_name or row.company_name) for row in rows if (row.supplier_name or row.company_name) and not row.email}),
        }
        return {
            "source_path": meta.get("source_path"),
            "file_updated_at": meta.get("file_updated_at"),
            "sheets_read": meta.get("sheets_read", []),
            "sheets_ignored": meta.get("sheets_ignored", []),
            "layouts": meta.get("layouts", {}),
            "summary": summary,
            "items": [_serialize(row, include_technical=include_technical) for row in page],
            "total": total,
            "limit": limit,
            "offset": offset,
            "has_more": offset + limit < total,
        }

    @staticmethod
    def export_csv(db: Session) -> dict[str, Any]:
        payload = PurchasesXlsxReconciliation.payload(db, limit=100000, include_technical=True)
        export_dir = Path(os.getenv("PURCHASES_EXPORT_DIR") or (Path(tempfile.gettempdir()) / "Portal-Vesper-Dados"))
        try:
            export_dir.mkdir(parents=True, exist_ok=True)
        except OSError:
            export_dir = Path(tempfile.gettempdir()) / "Portal-Vesper-Dados"
            export_dir.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        path = export_dir / f"compras_conferencia_planilha_{stamp}.csv"
        fields = ["sheet", "display_name", "variation", "supplier_name", "email", "spreadsheet_price", "portal_price", "situation"]
        with path.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=fields)
            writer.writeheader()
            for item in payload["items"]:
                writer.writerow({field: item.get(field) for field in fields})
        return {"path": str(path), "rows": len(payload["items"])}


class PurchasesXlsxSmartCatalog:
    @staticmethod
    def _ensure_family(db: Session, name: str, category: str | None) -> ProductFamily:
        clean_name = (_text(name) or "SEM FAMILIA").upper()[:255]
        family = db.query(ProductFamily).filter(func.lower(ProductFamily.name) == clean_name.lower()).first()
        if family:
            return family
        family = ProductFamily(name=clean_name, description=f"Categoria da planilha: {category}" if category else None)
        db.add(family)
        db.flush()
        return family

    @staticmethod
    def _ensure_product_item(db: Session, family: ProductFamily, row: ParsedPurchaseRow) -> ProductItem:
        canonical = _canonical_key(row.category, family.name, row.variation, row.unit, row.manufacturer or "")
        item = db.query(ProductItem).filter(ProductItem.canonical_key == canonical).first()
        if item:
            attrs = dict(item.attributes or {})
            attrs.update(row.attributes or {})
            attrs.setdefault("descricao_original_planilha", row.original_description)
            attrs.setdefault("nome_antigo", row.original_description)
            item.attributes = attrs
            item.family_id = family.id
            item.category = row.category[:100] if row.category else item.category
            item.updated_at = datetime.now(timezone.utc)
            return item
        sku = f"XLSX-{hashlib.sha1(canonical.encode('utf-8', errors='ignore')).hexdigest()[:12].upper()}"
        item = ProductItem(
            sku=sku,
            name=(row.display_name or row.variation)[:255],
            description=row.original_description,
            item_type="RAW_MATERIAL",
            unit_of_measure=row.unit or "un",
            category=row.category[:100] if row.category else None,
            family_id=family.id,
            canonical_key=canonical,
            attributes={
                **(row.attributes or {}),
                "variation_name": row.variation,
                "descricao_original_planilha": row.original_description,
                "categoria_planilha": row.category,
                "aba_planilha": row.sheet,
                "codigo_interno": row.code,
                "codigo_fabricante": row.manufacturer_code,
                "fabricante": row.manufacturer,
            },
            is_active=True,
        )
        db.add(item)
        db.flush()
        return item

    @staticmethod
    def _ensure_supplier(
        db: Session,
        name: str | None,
        email: str | None,
        phone: str | None,
        category: str | None,
    ) -> Supplier | None:
        clean_name = _clean_supplier(name)
        if not clean_name:
            return None
        email = email.strip().lower() if email else None
        phone = phone.strip() if phone else None
        supplier = None
        if email:
            supplier = (
                db.query(Supplier)
                .join(Person, Supplier.person_id == Person.id)
                .filter(or_(func.lower(Person.email) == email, func.lower(Supplier.preferred_contact_email) == email))
                .first()
            )
        if not supplier:
            norm_name = normalize_search_text(clean_name)
            suppliers = db.query(Supplier).options(joinedload(Supplier.person).joinedload(Person.contacts)).all()
            supplier = next(
                (
                    candidate
                    for candidate in suppliers
                    if candidate.person and normalize_search_text(candidate.person.name) == norm_name
                ),
                None,
            )
        if not supplier:
            person = Person(
                type="COMPANY",
                name=clean_name,
                email=email,
                phone=phone,
                is_active=True,
                notes="Fornecedor criado automaticamente pela sincronizacao da planilha de compras.",
            )
            db.add(person)
            db.flush()
            supplier = Supplier(
                person_id=person.id,
                supplier_code=f"XLSX-SUP-{hashlib.sha1(clean_name.encode('utf-8', errors='ignore')).hexdigest()[:10].upper()}",
                categories=[category] if category else None,
                preferred_contact_email=email,
                status="ACTIVE",
            )
            db.add(supplier)
            db.flush()
        else:
            person = supplier.person
            if person:
                if email and not person.email:
                    person.email = email
                if phone and not person.phone:
                    person.phone = phone
                person.updated_at = datetime.now(timezone.utc)
            if email and not supplier.preferred_contact_email:
                supplier.preferred_contact_email = email
            categories = list(supplier.categories or [])
            if category and category not in categories:
                categories.append(category)
                supplier.categories = categories
            supplier.updated_at = datetime.now(timezone.utc)

        if supplier.person and (email or phone):
            existing_contacts = supplier.person.contacts or []
            has_email = email and any((contact.email or "").lower() == email for contact in existing_contacts)
            has_phone = phone and any((contact.phone or "") == phone for contact in existing_contacts)
            if (email and not has_email) or (phone and not has_phone):
                db.add(PersonContact(
                    person_id=supplier.person.id,
                    name="Contato da planilha",
                    role="Compras",
                    email=email,
                    phone=phone,
                    is_primary=not bool(existing_contacts),
                ))
        return supplier

    @staticmethod
    def _set_active_reference(
        db: Session,
        item: ProductItem,
        supplier_id: uuid.UUID | None,
        price: Decimal,
        current_user_id: int,
        *,
        source_row_key: str,
        notes: str,
    ) -> PurchasePriceReference:
        now = datetime.now(timezone.utc)
        active_refs = db.query(PurchasePriceReference).filter(
            PurchasePriceReference.product_item_id == item.id,
            PurchasePriceReference.supplier_id == supplier_id,
            PurchasePriceReference.is_active == True,
        ).all()
        current = next((ref for ref in active_refs if ref.current_unit_price == price), None)
        if current and len(active_refs) == 1:
            current.updated_at = now
            return current
        for ref in active_refs:
            ref.is_active = False
            ref.updated_at = now

        evidence = PurchasePriceEvidence(
            source_type="IMPORTED_XLSX",
            supplier_id=supplier_id,
            product_item_id=item.id,
            document_number=f"Compras Nova.xlsx:{source_row_key}"[:100],
            document_date=now,
            unit_price=price,
            quantity=Decimal("1"),
            total_amount=price,
            currency="BRL",
            unit_of_measure=item.unit_of_measure,
            raw_summary=notes,
            notes=notes,
            created_by_user_id=current_user_id,
            created_at=now,
        )
        db.add(evidence)
        db.flush()
        history = PurchasePriceHistory(
            product_item_id=item.id,
            supplier_id=supplier_id,
            evidence_id=evidence.id,
            unit_price=price,
            quantity=Decimal("1"),
            total_amount=price,
            currency="BRL",
            unit_of_measure=item.unit_of_measure,
            observed_at=now,
            source_type="IMPORTED_XLSX",
            source_id=source_row_key[:100],
            created_by_user_id=current_user_id,
            created_at=now,
        )
        db.add(history)
        db.flush()
        reference = PurchasePriceReference(
            product_item_id=item.id,
            supplier_id=supplier_id,
            current_unit_price=price,
            currency="BRL",
            unit_of_measure=item.unit_of_measure,
            source_history_id=history.id,
            source_evidence_id=evidence.id,
            approved_by_user_id=current_user_id,
            approved_at=now,
            notes=notes,
            is_active=True,
            created_at=now,
            updated_at=now,
        )
        db.add(reference)
        db.flush()
        return reference

    @staticmethod
    def _choose_current_row(rows: list[ParsedPurchaseRow]) -> ParsedPurchaseRow:
        dated = sorted(rows, key=lambda row: _parse_observed_sort(row.updated_at), reverse=True)
        if dated and _parse_observed_sort(dated[0].updated_at) != datetime.min:
            return dated[0]
        priced = [row for row in rows if _price_for_offer(row) is not None]
        if priced:
            return min(priced, key=lambda row: _price_for_offer(row) or Decimal("999999999"))
        return rows[0]

    @staticmethod
    def sync_from_xlsx(db: Session, current_user: Any, path: Path | None = None) -> dict[str, Any]:
        source = path or PurchasesXlsxReconciliation.source_path()
        started = datetime.now(timezone.utc)
        run = PurchaseXlsxImportRun(
            source_path=str(source),
            file_modified_at=datetime.fromtimestamp(source.stat().st_mtime) if source.exists() else None,
            file_hash=_file_hash(source),
            status="running",
            started_at=started,
            created_by_user_id=getattr(current_user, "id", None),
        )
        db.add(run)
        db.flush()
        try:
            rows, meta = PurchasesXlsxReconciliation.parse_workbook(source)
            if meta.get("error"):
                raise FileNotFoundError(meta["error"])

            db.query(PurchaseCatalogSearchIndex).delete(synchronize_session=False)
            db.query(PurchaseSupplierPriceOffer).delete(synchronize_session=False)
            db.query(PurchaseXlsxCatalogRow).delete(synchronize_session=False)
            db.flush()

            grouped: dict[str, list[ParsedPurchaseRow]] = defaultdict(list)
            for row in rows:
                grouped[_canonical_key(row.category, row.family, row.variation, row.unit, row.manufacturer or "")].append(row)

            supplier_names: set[str] = set()
            supplier_email_names: set[str] = set()
            items_by_key: dict[str, ProductItem] = {}
            catalog_rows: list[PurchaseXlsxCatalogRow] = []
            offers: list[PurchaseSupplierPriceOffer] = []

            for key, group_rows in grouped.items():
                base = group_rows[0]
                family = PurchasesXlsxSmartCatalog._ensure_family(db, base.family, base.category)
                item = PurchasesXlsxSmartCatalog._ensure_product_item(db, family, base)
                items_by_key[key] = item
                current_row = PurchasesXlsxSmartCatalog._choose_current_row(group_rows)
                current_supplier: Supplier | None = None
                group_supplier_names: list[str] = []

                for row in group_rows:
                    supplier_name = _supplier_display_name(row)
                    if supplier_name:
                        supplier_names.add(supplier_name)
                    supplier = PurchasesXlsxSmartCatalog._ensure_supplier(db, supplier_name, row.email, row.phone, row.category)
                    if row.email and supplier_name:
                        supplier_email_names.add(supplier_name)
                    if supplier and supplier.person:
                        group_supplier_names.append(supplier.person.name)
                    if row is current_row:
                        current_supplier = supplier
                    price = _price_for_offer(row)
                    situation, issues = _status_from_row(row, has_email=bool(row.email), ambiguous=len(group_rows) > 1 and len({r.supplier_name for r in group_rows}) <= 1)
                    search_text = " ".join([
                        row.sheet,
                        row.category,
                        row.family,
                        row.subfamily or "",
                        row.variation,
                        row.original_description,
                        row.code or "",
                        row.manufacturer or "",
                        row.manufacturer_code or "",
                        supplier_name or "",
                        row.email or "",
                        row.phone or "",
                        " ".join(str(v) for v in (row.attributes or {}).values() if v),
                    ])
                    catalog_row = PurchaseXlsxCatalogRow(
                        import_run_id=run.id,
                        row_key=row.row_key,
                        sheet=row.sheet,
                        row_number=row.row_number,
                        parser_type=row.parser_type,
                        category=row.category,
                        family=row.family,
                        subfamily=row.subfamily,
                        group_name=row.subfamily,
                        variation=row.variation,
                        display_name=row.display_name,
                        original_description=row.original_description,
                        code=row.code,
                        manufacturer=row.manufacturer,
                        manufacturer_code=row.manufacturer_code,
                        product_item_id=item.id,
                        current_price=price,
                        raw_price=row.raw_price,
                        final_value=row.final_value,
                        price_type=row.price_type,
                        unit=row.unit,
                        source_updated_at=row.updated_at,
                        situation=situation,
                        issues=issues,
                        attributes=row.attributes,
                        technical_details={
                            "parser_type": row.parser_type,
                            "price_rule": "ultimo_por_data" if row is current_row and _parse_observed_sort(row.updated_at) != datetime.min else "menor_valor_final_ou_preco",
                            "origem": "Compras Nova.xlsx",
                            "aba": row.sheet,
                            "linha": row.row_number,
                            "row_key": row.row_key,
                        },
                        search_text=normalize_search_text(search_text),
                    )
                    db.add(catalog_row)
                    db.flush()
                    catalog_rows.append(catalog_row)

                    offer = PurchaseSupplierPriceOffer(
                        catalog_row_id=catalog_row.id,
                        product_item_id=item.id,
                        supplier_id=supplier.id if supplier else None,
                        supplier_name_snapshot=supplier_name,
                        email=row.email,
                        phone=row.phone,
                        raw_price=row.raw_price,
                        ipi=row.ipi,
                        adjustment=row.adjustment,
                        final_value=price,
                        currency="BRL",
                        unit=row.unit,
                        observed_at=row.updated_at,
                        source_row_key=row.row_key,
                        is_current_supplier=row is current_row,
                        is_consolidated=row is current_row,
                        status=situation,
                    )
                    db.add(offer)
                    offers.append(offer)
                    if supplier and price is not None:
                        PurchasesXlsxSmartCatalog._set_active_reference(
                            db,
                            item,
                            supplier.id,
                            price,
                            getattr(current_user, "id", 1),
                            source_row_key=row.row_key,
                            notes="Preco por fornecedor importado da planilha de compras.",
                        )

                if current_supplier:
                    for catalog_row in catalog_rows:
                        if catalog_row.product_item_id == item.id:
                            catalog_row.current_supplier_id = current_supplier.id
                current_price = _price_for_offer(current_row)
                if current_price is not None:
                    PurchasesXlsxSmartCatalog._set_active_reference(
                        db,
                        item,
                        None,
                        current_price,
                        getattr(current_user, "id", 1),
                        source_row_key=current_row.row_key,
                        notes="Preco atual da empresa definido pelo snapshot da planilha de compras.",
                    )
                db.add(PurchaseCatalogSearchIndex(
                    product_item_id=item.id,
                    family_id=family.id,
                    catalog_row_id=None,
                    category=base.category,
                    family=family.name,
                    variation=base.variation,
                    supplier_names=sorted(set(group_supplier_names)),
                    search_text=" ".join([base.category, family.name, base.variation, " ".join(group_supplier_names)]),
                    normalized_search_text=normalize_search_text(" ".join([
                        base.category,
                        family.name,
                        base.variation,
                        " ".join(group_supplier_names),
                        " ".join(row.original_description for row in group_rows),
                        " ".join(str(v) for row in group_rows for v in (row.attributes or {}).values() if v),
                    ])),
                    rank_hint=len(group_rows),
                ))

            contacts_registry = meta.get("contacts_registry") or {}
            for comp_norm, entry in contacts_registry.items():
                comp_name = entry["name"]
                email = entry["email"]
                phone = entry["phone"]
                PurchasesXlsxSmartCatalog._ensure_supplier(db, comp_name, email, phone, "Importado da planilha")

            run.status = "success"
            run.finished_at = datetime.now(timezone.utc)
            run.metrics = {
                "parser_version": SMART_CATALOG_PARSER_VERSION,
                "source_path": str(source),
                "sheets_ignored": meta.get("sheets_ignored", []),
                "layouts": meta.get("layouts", {}),
                "parsed_rows": len(rows),
                "catalog_rows": len(catalog_rows),
                "variations": len(items_by_key),
                "offers": len(offers),
                "suppliers_found": len(supplier_names),
                "suppliers_with_email": len(supplier_email_names),
                "needs_review": sum(1 for row in catalog_rows if row.situation not in {"Pronto", "Sem e-mail"}),
                "suppliers_without_email": len(supplier_names - supplier_email_names),
            }
            db.commit()
            return PurchasesXlsxSmartCatalog.sync_summary(db, run)
        except Exception as exc:
            run.status = "failed"
            run.finished_at = datetime.now(timezone.utc)
            run.error_message = str(exc)
            db.commit()
            raise

    @staticmethod
    def ensure_snapshot(db: Session, current_user: Any | None = None) -> None:
        exists = db.query(func.count(PurchaseXlsxCatalogRow.id)).scalar() or 0
        last_run = _latest_import_run(db)
        parser_version = (last_run.metrics or {}).get("parser_version") if last_run else None
        if exists and parser_version == SMART_CATALOG_PARSER_VERSION:
            return
        if current_user is None:
            return
        PurchasesXlsxSmartCatalog.sync_from_xlsx(db, current_user)

    @staticmethod
    def sync_summary(db: Session, run: PurchaseXlsxImportRun | None = None) -> dict[str, Any]:
        run = run or _latest_import_run(db)
        summary = PurchasesXlsxSmartCatalog.summary(db)
        return {
            "status": run.status if run else "not_synced",
            "message": "Catalogo de compras atualizado da planilha." if run and run.status == "success" else "Catalogo ainda nao sincronizado.",
            "run_id": run.id if run else None,
            "source_path": run.source_path if run else str(PurchasesXlsxReconciliation.source_path()),
            "started_at": run.started_at if run else None,
            "finished_at": run.finished_at if run else None,
            "metrics": run.metrics or {} if run else {},
            "summary": summary,
        }

    @staticmethod
    def summary(db: Session) -> dict[str, Any]:
        last_run = _latest_import_run(db)
        rows_count = db.query(func.count(PurchaseXlsxCatalogRow.id)).scalar() or 0
        offers_count = db.query(func.count(PurchaseSupplierPriceOffer.id)).scalar() or 0
        suppliers_without_email = db.query(func.count(PurchaseSupplierPriceOffer.id)).filter(
            or_(PurchaseSupplierPriceOffer.email == None, PurchaseSupplierPriceOffer.email == "")
        ).scalar() or 0
        needs_review = db.query(func.count(PurchaseXlsxCatalogRow.id)).filter(
            PurchaseXlsxCatalogRow.situation.notin_(["Pronto", "Sem e-mail"])
        ).scalar() or 0
        return {
            "families_count": db.query(func.count(ProductFamily.id)).scalar() or 0,
            "variations_count": db.query(func.count(PurchaseCatalogSearchIndex.product_item_id.distinct())).scalar() or 0,
            "catalog_rows_count": rows_count,
            "supplier_offers_count": offers_count,
            "current_prices_count": db.query(func.count(PurchasePriceReference.id)).filter(PurchasePriceReference.is_active == True).scalar() or 0,
            "needs_review_count": needs_review,
            "suppliers_without_email_count": suppliers_without_email,
            "last_sync_at": last_run.finished_at if last_run else None,
            "last_sync_status": last_run.status if last_run else "not_synced",
        }

    @staticmethod
    def families(db: Session, *, search: str | None = None, limit: int = 80, offset: int = 0) -> dict[str, Any]:
        limit = max(1, min(limit, 120))
        offset = max(0, offset)
        index_query = db.query(PurchaseCatalogSearchIndex)
        if search:
            matches = [
                idx.id
                for idx in index_query.all()
                if search_matches(idx.normalized_search_text, search)
            ]
            index_query = db.query(PurchaseCatalogSearchIndex).filter(PurchaseCatalogSearchIndex.id.in_(matches or [uuid.uuid4()]))
        indexes = index_query.order_by(desc(PurchaseCatalogSearchIndex.rank_hint), PurchaseCatalogSearchIndex.family.asc()).all()
        family_ids = []
        for idx in indexes:
            if idx.family_id and idx.family_id not in family_ids:
                family_ids.append(idx.family_id)
        total = len(family_ids)
        page_family_ids = family_ids[offset:offset + limit]
        rows: list[dict[str, Any]] = []
        for family_id in page_family_ids:
            family = db.query(ProductFamily).filter(ProductFamily.id == family_id).first()
            item_ids = [
                value[0]
                for value in db.query(PurchaseCatalogSearchIndex.product_item_id)
                .filter(PurchaseCatalogSearchIndex.family_id == family_id)
                .distinct()
                .all()
                if value[0]
            ]
            if not family:
                continue
            refs = db.query(PurchasePriceReference).filter(
                PurchasePriceReference.product_item_id.in_(item_ids or [uuid.uuid4()]),
                PurchasePriceReference.is_active == True,
                PurchasePriceReference.supplier_id == None,
            ).all()
            offers_supplier_ids = {
                value[0]
                for value in db.query(PurchaseSupplierPriceOffer.supplier_id)
                .filter(PurchaseSupplierPriceOffer.product_item_id.in_(item_ids or [uuid.uuid4()]))
                .distinct()
                .all()
                if value[0]
            }
            rows.append(_serialize_catalog_family({
                "family_id": family.id,
                "family_name": family.name,
                "description": family.description,
                "category": next((idx.category for idx in indexes if idx.family_id == family_id), None),
                "variations_count": len(item_ids),
                "suppliers_count": len(offers_supplier_ids),
                "min_price": min((ref.current_unit_price for ref in refs), default=None),
                "max_price": max((ref.current_unit_price for ref in refs), default=None),
                "last_updated_at": max((ref.updated_at for ref in refs), default=None),
                "variations": [],
            }))
        return {
            "items": rows,
            "total_variations": total,
            "total": total,
            "limit": limit,
            "offset": offset,
            "has_more": offset + limit < total,
            "summary": PurchasesXlsxSmartCatalog.summary(db),
        }

    @staticmethod
    def variations(db: Session, family_id: uuid.UUID, *, search: str | None = None, limit: int = 80, offset: int = 0) -> dict[str, Any]:
        limit = max(1, min(limit, 120))
        offset = max(0, offset)
        indexes = db.query(PurchaseCatalogSearchIndex).filter(PurchaseCatalogSearchIndex.family_id == family_id).all()
        if search:
            indexes = [idx for idx in indexes if search_matches(idx.normalized_search_text, search)]
        item_ids = []
        for idx in indexes:
            if idx.product_item_id and idx.product_item_id not in item_ids:
                item_ids.append(idx.product_item_id)
        total = len(item_ids)
        page_ids = item_ids[offset:offset + limit]
        items = db.query(ProductItem).options(joinedload(ProductItem.family)).filter(ProductItem.id.in_(page_ids or [uuid.uuid4()])).all()
        variations = [PurchasesXlsxSmartCatalog.serialize_variation(db, item) for item in items]
        family = db.query(ProductFamily).filter(ProductFamily.id == family_id).first()
        prices = [v["current_price"] for v in variations if v.get("current_price") is not None]
        return {
            "family_id": family.id if family else family_id,
            "family_name": family.name if family else "Familia",
            "description": family.description if family else None,
            "variations_count": total,
            "suppliers_count": len({v.get("supplier_id") for v in variations if v.get("supplier_id")}),
            "min_price": min(prices) if prices else None,
            "max_price": max(prices) if prices else None,
            "last_updated_at": max((v.get("last_updated_at") for v in variations if v.get("last_updated_at")), default=None),
            "variations": variations,
            "limit": limit,
            "offset": offset,
            "has_more": offset + limit < total,
        }

    @staticmethod
    def serialize_variation(db: Session, item: ProductItem) -> dict[str, Any]:
        ref = db.query(PurchasePriceReference).options(
            joinedload(PurchasePriceReference.supplier).joinedload(Supplier.person)
        ).filter(
            PurchasePriceReference.product_item_id == item.id,
            PurchasePriceReference.is_active == True,
            PurchasePriceReference.supplier_id == None,
        ).order_by(desc(PurchasePriceReference.updated_at)).first()
        current_offer = db.query(PurchaseSupplierPriceOffer).options(
            joinedload(PurchaseSupplierPriceOffer.supplier).joinedload(Supplier.person)
        ).filter(
            PurchaseSupplierPriceOffer.product_item_id == item.id,
            PurchaseSupplierPriceOffer.is_current_supplier == True,
        ).order_by(desc(PurchaseSupplierPriceOffer.updated_at)).first()
        supplier = current_offer.supplier if current_offer else ref.supplier if ref else None
        supplier_name = supplier.person.name if supplier and supplier.person else current_offer.supplier_name_snapshot if current_offer else None
        catalog = db.query(PurchaseXlsxCatalogRow).filter(PurchaseXlsxCatalogRow.product_item_id == item.id).first()
        history_count = db.query(func.count(PurchasePriceHistory.id)).filter(PurchasePriceHistory.product_item_id == item.id).scalar() or 0
        suppliers_count = db.query(func.count(PurchaseSupplierPriceOffer.supplier_id.distinct())).filter(
            PurchaseSupplierPriceOffer.product_item_id == item.id,
            PurchaseSupplierPriceOffer.supplier_id != None,
        ).scalar() or 0
        attrs = item.attributes or {}
        variation_name = attrs.get("variation_name") or item.name
        return {
            "id": item.id,
            "family_id": item.family_id,
            "family_name": item.family.name if item.family else catalog.family if catalog else None,
            "name": item.name,
            "short_name": variation_name,
            "variation_name": variation_name,
            "attributes": attrs,
            "unit_of_measure": item.unit_of_measure,
            "category": item.category,
            "current_price": ref.current_unit_price if ref else None,
            "currency": ref.currency if ref else "BRL",
            "supplier_id": supplier.id if supplier else None,
            "supplier_name": supplier_name,
            "last_updated_at": ref.updated_at if ref else None,
            "history_count": history_count,
            "suppliers_count": suppliers_count,
            "description": item.description,
            "situation": catalog.situation if catalog else "Pronto",
            "needs_review": bool(catalog and catalog.situation not in {"Pronto", "Sem e-mail"}),
            "technical_details": {
                "id": str(item.id),
                "sku": item.sku,
                "canonical_key": item.canonical_key,
                "origem": "Compras Nova.xlsx",
            },
        }

    @staticmethod
    def item_detail(db: Session, item_id: uuid.UUID) -> dict[str, Any]:
        item = db.query(ProductItem).options(joinedload(ProductItem.family)).filter(ProductItem.id == item_id).first()
        if not item:
            raise ValueError("Produto ou variacao nao encontrado.")
        data = PurchasesXlsxSmartCatalog.serialize_variation(db, item)
        data["offers"] = PurchasesXlsxSmartCatalog.supplier_offers(db, item_id)["items"]
        data["catalog_rows"] = [
            {
                "row_key": row.row_key,
                "sheet": row.sheet,
                "row_number": row.row_number,
                "situation": row.situation,
                "issues": row.issues or [],
            }
            for row in db.query(PurchaseXlsxCatalogRow).filter(PurchaseXlsxCatalogRow.product_item_id == item_id).limit(20).all()
        ]
        return data

    @staticmethod
    def supplier_offers(db: Session, item_id: uuid.UUID) -> dict[str, Any]:
        offers = db.query(PurchaseSupplierPriceOffer).options(
            joinedload(PurchaseSupplierPriceOffer.supplier).joinedload(Supplier.person)
        ).filter(PurchaseSupplierPriceOffer.product_item_id == item_id).order_by(
            desc(PurchaseSupplierPriceOffer.is_current_supplier),
            PurchaseSupplierPriceOffer.supplier_name_snapshot.asc(),
        ).all()
        return {"items": [_serialize_offer(offer) for offer in offers], "total": len(offers)}

    @staticmethod
    def update_supplier_offer_price(
        db: Session,
        item_id: uuid.UUID,
        offer_id: uuid.UUID,
        payload: Any,
        current_user: Any,
    ) -> dict[str, Any]:
        offer = db.query(PurchaseSupplierPriceOffer).filter(
            PurchaseSupplierPriceOffer.id == offer_id,
            PurchaseSupplierPriceOffer.product_item_id == item_id,
        ).first()
        if not offer:
            raise ValueError("Oferta de fornecedor nao encontrada para esta variacao.")
        item = db.query(ProductItem).options(joinedload(ProductItem.family)).filter(ProductItem.id == item_id).first()
        if not item:
            raise ValueError("Produto ou variacao nao encontrado.")
        new_price = _money(getattr(payload, "new_price", None))
        if new_price is None or new_price <= 0:
            raise ValueError("Informe um novo preco maior que zero.")
        old_price = _price_for_offer_obj(offer)
        now = datetime.now(timezone.utc)
        offer.final_value = new_price
        offer.updated_at = now
        offer.status = "Pronto"
        for other in db.query(PurchaseSupplierPriceOffer).filter(PurchaseSupplierPriceOffer.product_item_id == item_id).all():
            other.is_current_supplier = other.id == offer.id
            other.is_consolidated = other.id == offer.id

        PurchasesXlsxSmartCatalog._set_active_reference(
            db,
            item,
            offer.supplier_id,
            new_price,
            getattr(current_user, "id", 1),
            source_row_key=offer.source_row_key or str(offer.id),
            notes=getattr(payload, "notes", None) or "Preco atualizado na oferta do fornecedor.",
        )
        company_ref = PurchasesXlsxSmartCatalog._set_active_reference(
            db,
            item,
            None,
            new_price,
            getattr(current_user, "id", 1),
            source_row_key=offer.source_row_key or str(offer.id),
            notes="Preco atual da empresa atualizado a partir da oferta do fornecedor.",
        )
        difference_amount = None
        difference_percent = None
        if old_price is not None:
            difference_amount = (new_price - old_price).quantize(Decimal("0.01"))
            if old_price > 0:
                difference_percent = ((difference_amount / old_price) * Decimal("100")).quantize(Decimal("0.0001"))
        db.commit()
        supplier_name = offer.supplier.person.name if offer.supplier and offer.supplier.person else offer.supplier_name_snapshot
        return {
            "item_id": item.id,
            "family_name": item.family.name if item.family else None,
            "variation_name": (item.attributes or {}).get("variation_name") or item.name,
            "old_price": old_price,
            "new_price": new_price,
            "difference_amount": difference_amount,
            "difference_percent": difference_percent,
            "supplier_id": offer.supplier_id,
            "supplier_name": supplier_name,
            "offer_id": offer.id,
            "reference_id": company_ref.id,
            "updated_at": now,
        }

    @staticmethod
    def reconciliation_payload(
        db: Session,
        *,
        search: str | None = None,
        status_filter: str = "all",
        limit: int = 80,
        offset: int = 0,
        include_technical: bool = False,
    ) -> dict[str, Any]:
        rows_query = db.query(PurchaseXlsxCatalogRow).options(
            joinedload(PurchaseXlsxCatalogRow.current_supplier).joinedload(Supplier.person)
        )
        rows = rows_query.order_by(PurchaseXlsxCatalogRow.sheet, PurchaseXlsxCatalogRow.row_number).all()
        if search:
            rows = [row for row in rows if search_matches(row.search_text or " ".join([row.family, row.variation]), search)]
        if status_filter and status_filter != "all":
            mapping = {
                "divergent": "Preço diferente",
                "missing": "Faltando no Portal",
                "no_supplier": "Sem fornecedor",
                "no_email": "Sem e-mail",
                "ambiguous": "Variação ambígua",
                "price_diff": "Preço diferente",
                "ready": "Pronto",
                "review": "Precisa revisão",
            }
            expected = mapping.get(status_filter, status_filter)
            rows = [row for row in rows if row.situation == expected or expected in (row.issues or [])]
        total = len(rows)
        limit = max(1, min(limit, 200))
        offset = max(0, offset)
        page = rows[offset:offset + limit]
        last_run = _latest_import_run(db)
        summary = PurchasesXlsxSmartCatalog.summary(db)
        items = []
        for row in page:
            supplier = row.current_supplier
            supplier_name = supplier.person.name if supplier and supplier.person else None
            ref = db.query(PurchasePriceReference).filter(
                PurchasePriceReference.product_item_id == row.product_item_id,
                PurchasePriceReference.supplier_id == None,
                PurchasePriceReference.is_active == True,
            ).order_by(desc(PurchasePriceReference.updated_at)).first()
            data = {
                "row_key": row.row_key,
                "sheet": row.sheet,
                "row_number": row.row_number,
                "category": row.category,
                "family": row.family,
                "subfamily": row.subfamily,
                "variation": row.variation,
                "display_name": row.display_name,
                "supplier_name": supplier_name,
                "company_name": supplier_name,
                "email": supplier.preferred_contact_email if supplier else None,
                "phone": supplier.person.phone if supplier and supplier.person else None,
                "spreadsheet_price": _decimal_to_json(row.current_price or row.final_value or row.raw_price),
                "raw_price": _decimal_to_json(row.raw_price),
                "final_value": _decimal_to_json(row.final_value),
                "portal_price": _decimal_to_json(ref.current_unit_price if ref else None),
                "difference_amount": None,
                "difference_percent": None,
                "price_type": row.price_type,
                "unit": row.unit,
                "updated_at": row.source_updated_at,
                "situation": row.situation,
                "issues": row.issues or [],
                "portal_item_id": str(row.product_item_id) if row.product_item_id else None,
                "portal_supplier_id": str(row.current_supplier_id) if row.current_supplier_id else None,
                "portal_item_name": row.display_name,
                "portal_supplier_name": supplier_name,
                "ready_for_quote": bool(supplier and supplier.preferred_contact_email and row.situation in {"Pronto", "Sem e-mail"}),
            }
            if include_technical:
                data["technical_details"] = row.technical_details or {}
            items.append(data)
        return {
            "source_path": last_run.source_path if last_run else str(PurchasesXlsxReconciliation.source_path()),
            "file_updated_at": last_run.file_modified_at.isoformat() if last_run and last_run.file_modified_at else None,
            "sheets_read": (last_run.metrics or {}).get("sheets_read", []) if last_run else [],
            "sheets_ignored": (last_run.metrics or {}).get("sheets_ignored", []) if last_run else [],
            "layouts": (last_run.metrics or {}).get("layouts", {}) if last_run else {},
            "summary": {
                "sheets_read": len((last_run.metrics or {}).get("sheets_read", [])) if last_run else 0,
                "suppliers_found": (last_run.metrics or {}).get("suppliers_found", 0) if last_run else 0,
                "suppliers_with_email": (last_run.metrics or {}).get("suppliers_with_email", 0) if last_run else 0,
                "suppliers_with_phone": db.query(func.count(Supplier.id)).join(Person).filter(Person.phone != None).scalar() or 0,
                "suppliers_in_portal": db.query(func.count(Supplier.id)).scalar() or 0,
                "suppliers_missing": 0,
                "products_found": summary["catalog_rows_count"],
                "families_found": summary["families_count"],
                "variations_found": summary["variations_count"],
                "ambiguous_items": db.query(func.count(PurchaseXlsxCatalogRow.id)).filter(PurchaseXlsxCatalogRow.situation == "Variação ambígua").scalar() or 0,
                "prices_detected": summary["supplier_offers_count"],
                "prices_using_final_value": db.query(func.count(PurchaseXlsxCatalogRow.id)).filter(PurchaseXlsxCatalogRow.final_value != None).scalar() or 0,
                "prices_matching_portal": summary["current_prices_count"],
                "prices_divergent": db.query(func.count(PurchaseXlsxCatalogRow.id)).filter(PurchaseXlsxCatalogRow.situation == "Preço diferente").scalar() or 0,
                "needs_review": summary["needs_review_count"],
                "suppliers_without_email": summary["suppliers_without_email_count"],
            },
            "items": items,
            "total": total,
            "limit": limit,
            "offset": offset,
            "has_more": offset + limit < total,
        }


def _price_for_offer_obj(offer: PurchaseSupplierPriceOffer) -> Decimal | None:
    value = offer.final_value or offer.raw_price
    if value is None:
        return None
    return Decimal(str(value)).quantize(Decimal("0.01"))
