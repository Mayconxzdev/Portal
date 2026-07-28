import hashlib
import csv
import io
import json
import os
import re
import tempfile
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

from fastapi import HTTPException, UploadFile, status
from fastapi.responses import FileResponse as DownloadFileResponse
from sqlalchemy import func, or_
from sqlalchemy.orm import Session
from sqlalchemy.orm.attributes import set_committed_value

from app.core.permissions import PermissionLevel
from app.models.kanban import (
    File,
    FileModuleLink,
    KanbanActivity,
    KanbanBoard,
    KanbanBoardPermission,
    KanbanBoardView,
    KanbanCard,
    KanbanCardAssignee,
    KanbanCardAttachment,
    KanbanCardChecklist,
    KanbanCardChecklistItem,
    KanbanCardComment,
    KanbanCardLabel,
    KanbanColumn,
    KanbanCustomField,
    KanbanLabel,
    KanbanTVView,
)
from app.models.user import User
from app.modules.kanban.activity import record_activity
from app.modules.kanban.permissions import (
    BOARD_LEVEL_VALUES,
    BoardAccessLevel,
    get_board_access_level,
    require_board_level,
    require_module_level,
)
from app.modules.kanban.repository import get_board, get_card, get_column, get_custom_field
from app.modules.kanban.schemas import (
    CARD_PRIORITIES,
    DENSITIES,
    FIELD_TYPES,
    VIEW_TYPES,
    BoardCreate,
    BoardDuplicatePayload,
    BoardImportConfirmPayload,
    BoardPermissionUpsert,
    BoardViewCreate,
    BoardViewUpdate,
    BoardUpdate,
    CardCreate,
    CardDuplicatePayload,
    CardMovePayload,
    CardUpdate,
    ChecklistCreate,
    ChecklistItemCreate,
    ChecklistItemReorderPayload,
    ChecklistItemUpdate,
    ChecklistUpdate,
    ColumnCreate,
    ColumnReorderPayload,
    ColumnUpdate,
    CommentCreate,
    CommentUpdate,
    CustomFieldCreate,
    CustomFieldUpdate,
    LabelCreate,
    LabelUpdate,
    QuickCardPayload,
    TVConfigUpdate,
)


def slugify(value: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", value.strip().lower()).strip("-")
    return slug or "board"


UPLOAD_DIR = Path(__file__).resolve().parents[4] / "storage" / "uploads"
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
ALLOWED_EXTENSIONS = {
    ".pdf", ".png", ".jpg", ".jpeg", ".webp", ".gif", ".txt", ".csv",
    ".xls", ".xlsx", ".doc", ".docx", ".ppt", ".pptx", ".zip"
}
BLOCKED_EXTENSIONS = {".exe", ".bat", ".cmd", ".ps1", ".sh", ".js", ".msi", ".dll", ".scr", ".vbs"}
TV_LAYOUTS = {"COLUMNS", "URGENCY", "PRODUCTION", "COMPACT", "TV_LIST", "PRODUCTION_LIST"}
PRODUCTION_CUSTOM_KEYS = ["op", "cliente", "modelo", "tensao", "qtd", "inicio", "entrega", "setor", "pendencia", "etapa", "material"]
IMPORT_DIR = Path(tempfile.gettempdir()) / "portal-vesper-imports"


class KanbanService:
    @staticmethod
    def _require_global_admin(current_user: User) -> None:
        if not (current_user.role and current_user.role.name == "ADMIN"):
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Permissoes de quadros sao gerenciadas na Administracao por ADMIN global.")

    @staticmethod
    def decorate_card(card: KanbanCard) -> KanbanCard:
        card.labels = [link.label for link in getattr(card, "label_links", []) if link.label and link.label.is_active]
        card.checklist_total = sum(len(checklist.items) for checklist in getattr(card, "checklists", []))
        card.checklist_done = sum(1 for checklist in getattr(card, "checklists", []) for item in checklist.items if item.is_done)
        return card

    @staticmethod
    def attach_access(board: KanbanBoard, db: Session, user: User) -> KanbanBoard:
        board.access_level = get_board_access_level(db, board, user).value
        board.columns = sorted(board.columns, key=lambda col: col.position)
        set_committed_value(board, "custom_fields", sorted([field for field in board.custom_fields if field.is_active], key=lambda field: field.position))
        set_committed_value(board, "labels", sorted([label for label in board.labels if label.is_active], key=lambda label: label.name))
        board.views = sorted(getattr(board, "views", []), key=lambda view: (not view.is_default, view.view_type, view.name))
        for column in board.columns:
            column.cards = sorted(column.cards, key=lambda card: card.position)
            for card in column.cards:
                KanbanService.decorate_card(card)
        return board

    @staticmethod
    def list_boards(db: Session, current_user: User, include_archived: bool = False) -> List[KanbanBoard]:
        require_module_level(db, current_user, PermissionLevel.READ_ONLY)
        query = db.query(KanbanBoard)
        if not include_archived:
            query = query.filter(KanbanBoard.is_archived == False)

        if current_user.role and current_user.role.name == "ADMIN":
            boards = query.order_by(KanbanBoard.name).all()
            return [KanbanService.attach_access(board, db, current_user) for board in boards]

        permission_board_ids = [
            row[0] for row in db.query(KanbanBoardPermission.board_id).filter(
                or_(
                    KanbanBoardPermission.user_id == current_user.id,
                    KanbanBoardPermission.role_id == current_user.role_id if current_user.role_id else False,
                ),
                KanbanBoardPermission.access_level != "NO_ACCESS",
            ).all()
        ]

        boards = query.filter(
            or_(
                KanbanBoard.created_by_user_id == current_user.id,
                KanbanBoard.id.in_(permission_board_ids) if permission_board_ids else False,
            )
        ).order_by(KanbanBoard.name).all()
        return [KanbanService.attach_access(board, db, current_user) for board in boards]

    @staticmethod
    def default_view_payloads(board: KanbanBoard, preset: Optional[str] = None) -> List[Dict[str, Any]]:
        return [
            {"name": "Quadro Geral", "view_type": "BOARD", "is_default": True, "visible_columns": ["title", "op", "cliente", "modelo", "entrega", "responsaveis", "prioridade"]},
            {"name": "Lista Geral", "view_type": "LIST", "is_default": False, "density": "COMPACT", "visible_columns": ["op", "title", "status", "cliente", "modelo", "tensao", "qtd", "inicio", "entrega", "setor", "pendencia", "prioridade", "responsaveis"], "sort_by": "entrega"},
            {"name": "TV Lista", "view_type": "TV_LIST", "is_default": False, "density": "DENSE", "visible_columns": ["op", "cliente", "modelo", "qtd", "entrega", "setor", "pendencia"], "sort_by": "urgency_score", "font_scale": "large", "auto_scroll": True, "auto_scroll_seconds": 25},
            {"name": "TV Quadro", "view_type": "TV_BOARD", "is_default": False, "visible_columns": ["op", "cliente", "modelo", "entrega", "etapa"], "sort_by": "urgency_score"},
        ]

    @staticmethod
    def ensure_default_views(db: Session, board: KanbanBoard, user: User, preset: Optional[str] = None) -> None:
        if db.query(KanbanBoardView).filter(KanbanBoardView.board_id == board.id).first():
            return
        now = datetime.now(timezone.utc)
        for payload in KanbanService.default_view_payloads(board, preset):
            db.add(KanbanBoardView(
                board_id=board.id,
                created_by_user_id=user.id,
                created_at=now,
                updated_at=now,
                **payload,
            ))

    @staticmethod
    def ensure_production_fields(db: Session, board: KanbanBoard) -> None:
        production_fields = [
            ("OP", "op", "TEXT"),
            ("Cliente", "cliente", "TEXT"),
            ("Modelo", "modelo", "TEXT"),
            ("Tensao", "tensao", "TEXT"),
            ("Qtd", "qtd", "NUMBER"),
            ("Inicio", "inicio", "DATE"),
            ("Entrega", "entrega", "DATE"),
            ("Setor", "setor", "TEXT"),
            ("Pendencia", "pendencia", "TEXT"),
            ("Etapa", "etapa", "TEXT"),
            ("Material", "material", "TEXT"),
        ]
        for position, (name, key, field_type) in enumerate(production_fields):
            existing = db.query(KanbanCustomField).filter(KanbanCustomField.board_id == board.id, KanbanCustomField.key == key).first()
            if not existing:
                db.add(KanbanCustomField(
                    board_id=board.id,
                    name=name,
                    key=key,
                    field_type=field_type,
                    position=position,
                    options={"display": {"detail": True, "board": position < 5, "list": True, "tv": True, "filter": True}},
                ))

    @staticmethod
    def create_board(db: Session, payload: BoardCreate, current_user: User) -> KanbanBoard:
        require_module_level(db, current_user, PermissionLevel.NORMAL)
        slug = slugify(payload.slug or payload.name)
        if db.query(KanbanBoard).filter(KanbanBoard.slug == slug).first():
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Ja existe um board com este slug.")

        board = KanbanBoard(
            name=payload.name,
            slug=slug,
            description=payload.description,
            module_origin=payload.module_origin,
            color=payload.color,
            icon=payload.icon,
            created_by_user_id=current_user.id,
        )
        db.add(board)
        db.flush()

        db.add(KanbanBoardPermission(board_id=board.id, user_id=current_user.id, access_level="ADMIN"))
        defaults = [
            ("A Fazer", 0, False),
            ("Em andamento", 1, False),
            ("Concluido", 2, True),
        ]
        for name, position, is_done in defaults:
            db.add(KanbanColumn(board_id=board.id, name=name, position=position, is_done_column=is_done))

        # Criar etiquetas padrao
        default_labels = [
            ("Urgente", "#ef4444"),
            ("Alta", "#f59e0b"),
            ("M\u00e9dio", "#3b82f6"),
            ("Baixa", "#10b981"),
        ]
        for name, color in default_labels:
            db.add(KanbanLabel(board_id=board.id, name=name, color=color, is_active=True))

        # Sempre cria os campos de producao por padrao no modelo completo
        KanbanService.ensure_production_fields(db, board)
        KanbanService.ensure_default_views(db, board, current_user, payload.preset)
        record_activity(db, board.id, current_user, "board.created", metadata={"name": board.name})
        db.commit()
        refreshed = get_board(db, board.id)
        return KanbanService.attach_access(refreshed, db, current_user)

    @staticmethod
    def get_board_detail(db: Session, board_id: int, current_user: User) -> KanbanBoard:
        board = get_board(db, board_id)
        if not board:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Board nao encontrado.")
        require_board_level(db, board, current_user, BoardAccessLevel.READ_ONLY)
        return KanbanService.attach_access(board, db, current_user)

    @staticmethod
    def update_board(db: Session, board_id: int, payload: BoardUpdate, current_user: User) -> KanbanBoard:
        board = KanbanService.get_board_detail(db, board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.MANAGER)
        for field, value in payload.model_dump(exclude_unset=True).items():
            setattr(board, field, value)
        board.updated_at = datetime.now(timezone.utc)
        record_activity(db, board.id, current_user, "board.updated", metadata=payload.model_dump(exclude_unset=True))
        db.commit()
        return KanbanService.get_board_detail(db, board_id, current_user)

    @staticmethod
    def set_board_archived(db: Session, board_id: int, archived: bool, current_user: User) -> KanbanBoard:
        board = KanbanService.get_board_detail(db, board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.MANAGER)
        board.is_archived = archived
        board.updated_at = datetime.now(timezone.utc)
        record_activity(db, board.id, current_user, "board.archived" if archived else "board.restored")
        db.commit()
        return KanbanService.get_board_detail(db, board_id, current_user)

    @staticmethod
    def duplicate_board(db: Session, board_id: int, payload: BoardDuplicatePayload, current_user: User) -> KanbanBoard:
        source = KanbanService.get_board_detail(db, board_id, current_user)
        require_board_level(db, source, current_user, BoardAccessLevel.MANAGER)
        slug = slugify(payload.slug or payload.name)
        if db.query(KanbanBoard).filter(KanbanBoard.slug == slug).first():
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Ja existe um quadro com este identificador.")
        target = KanbanBoard(
            name=payload.name,
            slug=slug,
            description=source.description,
            module_origin=source.module_origin,
            color=source.color,
            icon=source.icon,
            created_by_user_id=current_user.id,
        )
        db.add(target)
        db.flush()
        db.add(KanbanBoardPermission(board_id=target.id, user_id=current_user.id, access_level="ADMIN"))
        column_map: Dict[int, int] = {}
        for column in sorted(source.columns, key=lambda item: item.position):
            cloned = KanbanColumn(
                board_id=target.id,
                name=column.name,
                position=column.position,
                color=column.color,
                wip_limit=column.wip_limit,
                is_done_column=column.is_done_column,
                is_archived=column.is_archived,
            )
            db.add(cloned)
            db.flush()
            column_map[column.id] = cloned.id
        if payload.copy_custom_fields:
            for field in source.custom_fields:
                db.add(KanbanCustomField(
                    board_id=target.id,
                    name=field.name,
                    key=field.key,
                    field_type=field.field_type,
                    options=field.options,
                    is_required=field.is_required,
                    is_active=field.is_active,
                    position=field.position,
                ))
        label_map: Dict[int, int] = {}
        if payload.copy_labels:
            for label in source.labels:
                cloned_label = KanbanLabel(board_id=target.id, name=label.name, color=label.color, is_active=label.is_active)
                db.add(cloned_label)
                db.flush()
                label_map[label.id] = cloned_label.id
        if payload.copy_views:
            for view in getattr(source, "views", []):
                db.add(KanbanBoardView(
                    board_id=target.id,
                    name=view.name,
                    view_type=view.view_type,
                    is_default=view.is_default,
                    density=view.density,
                    visible_columns=view.visible_columns,
                    column_order=view.column_order,
                    column_widths=view.column_widths,
                    filters=view.filters,
                    sort_by=view.sort_by,
                    group_by=view.group_by,
                    color_rules=view.color_rules,
                    font_scale=view.font_scale,
                    auto_scroll=view.auto_scroll,
                    auto_scroll_seconds=view.auto_scroll_seconds,
                    created_by_user_id=current_user.id,
                ))
        else:
            KanbanService.ensure_default_views(db, target, current_user)
        if payload.copy_cards:
            for card in db.query(KanbanCard).filter(KanbanCard.board_id == source.id).order_by(KanbanCard.position).all():
                cloned_card = KanbanCard(
                    board_id=target.id,
                    column_id=column_map[card.column_id],
                    title=card.title,
                    description=card.description,
                    position=card.position,
                    priority=card.priority,
                    status=card.status,
                    due_date=card.due_date,
                    assigned_to_user_id=card.assigned_to_user_id,
                    created_by_user_id=current_user.id,
                    custom_fields=card.custom_fields if payload.copy_custom_fields else None,
                    is_archived=card.is_archived,
                )
                db.add(cloned_card)
                db.flush()
                for link in getattr(card, "label_links", []):
                    if link.label_id in label_map:
                        db.add(KanbanCardLabel(card_id=cloned_card.id, label_id=label_map[link.label_id]))
        record_activity(db, source.id, current_user, "board.duplicated", metadata={"new_board_id": target.id, "new_board_name": target.name})
        record_activity(db, target.id, current_user, "board.created", metadata={"duplicated_from_board_id": source.id})
        db.commit()
        return KanbanService.get_board_detail(db, target.id, current_user)

    @staticmethod
    def delete_board(db: Session, board_id: int, current_user: User) -> Dict[str, str]:
        board = KanbanService.get_board_detail(db, board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.MANAGER)
        record_activity(db, board.id, current_user, "board.delete_blocked", metadata={"reason": "permanent_delete_deferred"})
        db.commit()
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Exclusao permanente de quadros foi adiada por seguranca. Use arquivar/restaurar.")

    @staticmethod
    def create_column(db: Session, board_id: int, payload: ColumnCreate, current_user: User) -> KanbanColumn:
        board = KanbanService.get_board_detail(db, board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.MANAGER)
        position = db.query(func.count(KanbanColumn.id)).filter(KanbanColumn.board_id == board_id, KanbanColumn.is_archived == False).scalar() or 0
        column = KanbanColumn(
            board_id=board_id,
            name=payload.name,
            color=payload.color,
            wip_limit=payload.wip_limit,
            is_done_column=payload.is_done_column,
            position=position,
        )
        db.add(column)
        db.flush()
        record_activity(db, board_id, current_user, "column.created", metadata={"column_id": column.id, "name": column.name})
        db.commit()
        db.refresh(column)
        return column

    @staticmethod
    def update_column(db: Session, column_id: int, payload: ColumnUpdate, current_user: User) -> KanbanColumn:
        column = get_column(db, column_id)
        if not column:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Coluna nao encontrada.")
        board = KanbanService.get_board_detail(db, column.board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.MANAGER)
        for field, value in payload.model_dump(exclude_unset=True).items():
            setattr(column, field, value)
        column.updated_at = datetime.now(timezone.utc)
        record_activity(db, column.board_id, current_user, "column.updated", metadata={"column_id": column.id})
        db.commit()
        db.refresh(column)
        return column

    @staticmethod
    def archive_column(db: Session, column_id: int, current_user: User) -> KanbanColumn:
        column = get_column(db, column_id)
        if not column:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Coluna nao encontrada.")
        board = KanbanService.get_board_detail(db, column.board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.MANAGER)
        active_cards = db.query(func.count(KanbanCard.id)).filter(
            KanbanCard.column_id == column.id,
            KanbanCard.is_archived == False,
        ).scalar() or 0
        if active_cards:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Nao e possivel arquivar uma coluna com cards ativos nesta fase. Mova ou arquive os cards primeiro.",
            )
        column.is_archived = True
        record_activity(db, column.board_id, current_user, "column.archived", metadata={"column_id": column.id})
        db.commit()
        db.refresh(column)
        return column

    @staticmethod
    def restore_column(db: Session, column_id: int, current_user: User) -> KanbanColumn:
        column = get_column(db, column_id)
        if not column:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Coluna nao encontrada.")
        board = KanbanService.get_board_detail(db, column.board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.MANAGER)
        column.is_archived = False
        column.updated_at = datetime.now(timezone.utc)
        record_activity(db, column.board_id, current_user, "column.restored", metadata={"column_id": column.id, "name": column.name})
        db.commit()
        db.refresh(column)
        return column

    @staticmethod
    def delete_column(db: Session, column_id: int, current_user: User) -> Dict[str, str]:
        column = get_column(db, column_id)
        if not column:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Coluna nao encontrada.")
        board = KanbanService.get_board_detail(db, column.board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.MANAGER)
        active_cards = db.query(func.count(KanbanCard.id)).filter(KanbanCard.column_id == column.id, KanbanCard.is_archived == False).scalar() or 0
        any_cards = db.query(func.count(KanbanCard.id)).filter(KanbanCard.column_id == column.id).scalar() or 0
        if active_cards or any_cards:
            raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Nao e possivel excluir coluna com cards. Mova ou arquive os cards e mantenha o historico.")
        metadata = {"column_id": column.id, "name": column.name}
        db.delete(column)
        record_activity(db, board.id, current_user, "column.deleted", metadata=metadata)
        db.commit()
        return {"status": "success"}

    @staticmethod
    def reorder_columns(db: Session, board_id: int, payload: ColumnReorderPayload, current_user: User) -> List[KanbanColumn]:
        board = KanbanService.get_board_detail(db, board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.MANAGER)
        column_ids = [item.column_id for item in payload.columns]
        columns = db.query(KanbanColumn).filter(KanbanColumn.board_id == board_id, KanbanColumn.id.in_(column_ids)).all()
        if len(columns) != len(set(column_ids)):
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="A reordenacao contem coluna invalida.")
        positions = {item.column_id: index for index, item in enumerate(sorted(payload.columns, key=lambda item: item.position))}
        for column in columns:
            column.position = positions[column.id]
        record_activity(db, board_id, current_user, "column.reordered", metadata={"columns": positions})
        db.commit()
        return db.query(KanbanColumn).filter(KanbanColumn.board_id == board_id).order_by(KanbanColumn.position).all()

    @staticmethod
    def normalize_card_positions(db: Session, column_id: int) -> None:
        cards = (
            db.query(KanbanCard)
            .filter(KanbanCard.column_id == column_id, KanbanCard.is_archived == False)
            .order_by(KanbanCard.position, KanbanCard.id)
            .all()
        )
        for index, card in enumerate(cards):
            card.position = index

    @staticmethod
    def validate_custom_fields(db: Session, board_id: int, values: Optional[Dict[str, Any]]) -> Optional[Dict[str, Any]]:
        if values is None:
            return None
        fields = db.query(KanbanCustomField).filter(KanbanCustomField.board_id == board_id, KanbanCustomField.is_active == True).all()
        field_map = {field.key: field for field in fields}
        unknown_keys = set(values.keys()) - set(field_map.keys())
        if unknown_keys:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Campos personalizados invalidos: {', '.join(sorted(unknown_keys))}.")
        missing_required = [field.name for field in fields if field.is_required and field.key not in values]
        if missing_required:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Campos obrigatorios ausentes: {', '.join(missing_required)}.")
        return values

    @staticmethod
    def create_card(db: Session, board_id: int, payload: CardCreate, current_user: User) -> KanbanCard:
        board = KanbanService.get_board_detail(db, board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.NORMAL)
        column = get_column(db, payload.column_id)
        if not column or column.board_id != board_id or column.is_archived:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Coluna invalida para este board.")
        if payload.priority not in CARD_PRIORITIES:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Prioridade invalida.")
        custom_fields = KanbanService.validate_custom_fields(db, board_id, payload.custom_fields)
        position = db.query(func.count(KanbanCard.id)).filter(KanbanCard.column_id == column.id, KanbanCard.is_archived == False).scalar() or 0
        card = KanbanCard(
            board_id=board_id,
            column_id=column.id,
            title=payload.title,
            description=payload.description,
            priority=payload.priority,
            status=payload.status,
            due_date=payload.due_date,
            assigned_to_user_id=payload.assigned_to_user_id,
            created_by_user_id=current_user.id,
            custom_fields=custom_fields,
            position=position,
        )
        db.add(card)
        db.flush()
        record_activity(db, board_id, current_user, "card.created", card_id=card.id, metadata={"title": card.title})
        
        from app.core.events import emit_event
        emit_event(
            db=db,
            event_type="kanban.card.created",
            aggregate_type="kanban_card",
            aggregate_id=str(card.id),
            module="kanban",
            payload={
                "card_id": card.id,
                "board_id": board_id,
                "title": card.title,
                "assigned_to_user_id": card.assigned_to_user_id,
                "created_by_user_id": current_user.id
            },
            actor_user_id=current_user.id
        )
        if card.assigned_to_user_id:
            emit_event(
                db=db,
                event_type="kanban.card.assigned",
                aggregate_type="kanban_card",
                aggregate_id=str(card.id),
                module="kanban",
                payload={
                    "card_id": card.id,
                    "board_id": board_id,
                    "title": card.title,
                    "assigned_user_id": card.assigned_to_user_id,
                    "assigned_by_user_id": current_user.id
                },
                actor_user_id=current_user.id
            )
            
        db.commit()
        db.refresh(card)
        return card

    @staticmethod
    def get_card_detail(db: Session, card_id: int, current_user: User) -> KanbanCard:
        card = get_card(db, card_id)
        if not card:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Card nao encontrado.")
        board = KanbanService.get_board_detail(db, card.board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.READ_ONLY)
        return KanbanService.decorate_card(card)

    @staticmethod
    def update_card(db: Session, card_id: int, payload: CardUpdate, current_user: User) -> KanbanCard:
        card = KanbanService.get_card_detail(db, card_id, current_user)
        board = KanbanService.get_board_detail(db, card.board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.NORMAL)
        values = payload.model_dump(exclude_unset=True)
        if "priority" in values and values["priority"] not in CARD_PRIORITIES:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Prioridade invalida.")
        if "custom_fields" in values:
            values["custom_fields"] = KanbanService.validate_custom_fields(db, card.board_id, values["custom_fields"])
        for field, value in values.items():
            setattr(card, field, value)
        activity_metadata = {}
        for k, v in values.items():
            if isinstance(v, datetime):
                activity_metadata[k] = v.isoformat()
            else:
                activity_metadata[k] = v
        record_activity(db, card.board_id, current_user, "card.updated", card_id=card.id, metadata=activity_metadata)
        db.commit()
        db.refresh(card)
        return card

    @staticmethod
    def move_card(db: Session, card_id: int, payload: CardMovePayload, current_user: User) -> KanbanCard:
        card = KanbanService.get_card_detail(db, card_id, current_user)
        board = KanbanService.get_board_detail(db, card.board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.NORMAL)
        target_column_id = payload.to_column_id or payload.column_id
        target_position = payload.new_position if payload.new_position is not None else payload.position
        if target_column_id is None:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Informe a coluna de destino.")
        target_column = get_column(db, target_column_id)
        if not target_column or target_column.board_id != card.board_id or target_column.is_archived:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Coluna de destino invalida.")
        old_column_id = card.column_id
        old_position = card.position
        old_cards = (
            db.query(KanbanCard)
            .filter(KanbanCard.column_id == old_column_id, KanbanCard.id != card.id, KanbanCard.is_archived == False)
            .order_by(KanbanCard.position, KanbanCard.id)
            .all()
        )
        for index, old_card in enumerate(old_cards):
            old_card.position = index

        target_cards = (
            db.query(KanbanCard)
            .filter(KanbanCard.column_id == target_column.id, KanbanCard.id != card.id, KanbanCard.is_archived == False)
            .order_by(KanbanCard.position, KanbanCard.id)
            .all()
        )
        insert_at = max(0, min(target_position if target_position is not None else len(target_cards), len(target_cards)))
        target_cards.insert(insert_at, card)
        card.column_id = target_column.id
        for index, target_card in enumerate(target_cards):
            target_card.position = index
        card.updated_at = datetime.now(timezone.utc)
        record_activity(
            db,
            card.board_id,
            current_user,
            "card.moved",
            card_id=card.id,
            metadata={
                "from_column_id": old_column_id,
                "to_column_id": target_column.id,
                "from_position": old_position,
                "position": card.position,
            },
        )
        
        from app.core.events import emit_event
        assigned_user_ids = [a.user_id for a in card.assignees] if getattr(card, "assignees", None) else []
        if not assigned_user_ids and card.assigned_to_user_id:
            assigned_user_ids = [card.assigned_to_user_id]
            
        from_col_name = db.query(KanbanColumn.name).filter(KanbanColumn.id == old_column_id).scalar() or "Origem"
        
        emit_event(
            db=db,
            event_type="kanban.card.moved",
            aggregate_type="kanban_card",
            aggregate_id=str(card.id),
            module="kanban",
            payload={
                "card_id": card.id,
                "board_id": card.board_id,
                "title": card.title,
                "from_column_id": old_column_id,
                "to_column_id": target_column.id,
                "from_column_name": from_col_name,
                "to_column_name": target_column.name,
                "assigned_to_user_ids": assigned_user_ids,
                "moved_by_user_id": current_user.id
            },
            actor_user_id=current_user.id
        )
        
        # Verifica transicao para coluna concluida
        old_col = db.query(KanbanColumn).filter(KanbanColumn.id == old_column_id).first()
        old_is_done = old_col.is_done_column if old_col else False
        if target_column.is_done_column and not old_is_done:
            emit_event(
                db=db,
                event_type="kanban.card.completed",
                aggregate_type="kanban_card",
                aggregate_id=str(card.id),
                module="kanban",
                payload={
                    "card_id": card.id,
                    "board_id": card.board_id,
                    "title": card.title,
                    "completed_by_user_id": current_user.id,
                    "assigned_to_user_ids": assigned_user_ids
                },
                actor_user_id=current_user.id
            )
            
        db.commit()
        db.refresh(card)
        return KanbanService.decorate_card(card)

    @staticmethod
    def set_card_archived(db: Session, card_id: int, archived: bool, current_user: User) -> KanbanCard:
        card = KanbanService.get_card_detail(db, card_id, current_user)
        board = KanbanService.get_board_detail(db, card.board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.NORMAL)
        card.is_archived = archived
        card.updated_at = datetime.now(timezone.utc)
        record_activity(db, card.board_id, current_user, "card.archived" if archived else "card.restored", card_id=card.id)
        db.commit()
        db.refresh(card)
        return card

    @staticmethod
    def duplicate_card(db: Session, card_id: int, payload: CardDuplicatePayload, current_user: User) -> KanbanCard:
        source = KanbanService.get_card_detail(db, card_id, current_user)
        board = KanbanService.get_board_detail(db, source.board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.NORMAL)
        target_column_id = payload.column_id or source.column_id
        target_column = get_column(db, target_column_id)
        if not target_column or target_column.board_id != source.board_id or target_column.is_archived:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Coluna de destino invalida.")
        position = db.query(func.count(KanbanCard.id)).filter(KanbanCard.column_id == target_column.id, KanbanCard.is_archived == False).scalar() or 0
        cloned = KanbanCard(
            board_id=source.board_id,
            column_id=target_column.id,
            title=payload.title or f"Copiar - {source.title}",
            description=source.description,
            position=position,
            priority=source.priority,
            status=source.status,
            due_date=source.due_date,
            assigned_to_user_id=source.assigned_to_user_id if payload.copy_assignees else None,
            created_by_user_id=current_user.id,
            custom_fields=source.custom_fields if payload.copy_custom_fields else None,
        )
        db.add(cloned)
        db.flush()
        if payload.copy_labels:
            for link in getattr(source, "label_links", []):
                db.add(KanbanCardLabel(card_id=cloned.id, label_id=link.label_id))
        if payload.copy_assignees:
            for assignee in getattr(source, "assignees", []):
                db.add(KanbanCardAssignee(card_id=cloned.id, user_id=assignee.user_id, assigned_by_user_id=current_user.id))
        if payload.copy_checklist:
            for checklist in getattr(source, "checklists", []):
                cloned_checklist = KanbanCardChecklist(
                    card_id=cloned.id,
                    title=checklist.title,
                    position=checklist.position,
                    created_by_user_id=current_user.id,
                )
                db.add(cloned_checklist)
                db.flush()
                for item in checklist.items:
                    db.add(KanbanCardChecklistItem(
                        checklist_id=cloned_checklist.id,
                        text=item.text,
                        is_done=False,
                        position=item.position,
                        created_by_user_id=current_user.id,
                    ))
        record_activity(db, source.board_id, current_user, "card.duplicated", card_id=source.id, metadata={"new_card_id": cloned.id, "title": cloned.title})
        db.commit()
        db.refresh(cloned)
        return KanbanService.decorate_card(cloned)

    @staticmethod
    def delete_card(db: Session, card_id: int, current_user: User) -> Dict[str, str]:
        card = KanbanService.get_card_detail(db, card_id, current_user)
        board = KanbanService.get_board_detail(db, card.board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.MANAGER)
        record_activity(db, board.id, current_user, "card.delete_blocked", card_id=card.id, metadata={"reason": "permanent_delete_deferred"})
        db.commit()
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Exclusao permanente de cards foi adiada por seguranca. Use arquivar/restaurar.")

    @staticmethod
    def list_custom_fields(db: Session, board_id: int, current_user: User) -> List[KanbanCustomField]:
        board = KanbanService.get_board_detail(db, board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.READ_ONLY)
        return db.query(KanbanCustomField).filter(KanbanCustomField.board_id == board_id).order_by(KanbanCustomField.position).all()

    @staticmethod
    def create_custom_field(db: Session, board_id: int, payload: CustomFieldCreate, current_user: User) -> KanbanCustomField:
        board = KanbanService.get_board_detail(db, board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.MANAGER)
        if payload.field_type not in FIELD_TYPES:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Tipo de campo personalizado invalido.")
        if db.query(KanbanCustomField).filter(KanbanCustomField.board_id == board_id, KanbanCustomField.key == payload.key).first():
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Ja existe um campo com esta chave neste board.")
        position = payload.position
        if position is None:
            position = db.query(func.count(KanbanCustomField.id)).filter(KanbanCustomField.board_id == board_id).scalar() or 0
        field = KanbanCustomField(
            board_id=board_id,
            name=payload.name,
            key=payload.key,
            field_type=payload.field_type,
            options={**(payload.options or {}), **({"display": payload.display} if payload.display else {})} or None,
            is_required=payload.is_required,
            position=position,
        )
        db.add(field)
        db.flush()
        record_activity(db, board_id, current_user, "custom_field.created", metadata={"field_id": field.id, "key": field.key})
        db.commit()
        db.refresh(field)
        return field

    @staticmethod
    def update_custom_field(db: Session, field_id: int, payload: CustomFieldUpdate, current_user: User) -> KanbanCustomField:
        field = get_custom_field(db, field_id)
        if not field:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Campo personalizado nao encontrado.")
        board = KanbanService.get_board_detail(db, field.board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.MANAGER)
        values = payload.model_dump(exclude_unset=True)
        if "key" in values and values["key"] != field.key:
            used = db.query(KanbanCard).filter(KanbanCard.board_id == field.board_id).all()
            has_usage = any((card.custom_fields or {}).get(field.key) is not None for card in used)
            if has_usage:
                raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Nao e possivel alterar a chave de um campo que ja possui dados.")
            if db.query(KanbanCustomField).filter(KanbanCustomField.board_id == field.board_id, KanbanCustomField.key == values["key"], KanbanCustomField.id != field.id).first():
                raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Ja existe um campo com esta chave neste quadro.")
        if "field_type" in values and values["field_type"] not in FIELD_TYPES:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Tipo de campo personalizado invalido.")
        if "field_type" in values and values["field_type"] != field.field_type:
            used = db.query(KanbanCard).filter(KanbanCard.board_id == field.board_id).all()
            has_usage = any((card.custom_fields or {}).get(field.key) is not None for card in used)
            if has_usage:
                raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Nao e possivel mudar o tipo de um campo que ja possui dados.")
        if "display" in values:
            options = dict(field.options or {})
            options["display"] = values.pop("display")
            values["options"] = options
        for name, value in values.items():
            setattr(field, name, value)
        field.updated_at = datetime.now(timezone.utc)
        record_activity(db, field.board_id, current_user, "custom_field.updated", metadata={"field_id": field.id})
        db.commit()
        db.refresh(field)
        return field

    @staticmethod
    def delete_custom_field(db: Session, field_id: int, current_user: User) -> Dict[str, Any]:
        field = get_custom_field(db, field_id)
        if not field:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Campo personalizado nao encontrado.")
        board = KanbanService.get_board_detail(db, field.board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.MANAGER)
        used = db.query(KanbanCard).filter(KanbanCard.board_id == field.board_id).all()
        has_usage = any((card.custom_fields or {}).get(field.key) is not None for card in used)
        if has_usage:
            field.is_active = False
            action = "custom_field.disabled"
        else:
            db.delete(field)
            action = "custom_field.deleted"
        record_activity(db, board.id, current_user, action, metadata={"field_id": field_id, "key": field.key})
        db.commit()
        return {"status": "success", "action": action}

    @staticmethod
    def _validate_view_payload(values: Dict[str, Any]) -> None:
        if "view_type" in values and values["view_type"] not in VIEW_TYPES:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Tipo de visualizacao invalido.")
        if "density" in values and values["density"] not in DENSITIES:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Densidade invalida.")

    @staticmethod
    def list_views(db: Session, board_id: int, current_user: User) -> List[KanbanBoardView]:
        board = KanbanService.get_board_detail(db, board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.READ_ONLY)
        KanbanService.ensure_default_views(db, board, current_user)
        db.commit()
        return db.query(KanbanBoardView).filter(KanbanBoardView.board_id == board_id).order_by(KanbanBoardView.is_default.desc(), KanbanBoardView.name).all()

    @staticmethod
    def create_view(db: Session, board_id: int, payload: BoardViewCreate, current_user: User) -> KanbanBoardView:
        board = KanbanService.get_board_detail(db, board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.MANAGER)
        values = payload.model_dump()
        KanbanService._validate_view_payload(values)
        if values.get("is_default"):
            db.query(KanbanBoardView).filter(KanbanBoardView.board_id == board_id).update({"is_default": False})
        view = KanbanBoardView(board_id=board_id, created_by_user_id=current_user.id, **values)
        db.add(view)
        db.flush()
        record_activity(db, board_id, current_user, "view.created", metadata={"view_id": view.id, "name": view.name, "view_type": view.view_type})
        db.commit()
        db.refresh(view)
        return view

    @staticmethod
    def _get_view(db: Session, view_id: int, current_user: User, required: BoardAccessLevel = BoardAccessLevel.READ_ONLY) -> KanbanBoardView:
        view = db.query(KanbanBoardView).filter(KanbanBoardView.id == view_id).first()
        if not view:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Visualizacao nao encontrada.")
        board = KanbanService.get_board_detail(db, view.board_id, current_user)
        require_board_level(db, board, current_user, required)
        return view

    @staticmethod
    def update_view(db: Session, view_id: int, payload: BoardViewUpdate, current_user: User) -> KanbanBoardView:
        view = KanbanService._get_view(db, view_id, current_user, BoardAccessLevel.MANAGER)
        values = payload.model_dump(exclude_unset=True)
        KanbanService._validate_view_payload(values)
        if values.get("is_default"):
            db.query(KanbanBoardView).filter(KanbanBoardView.board_id == view.board_id).update({"is_default": False})
        for field, value in values.items():
            setattr(view, field, value)
        view.updated_at = datetime.now(timezone.utc)
        record_activity(db, view.board_id, current_user, "view.updated", metadata={"view_id": view.id, "name": view.name})
        db.commit()
        db.refresh(view)
        return view

    @staticmethod
    def set_default_view(db: Session, view_id: int, current_user: User) -> KanbanBoardView:
        view = KanbanService._get_view(db, view_id, current_user, BoardAccessLevel.MANAGER)
        db.query(KanbanBoardView).filter(KanbanBoardView.board_id == view.board_id).update({"is_default": False})
        view.is_default = True
        view.updated_at = datetime.now(timezone.utc)
        record_activity(db, view.board_id, current_user, "view.default_changed", metadata={"view_id": view.id, "name": view.name})
        db.commit()
        db.refresh(view)
        return view

    @staticmethod
    def duplicate_view(db: Session, view_id: int, current_user: User) -> KanbanBoardView:
        view = KanbanService._get_view(db, view_id, current_user, BoardAccessLevel.MANAGER)
        cloned = KanbanBoardView(
            board_id=view.board_id,
            name=f"Copia - {view.name}",
            view_type=view.view_type,
            is_default=False,
            density=view.density,
            visible_columns=view.visible_columns,
            column_order=view.column_order,
            column_widths=view.column_widths,
            filters=view.filters,
            sort_by=view.sort_by,
            group_by=view.group_by,
            color_rules=view.color_rules,
            font_scale=view.font_scale,
            auto_scroll=view.auto_scroll,
            auto_scroll_seconds=view.auto_scroll_seconds,
            created_by_user_id=current_user.id,
        )
        db.add(cloned)
        db.flush()
        record_activity(db, view.board_id, current_user, "view.created", metadata={"view_id": cloned.id, "duplicated_from_view_id": view.id})
        db.commit()
        db.refresh(cloned)
        return cloned

    @staticmethod
    def delete_view(db: Session, view_id: int, current_user: User) -> Dict[str, str]:
        view = KanbanService._get_view(db, view_id, current_user, BoardAccessLevel.MANAGER)
        count = db.query(func.count(KanbanBoardView.id)).filter(KanbanBoardView.board_id == view.board_id).scalar() or 0
        if count <= 1:
            raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Nao e possivel remover a unica visualizacao do quadro.")
        if view.is_default:
            raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Defina outra visualizacao padrao antes de remover esta.")
        metadata = {"view_id": view.id, "name": view.name}
        db.delete(view)
        record_activity(db, view.board_id, current_user, "view.deleted", metadata=metadata)
        db.commit()
        return {"status": "success"}

    @staticmethod
    def list_board_activity(db: Session, board_id: int, current_user: User) -> List[KanbanActivity]:
        board = KanbanService.get_board_detail(db, board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.READ_ONLY)
        return db.query(KanbanActivity).filter(KanbanActivity.board_id == board_id).order_by(KanbanActivity.created_at.desc()).limit(100).all()

    @staticmethod
    def list_card_activity(db: Session, card_id: int, current_user: User) -> List[KanbanActivity]:
        card = KanbanService.get_card_detail(db, card_id, current_user)
        return db.query(KanbanActivity).filter(KanbanActivity.card_id == card.id).order_by(KanbanActivity.created_at.desc()).limit(100).all()

    @staticmethod
    def _get_card_for_write(db: Session, card_id: int, current_user: User) -> KanbanCard:
        card = KanbanService.get_card_detail(db, card_id, current_user)
        board = KanbanService.get_board_detail(db, card.board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.NORMAL)
        return card

    @staticmethod
    def _get_checklist(db: Session, checklist_id: int, current_user: User, required: BoardAccessLevel = BoardAccessLevel.READ_ONLY) -> KanbanCardChecklist:
        checklist = db.query(KanbanCardChecklist).filter(KanbanCardChecklist.id == checklist_id).first()
        if not checklist:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Checklist nao encontrado.")
        card = KanbanService.get_card_detail(db, checklist.card_id, current_user)
        board = KanbanService.get_board_detail(db, card.board_id, current_user)
        require_board_level(db, board, current_user, required)
        return checklist

    @staticmethod
    def _get_checklist_item(db: Session, item_id: int, current_user: User, required: BoardAccessLevel = BoardAccessLevel.READ_ONLY) -> KanbanCardChecklistItem:
        item = db.query(KanbanCardChecklistItem).filter(KanbanCardChecklistItem.id == item_id).first()
        if not item:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Item de checklist nao encontrado.")
        KanbanService._get_checklist(db, item.checklist_id, current_user, required)
        return item

    @staticmethod
    def list_checklists(db: Session, card_id: int, current_user: User) -> List[KanbanCardChecklist]:
        KanbanService.get_card_detail(db, card_id, current_user)
        return db.query(KanbanCardChecklist).filter(KanbanCardChecklist.card_id == card_id).order_by(KanbanCardChecklist.position).all()

    @staticmethod
    def create_checklist(db: Session, card_id: int, payload: ChecklistCreate, current_user: User) -> KanbanCardChecklist:
        card = KanbanService._get_card_for_write(db, card_id, current_user)
        position = db.query(func.count(KanbanCardChecklist.id)).filter(KanbanCardChecklist.card_id == card.id).scalar() or 0
        checklist = KanbanCardChecklist(card_id=card.id, title=payload.title, position=position, created_by_user_id=current_user.id)
        db.add(checklist)
        db.flush()
        record_activity(db, card.board_id, current_user, "checklist.created", card_id=card.id, metadata={"checklist_id": checklist.id})
        db.commit()
        db.refresh(checklist)
        return checklist

    @staticmethod
    def update_checklist(db: Session, checklist_id: int, payload: ChecklistUpdate, current_user: User) -> KanbanCardChecklist:
        checklist = KanbanService._get_checklist(db, checklist_id, current_user, BoardAccessLevel.NORMAL)
        values = payload.model_dump(exclude_unset=True)
        for field, value in values.items():
            setattr(checklist, field, value)
        checklist.updated_at = datetime.now(timezone.utc)
        record_activity(db, checklist.card.board_id, current_user, "checklist.updated", card_id=checklist.card_id, metadata={"checklist_id": checklist.id})
        db.commit()
        db.refresh(checklist)
        return checklist

    @staticmethod
    def delete_checklist(db: Session, checklist_id: int, current_user: User) -> Dict[str, str]:
        checklist = KanbanService._get_checklist(db, checklist_id, current_user, BoardAccessLevel.NORMAL)
        card_id = checklist.card_id
        board_id = checklist.card.board_id
        db.delete(checklist)
        record_activity(db, board_id, current_user, "checklist.deleted", card_id=card_id, metadata={"checklist_id": checklist_id})
        db.commit()
        return {"status": "success"}

    @staticmethod
    def create_checklist_item(db: Session, checklist_id: int, payload: ChecklistItemCreate, current_user: User) -> KanbanCardChecklistItem:
        checklist = KanbanService._get_checklist(db, checklist_id, current_user, BoardAccessLevel.NORMAL)
        position = db.query(func.count(KanbanCardChecklistItem.id)).filter(KanbanCardChecklistItem.checklist_id == checklist.id).scalar() or 0
        item = KanbanCardChecklistItem(checklist_id=checklist.id, text=payload.text, position=position, created_by_user_id=current_user.id)
        db.add(item)
        db.flush()
        record_activity(db, checklist.card.board_id, current_user, "checklist_item.created", card_id=checklist.card_id, metadata={"item_id": item.id})
        db.commit()
        db.refresh(item)
        return item

    @staticmethod
    def update_checklist_item(db: Session, item_id: int, payload: ChecklistItemUpdate, current_user: User) -> KanbanCardChecklistItem:
        item = KanbanService._get_checklist_item(db, item_id, current_user, BoardAccessLevel.NORMAL)
        old_done = item.is_done
        values = payload.model_dump(exclude_unset=True)
        if "is_done" in values:
            item.is_done = values["is_done"]
            item.completed_by_user_id = current_user.id if item.is_done else None
            item.completed_at = datetime.now(timezone.utc) if item.is_done else None
        if "text" in values:
            item.text = values["text"]
        if "position" in values:
            item.position = values["position"]
        item.updated_at = datetime.now(timezone.utc)
        action = "checklist_item.updated"
        if "is_done" in values and values["is_done"] != old_done:
            action = "checklist_item.checked" if values["is_done"] else "checklist_item.unchecked"
        record_activity(db, item.checklist.card.board_id, current_user, action, card_id=item.checklist.card_id, metadata={"item_id": item.id})
        db.commit()
        db.refresh(item)
        return item

    @staticmethod
    def delete_checklist_item(db: Session, item_id: int, current_user: User) -> Dict[str, str]:
        item = KanbanService._get_checklist_item(db, item_id, current_user, BoardAccessLevel.NORMAL)
        card_id = item.checklist.card_id
        board_id = item.checklist.card.board_id
        db.delete(item)
        record_activity(db, board_id, current_user, "checklist_item.deleted", card_id=card_id, metadata={"item_id": item_id})
        db.commit()
        return {"status": "success"}

    @staticmethod
    def reorder_checklist_items(db: Session, checklist_id: int, payload: ChecklistItemReorderPayload, current_user: User) -> List[KanbanCardChecklistItem]:
        checklist = KanbanService._get_checklist(db, checklist_id, current_user, BoardAccessLevel.NORMAL)
        ids = [item.item_id for item in payload.items]
        items = db.query(KanbanCardChecklistItem).filter(KanbanCardChecklistItem.checklist_id == checklist.id, KanbanCardChecklistItem.id.in_(ids)).all()
        if len(items) != len(set(ids)):
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Reordenacao contem item invalido.")
        positions = {item.item_id: index for index, item in enumerate(sorted(payload.items, key=lambda value: value.position))}
        for item in items:
            item.position = positions[item.id]
        record_activity(db, checklist.card.board_id, current_user, "checklist_item.reordered", card_id=checklist.card_id, metadata={"checklist_id": checklist.id})
        db.commit()
        return db.query(KanbanCardChecklistItem).filter(KanbanCardChecklistItem.checklist_id == checklist.id).order_by(KanbanCardChecklistItem.position).all()

    @staticmethod
    def list_comments(db: Session, card_id: int, current_user: User) -> List[KanbanCardComment]:
        KanbanService.get_card_detail(db, card_id, current_user)
        return db.query(KanbanCardComment).filter(KanbanCardComment.card_id == card_id, KanbanCardComment.deleted_at == None).order_by(KanbanCardComment.created_at).all()

    @staticmethod
    def create_comment(db: Session, card_id: int, payload: CommentCreate, current_user: User) -> KanbanCardComment:
        card = KanbanService._get_card_for_write(db, card_id, current_user)
        comment = KanbanCardComment(card_id=card.id, user_id=current_user.id, comment=payload.comment.strip())
        db.add(comment)
        db.flush()
        record_activity(db, card.board_id, current_user, "comment.created", card_id=card.id, metadata={"comment_id": comment.id})
        
        from app.core.events import emit_event
        assigned_user_ids = [a.user_id for a in card.assignees] if getattr(card, "assignees", None) else []
        if not assigned_user_ids and card.assigned_to_user_id:
            assigned_user_ids = [card.assigned_to_user_id]
            
        emit_event(
            db=db,
            event_type="kanban.card.comment.created",
            aggregate_type="kanban_card",
            aggregate_id=str(card.id),
            module="kanban",
            payload={
                "card_id": card.id,
                "board_id": card.board_id,
                "title": card.title,
                "comment_id": comment.id,
                "comment_text": comment.comment,
                "author_user_id": current_user.id,
                "assigned_to_user_ids": assigned_user_ids
            },
            actor_user_id=current_user.id
        )
        
        db.commit()
        db.refresh(comment)
        return comment

    @staticmethod
    def _get_comment_for_action(db: Session, comment_id: int, current_user: User) -> KanbanCardComment:
        comment = db.query(KanbanCardComment).filter(KanbanCardComment.id == comment_id, KanbanCardComment.deleted_at == None).first()
        if not comment:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Comentario nao encontrado.")
        card = KanbanService.get_card_detail(db, comment.card_id, current_user)
        board = KanbanService.get_board_detail(db, card.board_id, current_user)
        level = require_board_level(db, board, current_user, BoardAccessLevel.READ_ONLY)
        is_global_admin = bool(current_user.role and current_user.role.name == "ADMIN")
        if not is_global_admin and comment.user_id != current_user.id and BOARD_LEVEL_VALUES[level] < BOARD_LEVEL_VALUES[BoardAccessLevel.MANAGER]:
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Somente o autor ou gestor do board pode alterar este comentario.")
        return comment

    @staticmethod
    def update_comment(db: Session, comment_id: int, payload: CommentUpdate, current_user: User) -> KanbanCardComment:
        comment = KanbanService._get_comment_for_action(db, comment_id, current_user)
        comment.comment = payload.comment.strip()
        comment.edited_at = datetime.now(timezone.utc)
        comment.updated_at = datetime.now(timezone.utc)
        record_activity(db, comment.card.board_id, current_user, "comment.updated", card_id=comment.card_id, metadata={"comment_id": comment.id})
        db.commit()
        db.refresh(comment)
        return comment

    @staticmethod
    def delete_comment(db: Session, comment_id: int, current_user: User) -> Dict[str, str]:
        comment = KanbanService._get_comment_for_action(db, comment_id, current_user)
        comment.deleted_at = datetime.now(timezone.utc)
        record_activity(db, comment.card.board_id, current_user, "comment.deleted", card_id=comment.card_id, metadata={"comment_id": comment.id})
        db.commit()
        return {"status": "success"}

    @staticmethod
    def list_labels(db: Session, board_id: int, current_user: User, include_inactive: bool = False) -> List[KanbanLabel]:
        board = KanbanService.get_board_detail(db, board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.READ_ONLY)
        query = db.query(KanbanLabel).filter(KanbanLabel.board_id == board_id)
        if not include_inactive:
            query = query.filter(KanbanLabel.is_active == True)
        return query.order_by(KanbanLabel.is_active.desc(), KanbanLabel.name).all()

    @staticmethod
    def create_label(db: Session, board_id: int, payload: LabelCreate, current_user: User) -> KanbanLabel:
        board = KanbanService.get_board_detail(db, board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.MANAGER)
        label = KanbanLabel(board_id=board_id, name=payload.name, color=payload.color)
        db.add(label)
        db.flush()
        record_activity(db, board_id, current_user, "label.created", metadata={"label_id": label.id, "name": label.name})
        db.commit()
        db.refresh(label)
        return label

    @staticmethod
    def update_label(db: Session, label_id: int, payload: LabelUpdate, current_user: User) -> KanbanLabel:
        label = db.query(KanbanLabel).filter(KanbanLabel.id == label_id).first()
        if not label:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Etiqueta nao encontrada.")
        board = KanbanService.get_board_detail(db, label.board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.MANAGER)
        for field, value in payload.model_dump(exclude_unset=True).items():
            setattr(label, field, value)
        label.updated_at = datetime.now(timezone.utc)
        record_activity(db, label.board_id, current_user, "label.updated", metadata={"label_id": label.id})
        db.commit()
        db.refresh(label)
        return label

    @staticmethod
    def delete_label(db: Session, label_id: int, current_user: User) -> Dict[str, str]:
        label = db.query(KanbanLabel).filter(KanbanLabel.id == label_id).first()
        if not label:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Etiqueta nao encontrada.")
        board = KanbanService.get_board_detail(db, label.board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.MANAGER)
        in_use = db.query(KanbanCardLabel).filter(KanbanCardLabel.label_id == label.id).first()
        if in_use:
            label.is_active = False
            action = "label.disabled"
        else:
            db.delete(label)
            action = "label.deleted"
        record_activity(db, board.id, current_user, action, metadata={"label_id": label_id})
        db.commit()
        return {"status": "success", "action": action}

    @staticmethod
    def apply_label(db: Session, card_id: int, label_id: int, current_user: User) -> KanbanCard:
        card = KanbanService._get_card_for_write(db, card_id, current_user)
        label = db.query(KanbanLabel).filter(KanbanLabel.id == label_id, KanbanLabel.board_id == card.board_id, KanbanLabel.is_active == True).first()
        if not label:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Etiqueta nao encontrada para este board.")
        existing = db.query(KanbanCardLabel).filter(KanbanCardLabel.card_id == card.id, KanbanCardLabel.label_id == label.id).first()
        if not existing:
            db.add(KanbanCardLabel(card_id=card.id, label_id=label.id))
            record_activity(db, card.board_id, current_user, "label.applied", card_id=card.id, metadata={"label_id": label.id})
            db.commit()
        return KanbanService.get_card_detail(db, card.id, current_user)

    @staticmethod
    def remove_label(db: Session, card_id: int, label_id: int, current_user: User) -> KanbanCard:
        card = KanbanService._get_card_for_write(db, card_id, current_user)
        link = db.query(KanbanCardLabel).filter(KanbanCardLabel.card_id == card.id, KanbanCardLabel.label_id == label_id).first()
        if link:
            db.delete(link)
            record_activity(db, card.board_id, current_user, "label.removed", card_id=card.id, metadata={"label_id": label_id})
            db.commit()
        return KanbanService.get_card_detail(db, card.id, current_user)

    @staticmethod
    def list_assignees(db: Session, card_id: int, current_user: User) -> List[KanbanCardAssignee]:
        KanbanService.get_card_detail(db, card_id, current_user)
        return db.query(KanbanCardAssignee).filter(KanbanCardAssignee.card_id == card_id).order_by(KanbanCardAssignee.created_at).all()

    @staticmethod
    def add_assignee(db: Session, card_id: int, user_id: int, current_user: User) -> KanbanCard:
        card = KanbanService._get_card_for_write(db, card_id, current_user)
        target_user = db.query(User).filter(User.id == user_id, User.is_active == True).first()
        if not target_user:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Usuario nao encontrado ou inativo.")
        board = KanbanService.get_board_detail(db, card.board_id, current_user)
        if BOARD_LEVEL_VALUES[get_board_access_level(db, board, target_user)] < BOARD_LEVEL_VALUES[BoardAccessLevel.READ_ONLY]:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Usuario precisa ter acesso ao board para ser responsavel.")
        existing = db.query(KanbanCardAssignee).filter(KanbanCardAssignee.card_id == card.id, KanbanCardAssignee.user_id == user_id).first()
        if not existing:
            db.add(KanbanCardAssignee(card_id=card.id, user_id=user_id, assigned_by_user_id=current_user.id))
            if card.assigned_to_user_id is None:
                card.assigned_to_user_id = user_id
            record_activity(db, card.board_id, current_user, "assignee.added", card_id=card.id, metadata={"user_id": user_id})
            
            from app.core.events import emit_event
            emit_event(
                db=db,
                event_type="kanban.card.assigned",
                aggregate_type="kanban_card",
                aggregate_id=str(card.id),
                module="kanban",
                payload={
                    "card_id": card.id,
                    "board_id": card.board_id,
                    "title": card.title,
                    "assigned_user_id": user_id,
                    "assigned_by_user_id": current_user.id
                },
                actor_user_id=current_user.id
            )
            
            db.commit()
        return KanbanService.get_card_detail(db, card.id, current_user)

    @staticmethod
    def remove_assignee(db: Session, card_id: int, user_id: int, current_user: User) -> KanbanCard:
        card = KanbanService._get_card_for_write(db, card_id, current_user)
        link = db.query(KanbanCardAssignee).filter(KanbanCardAssignee.card_id == card.id, KanbanCardAssignee.user_id == user_id).first()
        if link:
            db.delete(link)
            db.flush()
            first = db.query(KanbanCardAssignee).filter(KanbanCardAssignee.card_id == card.id).order_by(KanbanCardAssignee.created_at).first()
            card.assigned_to_user_id = first.user_id if first else None
            record_activity(db, card.board_id, current_user, "assignee.removed", card_id=card.id, metadata={"user_id": user_id})
            db.commit()
        return KanbanService.get_card_detail(db, card.id, current_user)

    @staticmethod
    def search_cards(
        db: Session,
        board_id: int,
        current_user: User,
        text: Optional[str] = None,
        priority: Optional[str] = None,
        column_id: Optional[int] = None,
        assignee_id: Optional[int] = None,
        label_id: Optional[int] = None,
        overdue: bool = False,
        archived: bool = False,
        created_by_me: bool = False,
        assigned_to_me: bool = False,
        unassigned: bool = False,
    ) -> List[KanbanCard]:
        board = KanbanService.get_board_detail(db, board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.READ_ONLY)
        query = db.query(KanbanCard).filter(KanbanCard.board_id == board_id, KanbanCard.is_archived == archived)
        if text:
            like = f"%{text}%"
            query = query.filter(or_(KanbanCard.title.ilike(like), KanbanCard.description.ilike(like)))
        if priority:
            query = query.filter(KanbanCard.priority == priority)
        if column_id:
            query = query.filter(KanbanCard.column_id == column_id)
        if overdue:
            now = datetime.now(timezone.utc).replace(tzinfo=None)
            query = query.join(KanbanColumn, KanbanColumn.id == KanbanCard.column_id).filter(KanbanColumn.is_done_column == False, KanbanCard.due_date != None, KanbanCard.due_date < now)
        if created_by_me:
            query = query.filter(KanbanCard.created_by_user_id == current_user.id)
        if assignee_id:
            query = query.join(KanbanCardAssignee, KanbanCardAssignee.card_id == KanbanCard.id).filter(KanbanCardAssignee.user_id == assignee_id)
        if assigned_to_me:
            query = query.join(KanbanCardAssignee, KanbanCardAssignee.card_id == KanbanCard.id).filter(KanbanCardAssignee.user_id == current_user.id)
        if label_id:
            query = query.join(KanbanCardLabel, KanbanCardLabel.card_id == KanbanCard.id).filter(KanbanCardLabel.label_id == label_id)
        if unassigned:
            query = query.outerjoin(KanbanCardAssignee, KanbanCardAssignee.card_id == KanbanCard.id).filter(KanbanCardAssignee.id == None)
        cards = query.order_by(KanbanCard.updated_at.desc()).limit(200).all()
        return [KanbanService.decorate_card(card) for card in cards]

    @staticmethod
    def list_attachments(db: Session, card_id: int, current_user: User) -> List[KanbanCardAttachment]:
        KanbanService.get_card_detail(db, card_id, current_user)
        return db.query(KanbanCardAttachment).filter(KanbanCardAttachment.card_id == card_id, KanbanCardAttachment.deleted_at == None).order_by(KanbanCardAttachment.created_at.desc()).all()

    @staticmethod
    def _validate_upload(upload: UploadFile, data: bytes) -> str:
        filename = upload.filename or "arquivo"
        extension = Path(filename).suffix.lower()
        if extension in BLOCKED_EXTENSIONS or extension not in ALLOWED_EXTENSIONS:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Tipo de arquivo nao permitido.")
        if len(data) > MAX_UPLOAD_BYTES:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Arquivo acima do limite de 10 MB.")
        return extension

    @staticmethod
    def upload_attachment(db: Session, card_id: int, upload: UploadFile, current_user: User) -> KanbanCardAttachment:
        card = KanbanService._get_card_for_write(db, card_id, current_user)
        data = upload.file.read()
        extension = KanbanService._validate_upload(upload, data)
        UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
        stored_filename = f"{uuid.uuid4().hex}{extension}"
        storage_key = stored_filename
        path = UPLOAD_DIR / stored_filename
        with path.open("wb") as buffer:
            buffer.write(data)
        checksum = hashlib.sha256(data).hexdigest()
        file_row = File(
            original_filename=os.path.basename(upload.filename or "arquivo"),
            stored_filename=stored_filename,
            content_type=upload.content_type or "application/octet-stream",
            size_bytes=len(data),
            storage_provider="local",
            storage_bucket="backend/storage/uploads",
            storage_key=storage_key,
            checksum_sha256=checksum,
            uploaded_by_user_id=current_user.id,
        )
        db.add(file_row)
        db.flush()
        db.add(FileModuleLink(file_id=file_row.id, module_slug="kanban", entity_type="kanban_card", entity_id=card.id, created_by_user_id=current_user.id))
        attachment = KanbanCardAttachment(card_id=card.id, file_id=file_row.id, uploaded_by_user_id=current_user.id)
        db.add(attachment)
        db.flush()
        record_activity(db, card.board_id, current_user, "attachment.uploaded", card_id=card.id, metadata={"attachment_id": attachment.id, "filename": file_row.original_filename})
        db.commit()
        db.refresh(attachment)
        return attachment

    @staticmethod
    def _get_attachment(db: Session, attachment_id: int, current_user: User) -> KanbanCardAttachment:
        attachment = db.query(KanbanCardAttachment).filter(KanbanCardAttachment.id == attachment_id, KanbanCardAttachment.deleted_at == None).first()
        if not attachment:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Anexo nao encontrado.")
        KanbanService.get_card_detail(db, attachment.card_id, current_user)
        return attachment

    @staticmethod
    def download_attachment(db: Session, attachment_id: int, current_user: User) -> DownloadFileResponse:
        attachment = KanbanService._get_attachment(db, attachment_id, current_user)
        file_row = attachment.file
        path = UPLOAD_DIR / file_row.storage_key
        if not path.exists():
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Arquivo fisico nao encontrado.")
        return DownloadFileResponse(path=path, media_type=file_row.content_type, filename=file_row.original_filename)

    @staticmethod
    def delete_attachment(db: Session, attachment_id: int, current_user: User) -> Dict[str, str]:
        attachment = KanbanService._get_attachment(db, attachment_id, current_user)
        card = KanbanService.get_card_detail(db, attachment.card_id, current_user)
        board = KanbanService.get_board_detail(db, card.board_id, current_user)
        level = require_board_level(db, board, current_user, BoardAccessLevel.READ_ONLY)
        if attachment.uploaded_by_user_id != current_user.id and BOARD_LEVEL_VALUES[level] < BOARD_LEVEL_VALUES[BoardAccessLevel.MANAGER]:
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Somente quem anexou ou gestor do board pode remover este anexo.")
        attachment.deleted_at = datetime.now(timezone.utc)
        attachment.file.deleted_at = datetime.now(timezone.utc)
        record_activity(db, card.board_id, current_user, "attachment.deleted", card_id=card.id, metadata={"attachment_id": attachment.id})
        db.commit()
        return {"status": "success"}

    @staticmethod
    def default_tv_config(board_id: int) -> Dict[str, Any]:
        return {
            "id": None,
            "board_id": board_id,
            "name": "Padrao",
            "is_default": True,
            "layout_type": "COLUMNS",
            "refresh_interval_seconds": 30,
            "show_archived": False,
            "show_done_columns": True,
            "show_checklist_progress": True,
            "show_assignees": True,
            "show_labels": True,
            "show_due_date": True,
            "show_card_description": False,
            "group_by": "column",
            "sort_by": "position",
            "filters": {},
            "visible_custom_fields": PRODUCTION_CUSTOM_KEYS,
            "display_options": {
                "show_portal_name": False,
                "show_board_name": True,
                "show_layout_name": True,
                "show_clock": True,
                "show_date": False,
                "show_connection": True,
                "show_last_update": True,
                "show_exit_button": True,
                "show_refresh_button": True,
            },
            "kpi_options": {
                "show_kpis": True,
                "visible": ["total_active_cards", "critical_cards", "overdue_cards", "due_today_cards", "unassigned_cards"],
            },
            "layout_options": {
                "density": "comfortable",
                "font_scale": "large",
                "visible_rows": "auto",
                "short_labels": {},
                "auto_scroll": False,
                "auto_scroll_seconds": 18,
                "show_ranking": True,
                "ranking_limit": 8,
            },
            "external_mode_options": {
                "layout_type": "PRODUCTION_LIST",
                "display_options": {
                    "show_portal_name": False,
                    "show_board_name": False,
                    "show_layout_name": False,
                    "show_clock": False,
                    "show_date": False,
                    "show_connection": False,
                    "show_last_update": False,
                    "show_exit_button": False,
                    "show_refresh_button": False,
                },
                "kpi_options": {"show_kpis": False, "visible": []},
                "layout_options": {
                    "density": "factory",
                    "font_scale": "large",
                    "visible_rows": 16,
                    "show_ranking": False,
                    "ranking_limit": 0,
                    "auto_scroll": True,
                    "auto_scroll_seconds": 18,
                    "short_labels": {
                        "tensao": "V",
                        "quantidade": "Qtd",
                        "qtd": "Qtd",
                        "entrada": "Inicio",
                        "inicio": "Inicio",
                        "pendencia": "Pend.",
                        "entrega": "Entrega"
                    },
                },
            },
        }

    @staticmethod
    def get_tv_config(db: Session, board_id: int, current_user: User) -> Any:
        board = KanbanService.get_board_detail(db, board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.READ_ONLY)
        config = db.query(KanbanTVView).filter(KanbanTVView.board_id == board_id, KanbanTVView.is_default == True).first()
        return config or KanbanService.default_tv_config(board_id)

    @staticmethod
    def update_tv_config(db: Session, board_id: int, payload: TVConfigUpdate, current_user: User) -> KanbanTVView:
        board = KanbanService.get_board_detail(db, board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.MANAGER)
        config = db.query(KanbanTVView).filter(KanbanTVView.board_id == board_id, KanbanTVView.is_default == True).first()
        if not config:
            config = KanbanTVView(
                board_id=board_id,
                name="Padrao",
                is_default=True,
                created_by_user_id=current_user.id,
                visible_custom_fields=PRODUCTION_CUSTOM_KEYS,
                filters={},
                display_options=KanbanService.default_tv_config(board_id)["display_options"],
                kpi_options=KanbanService.default_tv_config(board_id)["kpi_options"],
                layout_options=KanbanService.default_tv_config(board_id)["layout_options"],
                external_mode_options=KanbanService.default_tv_config(board_id)["external_mode_options"],
            )
            db.add(config)
            db.flush()
        values = payload.model_dump(exclude_unset=True)
        if "layout_type" in values and values["layout_type"] not in TV_LAYOUTS:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Layout de TV invalido.")
        for field, value in values.items():
            setattr(config, field, value)
        config.updated_at = datetime.now(timezone.utc)
        record_activity(db, board_id, current_user, "tv_config.updated", metadata={"layout_type": config.layout_type})
        db.commit()
        db.refresh(config)
        return config

    @staticmethod
    def card_urgency(card: KanbanCard, now: Optional[datetime] = None) -> Dict[str, Any]:
        now = now or datetime.now(timezone.utc).replace(tzinfo=None)
        score = 0
        reasons: List[str] = []
        due = card.due_date
        if card.priority == "URGENT":
            score += 50
            reasons.append("Prioridade urgente")
        elif card.priority == "HIGH":
            score += 30
            reasons.append("Prioridade alta")
        if due:
            days = (due.date() - now.date()).days
            if days < 0:
                score += 40
                reasons.append(f"Vencido ha {abs(days)} dia(s)")
            elif days == 0:
                score += 25
                reasons.append("Vence hoje")
            elif days <= 3:
                score += 15
                reasons.append(f"Vence em {days} dia(s)")
        if not getattr(card, "assignees", []):
            score += 10
            reasons.append("Sem responsavel")
        stopped_days = max((now.date() - card.updated_at.date()).days, 0)
        if stopped_days > 3:
            score += 15
            reasons.append(f"Parado ha {stopped_days} dia(s)")
        if card.due_date and card.checklist_total and card.checklist_done < card.checklist_total and (card.due_date.date() - now.date()).days <= 3:
            score += 10
            reasons.append("Checklist incompleto proximo ao prazo")
        if any((label.name or "").lower() == "urgente" for label in getattr(card, "labels", [])):
            score += 20
            reasons.append("Etiqueta Urgente")
        return {
            "urgency_score": score,
            "urgency_reasons": reasons or ["Sem alerta"],
            "days_overdue": abs((due.date() - now.date()).days) if due and due.date() < now.date() else 0,
            "days_stopped": stopped_days,
        }

    @staticmethod
    def _cards_for_board(db: Session, board_id: int, include_archived: bool = False) -> List[KanbanCard]:
        query = db.query(KanbanCard).filter(KanbanCard.board_id == board_id)
        if not include_archived:
            query = query.filter(KanbanCard.is_archived == False)
        cards = query.order_by(KanbanCard.position, KanbanCard.id).all()
        return [KanbanService.decorate_card(card) for card in cards]

    @staticmethod
    def get_board_metrics(db: Session, board_id: int, current_user: User) -> Dict[str, Any]:
        board = KanbanService.get_board_detail(db, board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.READ_ONLY)
        now = datetime.now(timezone.utc).replace(tzinfo=None)
        cards = KanbanService._cards_for_board(db, board_id, include_archived=True)
        active = [card for card in cards if not card.is_archived]
        archived = [card for card in cards if card.is_archived]
        columns = {column.id: column for column in board.columns}
        by_column = {column.name: len([card for card in active if card.column_id == column.id]) for column in board.columns}
        by_priority = {priority: len([card for card in active if card.priority == priority]) for priority in ["LOW", "MEDIUM", "HIGH", "URGENT"]}
        overdue = [card for card in active if card.due_date and card.due_date.date() < now.date()]
        due_today = [card for card in active if card.due_date and card.due_date.date() == now.date()]
        due_soon = [card for card in active if card.due_date and 0 < (card.due_date.date() - now.date()).days <= 3]
        unassigned = [card for card in active if not getattr(card, "assignees", [])]
        no_due = [card for card in active if not card.due_date]
        stopped = [card for card in active if (now.date() - card.updated_at.date()).days > 3]
        done_columns = {column.id for column in board.columns if column.is_done_column}
        completed_7d = [card for card in active if card.column_id in done_columns and card.updated_at >= now - timedelta(days=7)]
        created_7d = [card for card in active if card.created_at >= now - timedelta(days=7)]
        checklist_cards = [card for card in active if card.checklist_total]
        checklist_avg = 0
        if checklist_cards:
            checklist_avg = round(sum(card.checklist_done / card.checklist_total for card in checklist_cards) / len(checklist_cards) * 100, 2)
        by_assignee: Dict[str, int] = {}
        for card in active:
            for assignee in getattr(card, "assignees", []):
                name = assignee.user.username if assignee.user else str(assignee.user_id)
                by_assignee[name] = by_assignee.get(name, 0) + 1
        by_label: Dict[str, int] = {}
        for card in active:
            for label in getattr(card, "labels", []):
                by_label[label.name] = by_label.get(label.name, 0) + 1
        critical = KanbanService.get_critical_cards(db, board_id, current_user, limit=100)
        return {
            "total_active_cards": len(active),
            "total_archived_cards": len(archived),
            "cards_by_column": by_column,
            "cards_by_priority": by_priority,
            "overdue_cards": len(overdue),
            "due_today_cards": len(due_today),
            "due_soon_cards": len(due_soon),
            "unassigned_cards": len(unassigned),
            "no_due_date_cards": len(no_due),
            "stopped_cards": len(stopped),
            "completed_last_7_days": len(completed_7d),
            "created_last_7_days": len(created_7d),
            "average_checklist_completion": checklist_avg,
            "cards_by_assignee": by_assignee,
            "cards_by_label": by_label,
            "critical_cards": len(critical),
            "board_urgency_score": sum(item["urgency_score"] for item in critical),
            "generated_at": now.isoformat(),
        }

    @staticmethod
    def get_critical_cards(db: Session, board_id: int, current_user: User, limit: int = 20) -> List[Dict[str, Any]]:
        board = KanbanService.get_board_detail(db, board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.READ_ONLY)
        columns = {column.id: column.name for column in board.columns}
        result = []
        for card in KanbanService._cards_for_board(db, board_id, include_archived=False):
            urgency = KanbanService.card_urgency(card)
            if urgency["urgency_score"] >= 25:
                result.append({
                    "id": card.id,
                    "title": card.title,
                    "priority": card.priority,
                    "due_date": card.due_date,
                    "column_id": card.column_id,
                    "column_name": columns.get(card.column_id),
                    "assignees": [{"id": item.user_id, "username": item.user.username if item.user else str(item.user_id)} for item in getattr(card, "assignees", [])],
                    **urgency,
                })
        return sorted(result, key=lambda item: item["urgency_score"], reverse=True)[:limit]

    @staticmethod
    def get_tv_data(db: Session, board_id: int, current_user: User) -> Dict[str, Any]:
        board = KanbanService.get_board_detail(db, board_id, current_user)
        config = KanbanService.get_tv_config(db, board_id, current_user)
        config_dict = config if isinstance(config, dict) else {
            key: getattr(config, key) for key in KanbanService.default_tv_config(board_id).keys() if hasattr(config, key)
        }
        defaults = KanbanService.default_tv_config(board_id)
        for key in ("display_options", "kpi_options", "layout_options", "external_mode_options"):
            config_dict[key] = {**defaults.get(key, {}), **(config_dict.get(key) or {})}
        show_archived = bool(config_dict.get("show_archived"))
        show_done = bool(config_dict.get("show_done_columns", True))
        sort_by = config_dict.get("sort_by") or "position"
        active_custom_fields = (
            db.query(KanbanCustomField)
            .filter(KanbanCustomField.board_id == board_id, KanbanCustomField.is_active == True)
            .order_by(KanbanCustomField.position)
            .all()
        )
        active_field_keys = {field.key for field in active_custom_fields}

        def serialize_tv_card(card: KanbanCard) -> Dict[str, Any]:
            urgency = KanbanService.card_urgency(card)
            return {
                "id": card.id,
                "title": card.title,
                "description": card.description,
                "column_id": card.column_id,
                "priority": card.priority,
                "status": card.status,
                "due_date": card.due_date,
                "created_at": card.created_at,
                "updated_at": card.updated_at,
                "is_archived": card.is_archived,
                "custom_fields": {key: value for key, value in (card.custom_fields or {}).items() if key in active_field_keys},
                "labels": [{"id": label.id, "name": label.name, "color": label.color} for label in getattr(card, "labels", [])],
                "assignees": [{"id": item.id, "user_id": item.user_id, "user": {"id": item.user_id, "username": item.user.username if item.user else str(item.user_id)}} for item in getattr(card, "assignees", [])],
                "checklist_total": getattr(card, "checklist_total", 0),
                "checklist_done": getattr(card, "checklist_done", 0),
                **urgency,
            }

        columns_payload = []
        for column in board.columns:
            if column.is_archived or (column.is_done_column and not show_done):
                continue
            cards = [card for card in column.cards if show_archived or not card.is_archived]
            for card in cards:
                urgency = KanbanService.card_urgency(card)
                card.urgency_score = urgency["urgency_score"]
                card.urgency_reasons = urgency["urgency_reasons"]
            if sort_by == "due_date":
                cards = sorted(cards, key=lambda card: card.due_date or datetime.max)
            elif sort_by == "priority":
                order = {"URGENT": 0, "HIGH": 1, "MEDIUM": 2, "LOW": 3}
                cards = sorted(cards, key=lambda card: order.get(card.priority, 9))
            elif sort_by == "updated_at":
                cards = sorted(cards, key=lambda card: card.updated_at, reverse=True)
            elif sort_by == "urgency_score":
                cards = sorted(cards, key=lambda card: getattr(card, "urgency_score", 0), reverse=True)
            columns_payload.append({
                "id": column.id,
                "board_id": column.board_id,
                "name": column.name,
                "position": column.position,
                "color": column.color,
                "wip_limit": column.wip_limit,
                "is_done_column": column.is_done_column,
                "is_archived": column.is_archived,
                "created_at": column.created_at,
                "updated_at": column.updated_at,
                "cards": [serialize_tv_card(card) for card in cards],
            })
        cards_payload = [card for column in columns_payload for card in column["cards"]]
        return {
            "board": {
                "id": board.id,
                "name": board.name,
                "slug": board.slug,
                "description": board.description,
                "color": board.color,
                "icon": board.icon,
                "access_level": getattr(board, "access_level", None),
                "custom_fields": active_custom_fields,
            },
            "columns": columns_payload,
            "cards": cards_payload,
            "metrics": KanbanService.get_board_metrics(db, board_id, current_user),
            "critical_cards": KanbanService.get_critical_cards(db, board_id, current_user),
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "connection_mode": "online",
            "config": config_dict,
            "production_fields": [key for key in PRODUCTION_CUSTOM_KEYS if key in active_field_keys],
        }

    @staticmethod
    def get_list_data(db: Session, board_id: int, current_user: User, view_id: Optional[int] = None) -> Dict[str, Any]:
        board = KanbanService.get_board_detail(db, board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.READ_ONLY)
        views = KanbanService.list_views(db, board_id, current_user)
        selected = None
        if view_id:
            selected = next((view for view in views if view.id == view_id), None)
            if not selected:
                raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Visualizacao nao encontrada neste quadro.")
        if not selected:
            selected = next((view for view in views if view.is_default and view.view_type in {"LIST", "PRODUCTION_LIST"}), None)
            selected = selected or next((view for view in views if view.view_type in {"LIST", "PRODUCTION_LIST"}), None)
        default_columns = ["title", "status", "priority", "due_date", "assignees", "labels", "checklist", "updated_at"]
        visible_columns = selected.visible_columns if selected and selected.visible_columns else default_columns
        system_columns = {"title", "status", "priority", "due_date", "assignees", "labels", "checklist", "updated_at", "description"}
        active_custom_fields = (
            db.query(KanbanCustomField)
            .filter(KanbanCustomField.board_id == board_id, KanbanCustomField.is_active == True)
            .order_by(KanbanCustomField.position)
            .all()
        )
        active_field_keys = {field.key for field in active_custom_fields}
        visible_columns = [column for column in visible_columns if column in system_columns or column in active_field_keys]
        cards = []
        column_names = {column.id: column.name for column in board.columns}
        for card in KanbanService._cards_for_board(db, board_id, include_archived=bool((selected.filters or {}).get("archived")) if selected and selected.filters else False):
            urgency = KanbanService.card_urgency(card)
            custom_values = {key: value for key, value in (card.custom_fields or {}).items() if key in active_field_keys}
            cards.append({
                "id": card.id,
                "title": card.title,
                "description": card.description,
                "column_id": card.column_id,
                "status": column_names.get(card.column_id),
                "priority": card.priority,
                "due_date": card.due_date,
                "updated_at": card.updated_at,
                "is_archived": card.is_archived,
                "custom_fields": custom_values,
                "labels": [{"id": label.id, "name": label.name, "color": label.color} for label in getattr(card, "labels", [])],
                "assignees": [{"id": item.user_id, "username": item.user.username if item.user else str(item.user_id)} for item in getattr(card, "assignees", [])],
                "checklist_done": getattr(card, "checklist_done", 0),
                "checklist_total": getattr(card, "checklist_total", 0),
                **urgency,
            })
        return {
            "board": board,
            "cards": cards,
            "columns": board.columns,
            "labels": board.labels,
            "custom_fields": active_custom_fields,
            "views": views,
            "list_config": selected,
            "visible_columns": visible_columns,
            "column_order": selected.column_order if selected else visible_columns,
            "column_widths": selected.column_widths if selected else {},
            "generated_at": datetime.now(timezone.utc).isoformat(),
        }

    @staticmethod
    def list_board_permissions(db: Session, board_id: int, current_user: User) -> List[KanbanBoardPermission]:
        KanbanService._require_global_admin(current_user)
        board = KanbanService.get_board_detail(db, board_id, current_user)
        permissions = db.query(KanbanBoardPermission).filter(KanbanBoardPermission.board_id == board_id).order_by(KanbanBoardPermission.id).all()
        for perm in permissions:
            username = None
            user_email = None
            role_name = None
            if perm.user_id:
                u = db.query(User).filter(User.id == perm.user_id).first()
                if u:
                    username = u.username
                    user_email = u.email
            if perm.role_id:
                from app.models.role import Role
                r = db.query(Role).filter(Role.id == perm.role_id).first()
                if r:
                    role_name = r.name
            setattr(perm, 'username', username)
            setattr(perm, 'user_email', user_email)
            setattr(perm, 'role_name', role_name)
        return permissions

    @staticmethod
    def upsert_board_permission(db: Session, board_id: int, payload: BoardPermissionUpsert, current_user: User) -> KanbanBoardPermission:
        KanbanService._require_global_admin(current_user)
        board = KanbanService.get_board_detail(db, board_id, current_user)
        if payload.access_level not in {level.value for level in BoardAccessLevel}:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Nivel de acesso invalido.")
        if bool(payload.user_id) == bool(payload.role_id):
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Informe um usuario ou um perfil, mas nao ambos.")
        if payload.user_id and not db.query(User).filter(User.id == payload.user_id, User.is_active == True).first():
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Usuario nao encontrado ou inativo.")
        from app.core.audit import log_action
        query = db.query(KanbanBoardPermission).filter(KanbanBoardPermission.board_id == board_id)
        query = query.filter(KanbanBoardPermission.user_id == payload.user_id) if payload.user_id else query.filter(KanbanBoardPermission.role_id == payload.role_id)
        permission = query.first()
        if permission:
            permission.access_level = payload.access_level
            permission.updated_at = datetime.now(timezone.utc)
            action = "kanban.board_access.updated"
        else:
            permission = KanbanBoardPermission(board_id=board_id, user_id=payload.user_id, role_id=payload.role_id, access_level=payload.access_level)
            db.add(permission)
            db.flush()
            action = "kanban.board_access.granted"
        log_action(
            db=db,
            user_id=current_user.id,
            action=action,
            module="kanban",
            details={
                "permission_id": permission.id,
                "board_id": board_id,
                "access_level": payload.access_level,
                "user_id": payload.user_id,
                "role_id": payload.role_id
            },
            commit=False
        )
        record_activity(db, board_id, current_user, "access.updated", metadata={"permission_id": permission.id, "access_level": payload.access_level, "user_id": payload.user_id, "role_id": payload.role_id})
        db.commit()
        db.refresh(permission)
        username = None
        user_email = None
        role_name = None
        if permission.user_id:
            u = db.query(User).filter(User.id == permission.user_id).first()
            if u:
                username = u.username
                user_email = u.email
        if permission.role_id:
            from app.models.role import Role
            r = db.query(Role).filter(Role.id == permission.role_id).first()
            if r:
                role_name = r.name
        setattr(permission, 'username', username)
        setattr(permission, 'user_email', user_email)
        setattr(permission, 'role_name', role_name)
        return permission

    @staticmethod
    def delete_board_permission(db: Session, board_id: int, permission_id: int, current_user: User) -> Dict[str, str]:
        KanbanService._require_global_admin(current_user)
        board = KanbanService.get_board_detail(db, board_id, current_user)
        permission = db.query(KanbanBoardPermission).filter(KanbanBoardPermission.board_id == board_id, KanbanBoardPermission.id == permission_id).first()
        if not permission:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Permissao nao encontrada.")
        if permission.user_id == board.created_by_user_id and permission.access_level == "ADMIN":
            raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Nao remova o acesso ADMIN do criador do quadro.")
        from app.core.audit import log_action
        db.delete(permission)
        log_action(
            db=db,
            user_id=current_user.id,
            action="kanban.board_access.removed",
            module="kanban",
            details={"permission_id": permission_id, "board_id": board_id},
            commit=False
        )
        record_activity(db, board_id, current_user, "access.removed", metadata={"permission_id": permission_id})
        db.commit()
        return {"status": "success"}

    @staticmethod
    async def preview_import(db: Session, board_id: int, upload: UploadFile, current_user: User) -> Dict[str, Any]:
        board = KanbanService.get_board_detail(db, board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.NORMAL)
        suffix = Path(upload.filename or "").suffix.lower()
        if suffix not in {".csv", ".xlsx"}:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Importe apenas arquivos CSV ou XLSX.")
        data = await upload.read()
        if len(data) > MAX_UPLOAD_BYTES:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Arquivo acima do limite de 10 MB.")
        rows: List[Dict[str, Any]] = []
        if suffix == ".csv":
            text = data.decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(text))
            rows = [dict(row) for row in reader]
            headers = reader.fieldnames or []
        else:
            try:
                from openpyxl import load_workbook
            except ImportError as exc:
                raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail="Dependencia openpyxl nao instalada.") from exc
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            sheet = workbook.active
            header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), [])
            headers = [str(cell).strip() for cell in header_row if cell is not None]
            for values in sheet.iter_rows(min_row=2, values_only=True):
                row = {headers[index]: values[index] for index in range(min(len(headers), len(values)))}
                if any(value is not None and str(value).strip() != "" for value in row.values()):
                    rows.append(row)
        if not headers:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Arquivo sem cabecalho.")
        import_id = uuid.uuid4().hex
        IMPORT_DIR.mkdir(parents=True, exist_ok=True)
        (IMPORT_DIR / f"{import_id}.json").write_text(json.dumps({"headers": headers, "rows": rows}, default=str), encoding="utf-8")
        record_activity(db, board_id, current_user, "import.previewed", metadata={"filename": upload.filename, "rows": len(rows)})
        db.commit()
        return {"import_id": import_id, "filename": upload.filename or "arquivo", "headers": headers, "rows": rows[:25], "total_rows": len(rows)}

    @staticmethod
    def confirm_import(db: Session, board_id: int, payload: BoardImportConfirmPayload, current_user: User) -> Dict[str, Any]:
        board = KanbanService.get_board_detail(db, board_id, current_user)
        require_board_level(db, board, current_user, BoardAccessLevel.NORMAL)
        path = IMPORT_DIR / f"{payload.import_id}.json"
        if not path.exists():
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Previa de importacao expirada ou inexistente.")
        stored = json.loads(path.read_text(encoding="utf-8"))
        rows: List[Dict[str, Any]] = stored["rows"]
        custom_fields = {field.key: field for field in board.custom_fields if field.is_active}
        columns_by_name = {column.name.lower(): column for column in board.columns if not column.is_archived}
        default_column_id = payload.default_column_id or (board.columns[0].id if board.columns else None)
        if not default_column_id:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Quadro sem coluna disponivel para importacao.")
        created = 0
        updated = 0
        ignored = 0
        for row in rows:
            mapped: Dict[str, Any] = {}
            custom_values: Dict[str, Any] = {}
            for source, target in payload.mapping.items():
                value = row.get(source)
                if value is None or str(value).strip() == "":
                    continue
                if target.startswith("custom."):
                    key = target.split(".", 1)[1]
                    if key in custom_fields:
                        custom_values[key] = value
                else:
                    mapped[target] = value
            title = str(mapped.get("title") or custom_values.get("op") or "Card importado").strip()
            duplicate_card = None
            if payload.duplicate_field:
                duplicate_value = mapped.get(payload.duplicate_field) or custom_values.get(payload.duplicate_field)
                if duplicate_value is not None:
                    duplicate_card = next((card for column in board.columns for card in column.cards if (card.custom_fields or {}).get(payload.duplicate_field) == duplicate_value or getattr(card, payload.duplicate_field, None) == duplicate_value), None)
            if duplicate_card and payload.duplicate_strategy == "ignore":
                ignored += 1
                continue
            target_column_id = default_column_id
            if "column" in mapped and str(mapped["column"]).lower() in columns_by_name:
                target_column_id = columns_by_name[str(mapped["column"]).lower()].id
            due_date = mapped.get("due_date")
            parsed_due = None
            if due_date:
                try:
                    parsed_due = datetime.fromisoformat(str(due_date))
                except ValueError:
                    parsed_due = None
            if duplicate_card and payload.duplicate_strategy == "update":
                duplicate_card.title = title
                duplicate_card.description = mapped.get("description", duplicate_card.description)
                duplicate_card.priority = str(mapped.get("priority", duplicate_card.priority)).upper()
                duplicate_card.due_date = parsed_due or duplicate_card.due_date
                duplicate_card.custom_fields = {**(duplicate_card.custom_fields or {}), **custom_values}
                duplicate_card.updated_at = datetime.now(timezone.utc)
                updated += 1
                continue
            card = KanbanService.create_card(db, board_id, CardCreate(
                column_id=target_column_id,
                title=title,
                description=mapped.get("description"),
                priority=str(mapped.get("priority", "MEDIUM")).upper() if str(mapped.get("priority", "MEDIUM")).upper() in CARD_PRIORITIES else "MEDIUM",
                due_date=parsed_due,
                custom_fields=custom_values or None,
            ), current_user)
            created += 1
        path.unlink(missing_ok=True)
        record_activity(db, board_id, current_user, "import.completed", metadata={"created": created, "updated": updated, "ignored": ignored})
        db.commit()
        return {"status": "success", "created": created, "updated": updated, "ignored": ignored}

    @staticmethod
    def quick_card(db: Session, payload: QuickCardPayload, current_user: User) -> Dict[str, Any]:
        if not payload.confirm:
            return {"requires_confirmation": True, "parsed": KanbanService.parse_quick_card(payload.text, payload.board_id, payload.column_id)}
        parsed = KanbanService.parse_quick_card(payload.text, payload.board_id, payload.column_id)
        board = None
        if parsed["board_id"]:
            board = KanbanService.get_board_detail(db, parsed["board_id"], current_user)
        elif parsed["board_slug"]:
            board = next((item for item in KanbanService.list_boards(db, current_user) if item.slug == parsed["board_slug"]), None)
        if not board:
            boards = KanbanService.list_boards(db, current_user)
            board = boards[0] if boards else None
        if not board:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Nenhum board acessivel para criar o card.")
        require_board_level(db, board, current_user, BoardAccessLevel.NORMAL)
        column_id = parsed["column_id"] or (board.columns[0].id if board.columns else None)
        if not column_id:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Board sem coluna disponivel.")
        card = KanbanService.create_card(db, board.id, CardCreate(
            column_id=column_id,
            title=parsed["title"],
            priority=parsed["priority"],
            due_date=parsed["due_date"],
        ), current_user)
        return {"requires_confirmation": False, "card": card, "board_id": board.id}

    @staticmethod
    def parse_quick_card(text: str, board_id: Optional[int], column_id: Optional[int]) -> Dict[str, Any]:
        clean = text.strip()
        board_slug = None
        lowered = clean.lower()
        for prefix, slug in [("producao:", "producao"), ("produção:", "producao"), ("ti:", "ti"), ("projetos:", "projetos"), ("card:", None)]:
            if lowered.startswith(prefix):
                board_slug = slug
                clean = clean[len(prefix):].strip()
                break
        priority = "MEDIUM"
        if "prioridade urgente" in lowered or "urgente" in lowered:
            priority = "URGENT"
        elif "prioridade alta" in lowered or "alta" in lowered:
            priority = "HIGH"
        due_date = None
        now = datetime.now(timezone.utc).replace(tzinfo=None)
        if "prazo amanhã" in lowered or "prazo amanha" in lowered:
            due_date = now + timedelta(days=1)
        elif "prazo hoje" in lowered or "hoje" in lowered:
            due_date = now
        return {"title": clean or text.strip(), "board_id": board_id, "board_slug": board_slug, "column_id": column_id, "priority": priority, "due_date": due_date}

    @staticmethod
    def get_summary(db: Session, current_user: User) -> Dict[str, Any]:
        boards = KanbanService.list_boards(db, current_user, include_archived=False)
        board_ids = [board.id for board in boards]
        if not board_ids:
            return {
                "active_boards": 0,
                "total_cards": 0,
                "open_cards": 0,
                "overdue_cards": 0,
                "assigned_to_me": 0,
                "unassigned_cards": 0,
                "high_priority_cards": 0,
                "recently_updated_cards": 0,
                "cards_by_priority": {"LOW": 0, "MEDIUM": 0, "HIGH": 0, "URGENT": 0},
                "cards_by_board": {},
                "simple_pending": 0,
            }

        cards_query = db.query(KanbanCard).filter(KanbanCard.board_id.in_(board_ids), KanbanCard.is_archived == False)
        open_cards = cards_query.join(KanbanColumn, KanbanColumn.id == KanbanCard.column_id).filter(KanbanColumn.is_done_column == False).count()
        now = datetime.now(timezone.utc).replace(tzinfo=None)
        overdue_cards = cards_query.join(KanbanColumn, KanbanColumn.id == KanbanCard.column_id).filter(
            KanbanColumn.is_done_column == False,
            KanbanCard.due_date != None,
            KanbanCard.due_date < now,
        ).count()
        by_priority = {"LOW": 0, "MEDIUM": 0, "HIGH": 0, "URGENT": 0}
        for priority, count in db.query(KanbanCard.priority, func.count(KanbanCard.id)).filter(
            KanbanCard.board_id.in_(board_ids),
            KanbanCard.is_archived == False,
        ).group_by(KanbanCard.priority).all():
            by_priority[priority] = count
        total_cards = cards_query.count()
        assigned_to_me = cards_query.join(KanbanCardAssignee, KanbanCardAssignee.card_id == KanbanCard.id).filter(KanbanCardAssignee.user_id == current_user.id).count()
        unassigned_cards = cards_query.outerjoin(KanbanCardAssignee, KanbanCardAssignee.card_id == KanbanCard.id).filter(KanbanCardAssignee.id == None).count()
        high_priority_cards = cards_query.filter(KanbanCard.priority.in_(["HIGH", "URGENT"])).count()
        recent_cutoff = datetime.now(timezone.utc).replace(tzinfo=None)
        recently_updated_cards = cards_query.filter(KanbanCard.updated_at >= recent_cutoff.replace(hour=0, minute=0, second=0, microsecond=0)).count()
        cards_by_board = {
            name: count for name, count in db.query(KanbanBoard.name, func.count(KanbanCard.id))
            .join(KanbanCard, KanbanCard.board_id == KanbanBoard.id)
            .filter(KanbanBoard.id.in_(board_ids), KanbanCard.is_archived == False)
            .group_by(KanbanBoard.name)
            .all()
        }
        return {
            "active_boards": len(boards),
            "total_cards": total_cards,
            "open_cards": open_cards,
            "overdue_cards": overdue_cards,
            "assigned_to_me": assigned_to_me,
            "unassigned_cards": unassigned_cards,
            "high_priority_cards": high_priority_cards,
            "recently_updated_cards": recently_updated_cards,
            "cards_by_priority": by_priority,
            "cards_by_board": cards_by_board,
            "simple_pending": open_cards + overdue_cards,
        }
