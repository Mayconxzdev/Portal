import openpyxl
import hashlib
import uuid
import re
import unicodedata
from datetime import datetime, timezone
from typing import Optional, Dict, Any
from sqlalchemy.orm import Session
from app.models.stock import (
    StockCatalogImportRun,
    StockCatalogTreeNode,
    StockCatalogItem,
    StockCatalogItemSpec,
    StockCatalogSupplier,
    StockCatalogOffer,
    StockCatalogPriceHistory,
    StockCatalogSearchIndex,
    StockCatalogReviewQueue,
    StockCatalogAlert,
    StockCatalogCybersulProduct,
    StockCatalogLink
)
from app.modules.stock.parser import _text, _norm, _money, _email, _phone
from app.modules.stock.measure_utils import normalize_measure, strip_measurements_from_name, is_invalid_product_name, humanize_product_name

def _yes(value: Any) -> bool:
    text = _norm(value)
    return text in {"sim", "s", "true", "1", "yes", "active", "ativo"}


SUPPLIER_SUSPECT_TERMS = {
    "aco",
    "aço",
    "aluminio",
    "alumínio",
    "cantoneira",
    "chapa",
    "ferro",
    "inox",
    "material",
    "parafuso",
    "produto",
    "rodizio",
    "rodízio",
    "tubo",
    "vergalhao",
    "vergalhão",
}


def _is_suspicious_supplier_name(value: Optional[str]) -> bool:
    normalized = _norm(value)
    if not normalized:
        return True
    if "@" in normalized or any(term in normalized for term in ["ltda", "eireli", "comerc", "dist", "parafus", "ferragens"]):
        return False
    tokens = set(normalized.split())
    return bool(tokens & SUPPLIER_SUSPECT_TERMS)

def normalize_category_title(title: str) -> str:
    if not title:
        return ""
    title = " ".join(title.split())
    title_lower = title.lower()
    
    if "nao-usar" in title_lower or "nao usar" in title_lower:
        return ""
        
    title = re.sub(r'^[0-9.]+\s+', '', title)
    title = re.sub(r'^[XZxz]\s+', '', title)
    
    title_lower = title.lower()
    
    mappings = {
        "rodizios": "Rodízios",
        "rodizios inox": "Rodízios Inox",
        "mat. eletrico ex": "Material Elétrico Ex",
        "mat. p solda": "Material para Solda",
        "materias-primas": "Matéria-prima",
        "materia- prima": "Matéria-prima",
        "materia prima": "Matéria-prima",
        "materias primas": "Matéria-prima",
        "pvc": "PVC",
        "exaustor ex vesper": "Exaustor Vesper Ex",
        "climatizador vesper": "Climatizador Vesper",
        "subpecas": "Subpeças",
        "pecas de consumo": "Peças de Consumo",
        "ferramentas de precisao": "Ferramentas de Precisão",
        "dutos flexiveis": "Dutos Flexíveis",
        "rodas e rodizios": "Rodas e Rodízios"
    }
    
    def strip_accents(s):
       return ''.join(c for c in unicodedata.normalize('NFD', s)
                      if unicodedata.category(c) != 'Mn')
                      
    norm_key = strip_accents(title_lower)
    if norm_key in mappings:
        return mappings[norm_key]
        
    words = title.split()
    new_words = []
    for w in words:
        wl = w.lower()
        if wl in ["e", "de", "para", "o", "a", "os", "as", "com", "sem", "p/", "pp", "ex", "vesper", "inox", "pvc", "unc", "c/"]:
            if wl == "ex":
                new_words.append("Ex")
            elif wl == "vesper":
                new_words.append("Vesper")
            elif wl == "inox":
                new_words.append("Inox")
            elif wl == "pvc":
                new_words.append("PVC")
            elif wl in ["e", "de", "para", "o", "a", "os", "as", "com", "sem", "p/", "c/"]:
                new_words.append(wl)
            else:
                new_words.append(w.upper())
        else:
            new_words.append(w.capitalize())
            
    res = " ".join(new_words)
    if "Materia- prima" in res or "Materia prima" in res or "Matéria Prima" in res:
        res = "Matéria-prima"
    if "Conexões De Plástico / Metal" in res:
        res = "Conexões Plástico/Metal"
        
    return res

CANONICAL_CATEGORIES = {
    "PRODUTO_FINAL": "Produto Final",
    "PRODUTO_REVENDA": "Produto Revenda",
    "MATERIA_PRIMA": "Matéria-prima",
    "MATERIAL_ELETRICO": "Material Elétrico",
    "MATERIAL_ELETRICO_EX": "Material Elétrico Ex",
    "FIXADORES": "Fixadores",
    "FIXADORES_INOX": "Fixadores Inox",
    "TUBOS": "Tubos",
    "CHAPAS": "Chapas",
    "BARRAS_CANT_TARUGOS": "Barras, Cantoneiras e Tarugos",
    "RODAS_RODIZIOS": "Rodas e Rodízios",
    "CONEXOES_ALTA_PRESSAO": "Conexões Alta Pressão",
    "CONEXOES_PLASTICO_METAL": "Conexões Plástico/Metal",
    "HIDRAULICA": "Hidráulica",
    "ABRASIVOS": "Abrasivos",
    "TINTAS": "Tintas",
    "ADESIVOS_PLACAS": "Adesivos, Placas e Etiquetas",
    "MATERIAL_ESCRITORIO": "Material de Escritório",
    "MATERIAL_LIMPEZA": "Material de Limpeza",
    "MATERIAL_EMBALAGEM": "Material de Embalagem",
    "BENEFICIAMENTO": "Beneficiamento",
    "MAO_DE_OBRA": "Mão de Obra",
    "PECAS_CONSUMO": "Peças de Consumo",
    "OUTROS": "Outros"
}

def get_canonical_category(source_sheet: str, family: str) -> str:
    text = f"{source_sheet or ''} {family or ''}".lower()
    
    if "nao-usar" in text or "nao usar" in text:
        return "NAO_USAR"
        
    if "climatizador" in text or "exaustor" in text or "ventilador" in text or "qualitas" in text:
        return "PRODUTO_FINAL"
    if "revenda" in text:
        return "PRODUTO_REVENDA"
    if "eletrico ex" in text or "eletrica ex" in text or "cpex" in text or "plugue ceag" in text or "condulete" in text:
        return "MATERIAL_ELETRICO_EX"
    if "eletrico" in text or "eletrica" in text or "cabo" in text or "fio" in text or "plugue" in text or "tomada" in text or "prensa-cabo" in text or "conduite" in text or "sinaleira" in text or "contator" in text or "cooper" in text:
        return "MATERIAL_ELETRICO"
    if "solda" in text or "eletrodo" in text or "materia prima" in text or "materia-prima" in text or "materia- prima" in text:
        return "MATERIA_PRIMA"
    if "fixadores inox" in text or "inox" in text and ("parafuso" in text or "porca" in text or "rebite" in text):
        return "FIXADORES_INOX"
    if "fixador" in text or "parafuso" in text or "porca" in text or "rebite" in text or "prego" in text:
        return "FIXADORES"
    if "tubo" in text:
        return "TUBOS"
    if "chapa" in text:
        return "CHAPAS"
    if "barra" in text or "cantoneira" in text or "tarugo" in text or "eslinga" in text:
        return "BARRAS_CANT_TARUGOS"
    if "rodiz" in text or "roda" in text:
        return "RODAS_RODIZIOS"
    if "conexao plastico" in text or "conexao metal" in text or "conexoes de plastico" in text:
        return "CONEXOES_PLASTICO_METAL"
    if "conexao" in text or "conexoes" in text or "niple" in text or "cotovelo" in text or "flange" in text or "valvula" in text:
        return "CONEXOES_ALTA_PRESSAO"
    if "hidraul" in text or "lubrifil" in text or "mancal" in text or "correia" in text or "mola" in text:
        return "HIDRAULICA"
    if "abrasiv" in text or "disco" in text or "lente" in text or "liquido penetrante" in text:
        return "ABRASIVOS"
    if "tinta" in text or "spray" in text or "colar" in text or "cola" in text or "adesivo" in text:
        return "TINTAS"
    if "placa" in text or "etiqueta" in text or "plaqueta" in text or "identificador" in text:
        return "ADESIVOS_PLACAS"
    if "escritorio" in text:
        return "MATERIAL_ESCRITORIO"
    if "limpeza" in text:
        return "MATERIAL_LIMPEZA"
    if "embalagem" in text or "stretch" in text or "fita" in text or "papel velumoid" in text or "pallet" in text or "grampo" in text or "selo" in text:
        return "MATERIAL_EMBALAGEM"
    if "beneficiamento" in text:
        return "BENEFICIAMENTO"
    if "mao de obra" in text or "aluguel" in text:
        return "MAO_DE_OBRA"
    if "consumo" in text or "subpecas" in text or "pecas" in text:
        return "PECAS_CONSUMO"
        
    return "OUTROS"

class StockCatalogUnifiedImporter:
    @staticmethod
    def import_unified(filepath: str, db: Session, user_id: Optional[int] = None) -> StockCatalogImportRun:
        """
        Importa todo o catalogo a partir da planilha unificada mestre.
        Mapeia status de qualidade (READY, ENRICHED, NEEDS_REVIEW, QUARANTINED, HIDDEN, INACTIVE)
        e organiza a Central de Saneamento.
        """
        # 1. Calcular hash do arquivo
        hasher = hashlib.sha256()
        with open(filepath, "rb") as f:
            for chunk in iter(lambda: f.read(65536), b""):
                hasher.update(chunk)
        file_hash = hasher.hexdigest()
        filename = filepath.split("\\")[-1].split("/")[-1]

        # 2. Criar import run no banco
        run = StockCatalogImportRun(
            source_type="UNIFIED_Mestre",
            source_path=filepath,
            source_filename=filename,
            source_hash=file_hash,
            status="RUNNING",
            total_rows=0,
            total_items=0,
            total_offers=0,
            total_errors=0,
            total_review=0,
            triggered_by_user_id=user_id
        )
        db.add(run)
        db.commit()

        try:
            wb = openpyxl.load_workbook(filepath, data_only=True)

            for model in (
                StockCatalogAlert,
                StockCatalogReviewQueue,
                StockCatalogSearchIndex,
                StockCatalogPriceHistory,
                StockCatalogOffer,
                StockCatalogLink,
                StockCatalogItemSpec,
                StockCatalogItem,
                StockCatalogTreeNode,
                StockCatalogCybersulProduct,
            ):
                db.query(model).delete(synchronize_session=False)
            db.flush()

            # --- PARTE 1: IMPORTAR FORNECEDORES ---
            print("Importando fornecedores...")
            ws_forn = wb["FORNECEDORES"]
            forn_rows = list(ws_forn.iter_rows(values_only=True))
            forn_header = [str(x).strip().lower() if x is not None else "" for x in forn_rows[0]]
            
            forn_map = {} # nome_normalizado -> Supplier
            for r_idx in range(1, len(forn_rows)):
                row = forn_rows[r_idx]
                if not row or not row[1]:
                    continue
                
                name_raw = _text(row[1])
                norm_name = _norm(name_raw)
                email_raw = _email(row[4]) if len(row) > 4 else None
                phone_raw = _phone(row[5]) if len(row) > 5 else None
                contact_raw = _text(row[6]) if len(row) > 6 else None
                
                supplier = db.query(StockCatalogSupplier).filter(
                    StockCatalogSupplier.normalized_name == norm_name
                ).first()
                
                if not supplier:
                    supplier = StockCatalogSupplier(
                        name=name_raw[:255],
                        normalized_name=norm_name[:255],
                        email=email_raw[:255] if email_raw else None,
                        phone=phone_raw[:50] if phone_raw else None,
                        contact_name=contact_raw[:100] if contact_raw else None,
                        source="UNIFIED_IMPORT",
                        active=True
                    )
                    db.add(supplier)
                    db.flush()
                else:
                    if email_raw:
                        supplier.email = email_raw
                    if phone_raw:
                        supplier.phone = phone_raw
                    if contact_raw:
                        supplier.contact_name = contact_raw
                
                forn_map[norm_name] = supplier

            # --- PARTE 2: RECONSTRUIR ÁRVORE E ITENS ---
            print("Importando itens do catálogo mestre...")
            ws_mestre = wb["CATALOGO_MESTRE"]
            mestre_rows = list(ws_mestre.iter_rows(values_only=True))
            m_header = [str(x).strip().lower() if x is not None else "" for x in mestre_rows[0]]
            
            # Map headers to index
            col_map = {name: idx for idx, name in enumerate(m_header)}
            
            total_items = 0
            total_rows = 0
            total_review = 0
            
            tree_cache = {}
            items_cache = {}

            from app.modules.stock.measure_utils import classify_string, StringClassification
            from app.modules.stock.parser_context import ParserContext, is_attribute_only

            ctx = ParserContext()
            current_sheet = None
            current_family = None

            for r_idx in range(1, len(mestre_rows)):
                row = mestre_rows[r_idx]
                if not row or len(row) < 3:
                    continue
                
                total_rows += 1
                
                # Campos da planilha mestre
                status_imp = _text(row[col_map.get("status_importacao", 1)])
                precisa_rev = _text(row[col_map.get("precisa_revisao", 2)])
                motivo_rev = _text(row[col_map.get("motivo_revisao", 3)])
                conf_vinculo = row[col_map.get("confianca_vinculo", 4)]
                cod_cybersul = _text(row[col_map.get("codigo_cybersul", 5)])
                cod_antigo = _text(row[col_map.get("codigo_antigo", 6)])
                desc_cybersul = _text(row[col_map.get("descricao_oficial_cybersul", 7)])
                compl_cybersul = _text(row[col_map.get("complemento_cybersul", 8)])
                unidade = _text(row[col_map.get("unidade", 9)]) or "UN"
                grupo_cybersul = _text(row[col_map.get("grupo_cybersul", 10)])
                ncm = _text(row[col_map.get("ncm", 11)])
                ativo_cybersul = _text(row[col_map.get("ativo_cybersul", 12)])
                
                saldo_vesp = _money(row[col_map.get("saldo_vesper", 13)]) or 0.0
                saldo_vent = _money(row[col_map.get("saldo_ventrio", 14)]) or 0.0
                saldo_tot = _money(row[col_map.get("saldo_total", 15)]) or 0.0
                
                fam_comercial = _text(row[col_map.get("familia_comercial", 16)])
                prod_comercial = _text(row[col_map.get("produto_comercial", 17)])
                nome_limpo = _text(row[col_map.get("nome_limpo_portal", 18)])
                variacao = _text(row[col_map.get("variacao", 19)])
                medida_norm = _text(row[col_map.get("medida_normalizada", 20)])
                especificacao = _text(row[col_map.get("especificacao", 21)])
                
                forn_cyber = _text(row[col_map.get("fornecedor_cybersul", 22)])
                forn_principal = _text(row[col_map.get("fornecedor_principal", 23)])
                preco_val = _money(row[col_map.get("preco_ultima_compra_ou_cotacao", 25)])
                moeda = _text(row[col_map.get("moeda", 26)]) or "BRL"
                origem_preco = _text(row[col_map.get("origem_preco", 27)])
                
                aba_origem = _text(row[col_map.get("aba_origem_compras_nova", 31)])
                linha_origem = row[col_map.get("linha_origem_compras_nova", 32)]
                source_hash = _text(row[col_map.get("source_hash", 36)])
                
                sheet_title = aba_origem or grupo_cybersul or "Outros"
                family_title = fam_comercial or grupo_cybersul or "Geral"
                
                # Reset state trackers on sheet/family change
                if sheet_title != current_sheet or family_title != current_family:
                    current_sheet = sheet_title
                    current_family = family_title
                    ctx.reset_on_family(family_title)
                    ctx.last_valid_category = sheet_title

                prod_class = classify_string(prod_comercial)
                var_class = classify_string(variacao)
                spec_class = classify_string(especificacao)

                is_quarantine_row = False
                row_raw_content = " ".join([str(val) for val in row if val is not None]).lower()
                if any(k in row_raw_content for k in ["subtotal", "total", "média", "media", "fórmula", "formula", "#nome?"]):
                    is_quarantine_row = True
                if prod_class in {StringClassification.FORMULA_ERROR, StringClassification.SUBTOTAL, StringClassification.HEADER}:
                    is_quarantine_row = True

                is_prod_invalid = is_invalid_product_name(prod_comercial)

                if prod_comercial and prod_class == StringClassification.PRODUCT_NAME:
                    thickness_match = re.search(
                        r"\b(\d+(?:[.,]\d+)?\s*mm|\d+/\d+\s*mm|\d+/\d+\")", prod_comercial, re.IGNORECASE
                    )
                    if thickness_match:
                        ctx.apply_thickness(thickness_match.group(1))
                    ctx.last_valid_product = humanize_product_name(prod_comercial)
                elif prod_comercial and is_attribute_only(prod_class):
                    from app.modules.stock.parser import _apply_row_classification_to_context
                    _apply_row_classification_to_context(ctx, prod_class, prod_comercial)

                if variacao and var_class == StringClassification.DIMENSION:
                    ctx.apply_dimension(variacao)
                elif variacao and var_class == StringClassification.WEIGHT:
                    ctx.apply_weight(variacao)
                elif variacao and var_class == StringClassification.LENGTH_GROUP:
                    ctx.reset_on_length_group(variacao)

                if especificacao and spec_class == StringClassification.WEIGHT:
                    ctx.apply_weight(especificacao)
                elif especificacao and spec_class == StringClassification.DIMENSION:
                    ctx.apply_dimension(especificacao)

                prod_clean = ctx.last_valid_product
                if is_prod_invalid or not prod_comercial:
                    if ctx.last_valid_product:
                        prod_clean = ctx.last_valid_product
                    else:
                        fallback_cand = family_title
                        for cand in [desc_cybersul, nome_limpo]:
                            if cand and not is_invalid_product_name(cand):
                                fallback_cand = cand
                                break
                        prod_clean = humanize_product_name(fallback_cand)

                if is_prod_invalid and not ctx.last_valid_product and not cod_cybersul:
                    is_quarantine_row = True

                var_raw = variacao
                is_var_invalid = var_class in {
                    StringClassification.FORMULA_ERROR,
                    StringClassification.HEADER,
                    StringClassification.SUBTOTAL,
                    StringClassification.EMPTY,
                    StringClassification.APPLICATION_NOTE
                }

                var_clean = var_raw
                if is_var_invalid or not var_raw:
                    if ctx.last_valid_dimension:
                        var_clean = ctx.last_valid_dimension
                    else:
                        var_clean = ctx.last_valid_thickness or "Geral"
                elif var_class == StringClassification.DIMENSION:
                    var_clean = variacao

                display_name = str(prod_clean or "")
                if ctx.last_valid_thickness and ctx.last_valid_thickness not in display_name:
                    display_name += f" - {ctx.last_valid_thickness}"
                if var_clean and var_clean != ctx.last_valid_thickness and var_clean not in display_name:
                    display_name += f" - {var_clean}"

                base_name = prod_clean
                product_title = prod_clean or (display_name if display_name else family_title)

                extracted_weight = ctx.last_valid_weight
                if not extracted_weight:
                    for candidate in [prod_comercial, variacao, especificacao]:
                        if candidate:
                            match = re.search(
                                r"(\d+(?:[.,]\d+)?\s*kg/(?:m|barra|ch|pç|pc|barra|chapa|pça|t|g)|\d+(?:[.,]\d+)?\s*kg\b)",
                                candidate,
                                re.IGNORECASE,
                            )
                            if match:
                                extracted_weight = match.group(1).strip()
                                break

                # Nomes crus e limpos para classificação e exibição
                display_name = str(prod_clean or "")
                if ctx.last_valid_thickness and ctx.last_valid_thickness not in display_name:
                    display_name += f" - {ctx.last_valid_thickness}"
                if var_clean and var_clean != ctx.last_valid_thickness and var_clean not in display_name:
                    display_name += f" - {var_clean}"

                base_name = prod_clean
                product_title = prod_clean or (display_name if display_name else family_title)
                
                # Calcular categoria canônica
                canonical_key = get_canonical_category(sheet_title, family_title)
                canonical_display = CANONICAL_CATEGORIES.get(canonical_key, "Outros")
                
                # Normalizar títulos de árvore
                norm_sheet = normalize_category_title(sheet_title) or "Outros"
                norm_family = normalize_category_title(family_title) or "Geral"
                norm_product = normalize_category_title(product_title) or display_name
                
                # Visibilidade comum da categoria
                is_visible_common = (canonical_key != "NAO_USAR" and "nao-usar" not in norm_sheet.lower() and "nao-usar" not in norm_family.lower())
                
                # Nó de Aba (depth 0)
                sheet_key = sheet_title
                if sheet_key not in tree_cache:
                    node = db.query(StockCatalogTreeNode).filter(
                        StockCatalogTreeNode.node_type == "sheet",
                        StockCatalogTreeNode.title == sheet_title
                    ).first()
                    if not node:
                        node = StockCatalogTreeNode(
                            import_run_id=run.id,
                            source_type="UNIFIED",
                            source_sheet=sheet_title[:100],
                            parent_id=None,
                            node_type="sheet",
                            title=sheet_title[:255],
                            normalized_title=_norm(sheet_title)[:255],
                            path=sheet_title[:1000],
                            depth=0,
                            canonical_category=canonical_key,
                            canonical_category_display=canonical_display,
                            is_visible_to_common_user=is_visible_common
                        )
                        db.add(node)
                        db.flush()
                    else:
                        node.canonical_category = canonical_key
                        node.canonical_category_display = canonical_display
                        node.is_visible_to_common_user = is_visible_common
                        node.import_run_id = run.id
                        db.flush()
                    tree_cache[sheet_key] = node
                sheet_node = tree_cache[sheet_key]

                # Nó de Família (depth 1)
                family_key = (sheet_title, family_title)
                if family_key not in tree_cache:
                    node = db.query(StockCatalogTreeNode).filter(
                        StockCatalogTreeNode.node_type == "family",
                        StockCatalogTreeNode.title == family_title,
                        StockCatalogTreeNode.parent_id == sheet_node.id
                    ).first()
                    if not node:
                        node = StockCatalogTreeNode(
                            import_run_id=run.id,
                            source_type="UNIFIED",
                            source_sheet=sheet_title[:100],
                            parent_id=sheet_node.id,
                            node_type="family",
                            title=family_title[:255],
                            normalized_title=_norm(family_title)[:255],
                            path=f"{sheet_title} > {family_title}"[:1000],
                            depth=1,
                            canonical_category=canonical_key,
                            canonical_category_display=canonical_display,
                            is_visible_to_common_user=is_visible_common
                        )
                        db.add(node)
                        db.flush()
                    else:
                        node.canonical_category = canonical_key
                        node.canonical_category_display = canonical_display
                        node.is_visible_to_common_user = is_visible_common
                        node.import_run_id = run.id
                        db.flush()
                    tree_cache[family_key] = node
                family_node = tree_cache[family_key]

                # Nó de Produto (depth 2)
                product_title_str = product_title or family_title
                product_key = (sheet_title, family_title, product_title_str)
                if product_key not in tree_cache:
                    node = db.query(StockCatalogTreeNode).filter(
                        StockCatalogTreeNode.node_type == "product",
                        StockCatalogTreeNode.title == product_title_str,
                        StockCatalogTreeNode.parent_id == family_node.id
                    ).first()
                    if not node:
                        node = StockCatalogTreeNode(
                            import_run_id=run.id,
                            source_type="UNIFIED",
                            source_sheet=sheet_title[:100],
                            parent_id=family_node.id,
                            node_type="product",
                            title=product_title_str[:255],
                            normalized_title=_norm(product_title_str)[:255],
                            path=f"{sheet_title} > {family_title} > {product_title_str}"[:1000],
                            depth=2,
                            canonical_category=canonical_key,
                            canonical_category_display=canonical_display,
                            is_visible_to_common_user=is_visible_common
                        )
                        db.add(node)
                        db.flush()
                    else:
                        node.canonical_category = canonical_key
                        node.canonical_category_display = canonical_display
                        node.is_visible_to_common_user = is_visible_common
                        node.import_run_id = run.id
                        db.flush()
                    tree_cache[product_key] = node
                product_node = tree_cache[product_key]

                # --- B. Importar Produto do Cybersul ---
                cyber_prod = None
                if cod_cybersul:
                    cyber_prod = db.query(StockCatalogCybersulProduct).filter(
                        StockCatalogCybersulProduct.cybersul_code == cod_cybersul
                    ).first()
                    
                    if not cyber_prod:
                        cyber_prod = StockCatalogCybersulProduct(
                            cybersul_code=cod_cybersul,
                            old_code=cod_antigo,
                            description=desc_cybersul,
                            normalized_description=_norm(desc_cybersul),
                            complement=compl_cybersul,
                            unit=unidade,
                            balance_vesper=saldo_vesp,
                            balance_ventrio=saldo_vent,
                            balance_total=saldo_tot,
                            cost_price=preco_val if "Cybersul" in origem_preco else None,
                            supplier_name=forn_cyber or forn_principal,
                            ncm=ncm,
                            group_name=grupo_cybersul,
                            active=_yes(ativo_cybersul),
                            source_row=r_idx + 1,
                            import_run_id=run.id
                        )
                        db.add(cyber_prod)
                        db.flush()
                    else:
                        cyber_prod.balance_vesper = saldo_vesp
                        cyber_prod.balance_ventrio = saldo_vent
                        cyber_prod.balance_total = saldo_tot
                        cyber_prod.active = _yes(ativo_cybersul)
                        cyber_prod.import_run_id = run.id
                        db.flush()

                # --- C. Importar Item do Catálogo com Regras de UX ---
                display_name_clean = str(display_name or prod_clean or family_title)
                base_name_clean = prod_clean or display_name_clean or family_title
                measure_norm = normalize_measure(var_clean, medida_norm, especificacao, display_name_clean)
                measure_key = measure_norm.canonical_key
                measure_display = measure_norm.display
                measure_kind = measure_norm.kind
                measure_aliases = measure_norm.aliases
                canonical_identity_key = _norm(
                    "|".join(
                        [
                            canonical_key or "",
                            norm_sheet or "",
                            norm_family or "",
                            norm_product or "",
                            measure_key or _norm(var_clean or especificacao or display_name_clean),
                        ]
                    )
                )[:500]
                
                # Identificar ruídos e junk do parser
                name_stripped = (display_name or "").strip()
                is_junk = False
                if name_stripped.isdigit():
                    is_junk = True
                if name_stripped.lower() in ("diametro o mm", "diametro ø mm", "material / descricao", "material / descrição", "nao-usar", "nao usar"):
                    is_junk = True

                # Determinando a identidade
                identity_hash = source_hash or hashlib.sha256(
                    f"{sheet_title}|{family_title}|{base_name_clean}|{var_clean or ''}|{especificacao or ''}|{cod_cybersul or ''}".encode("utf-8")
                ).hexdigest()

                # Determinar o Status de Qualidade e Visibilidade
                quality_status = "READY"
                visibility_scope = "COMMON"
                review_type = "none"
                is_operational = True
                is_searchable_common = True
                is_tree_visible_common = True
                needs_review_bool = (precisa_rev == "SIM")
                is_active = True
                supplier_suspicious = _is_suspicious_supplier_name(forn_principal) if forn_principal else False
                
                if is_quarantine_row or is_junk:
                    quality_status = "QUARANTINED"
                    visibility_scope = "HIDDEN"
                    review_type = "quarantined_invalid_row" if is_quarantine_row else "parser_junk"
                    is_operational = False
                    is_searchable_common = False
                    is_tree_visible_common = False
                    needs_review_bool = True # Mantém pendência no saneamento
                elif status_imp == "INATIVO_CYBERSUL" or _norm(ativo_cybersul) in ("n", "d", "inativo"):
                    quality_status = "INACTIVE"
                    visibility_scope = "ADMIN_ONLY"
                    review_type = "inactive_cybersul"
                    is_active = False
                    is_operational = False
                    is_searchable_common = False
                    is_tree_visible_common = False
                    needs_review_bool = False # não gera pendência manual para resolver
                elif not cod_cybersul:
                    quality_status = "NEEDS_REVIEW"
                    visibility_scope = "COMMON"
                    review_type = "missing_cybersul_code"
                    is_searchable_common = True
                    is_tree_visible_common = True
                    is_operational = True
                elif supplier_suspicious:
                    quality_status = "READY"
                    visibility_scope = "COMMON"
                    review_type = "suspicious_supplier"
                    needs_review_bool = True
                elif conf_vinculo and float(conf_vinculo) >= 90.0:
                    # Auto-aceitar vínculos de alta confiança sem gerar pendência manual
                    if needs_review_bool and ("confianca" in _norm(motivo_rev) or "vinculo" in _norm(motivo_rev) or not motivo_rev):
                        needs_review_bool = False
                    quality_status = "ENRICHED" if status_imp == "CYBERSUL_ENRIQUECIDO_COMPRAS" else "READY"
                elif needs_review_bool:
                    if "fornecedor" in _norm(motivo_rev):
                        if "email" in _norm(motivo_rev):
                            # Warning de email não bloqueia catálogo operacional
                            quality_status = "ENRICHED" if status_imp == "CYBERSUL_ENRIQUECIDO_COMPRAS" else "READY"
                            review_type = "missing_supplier_email"
                        else:
                            quality_status = "READY"
                            visibility_scope = "COMMON"
                            review_type = "missing_supplier"
                            is_operational = True
                    elif any(term in _norm(motivo_rev) for term in ["duplicado", "preco suspeito", "preco_suspeito", "valor_final_usado_errado"]):
                        quality_status = "QUARANTINED"
                        visibility_scope = "HIDDEN"
                        review_type = "possible_duplicate"
                        is_searchable_common = False
                        is_tree_visible_common = False
                        is_operational = False
                    else:
                        quality_status = "NEEDS_REVIEW"
                        visibility_scope = "ADMIN_ONLY"
                        review_type = "low_confidence"
                        is_operational = False
                else:
                    if status_imp == "CYBERSUL_ENRIQUECIDO_COMPRAS":
                        quality_status = "ENRICHED"
                    else:
                        quality_status = "READY"

                item = None
                if canonical_identity_key:
                    item = db.query(StockCatalogItem).filter(
                        StockCatalogItem.canonical_identity_key == canonical_identity_key
                    ).first()
                if not item:
                    item = db.query(StockCatalogItem).filter(
                        StockCatalogItem.identity_hash == identity_hash
                    ).first()

                if not item:
                    total_items += 1
                    item = StockCatalogItem(
                        import_run_id=run.id,
                        source_sheet=sheet_title[:100],
                        family_node_id=family_node.id,
                        product_node_id=product_node.id,
                        display_name=display_name_clean[:255],
                        base_name=base_name_clean[:255],
                        normalized_name=_norm(display_name_clean)[:255],
                        variation_label=var_clean[:100] if var_clean else None,
                        normalized_measure=medida_norm[:100] if medida_norm else None,
                        canonical_measure_key=measure_key[:160] if measure_key else None,
                        measure_display=measure_display[:160] if measure_display else None,
                        measure_kind=measure_kind[:40] if measure_kind else None,
                        measure_aliases_json=measure_aliases,
                        canonical_identity_key=canonical_identity_key,
                        specification_text=especificacao,
                        identity_hash=identity_hash,
                        internal_code=cod_cybersul[:100] if cod_cybersul else None,
                        cybersul_product_id=cyber_prod.id if cyber_prod else None,
                        active=is_active,
                        quality_status=quality_status,
                        needs_review=needs_review_bool,
                        review_reason=motivo_rev,
                        
                        # Novas colunas
                        canonical_category=canonical_key,
                        canonical_category_display=canonical_display,
                        visibility_scope=visibility_scope,
                        review_type=review_type,
                        is_operational=is_operational,
                        is_searchable_common=is_searchable_common,
                        is_tree_visible_common=is_tree_visible_common,
                        is_offer_active=True,
                        is_category_only=False,
                        is_parser_junk=is_junk,
                        metadata_json={
                            "source_hash": source_hash,
                            "status_importacao": status_imp,
                            "origem_preco": origem_preco,
                            "linha_origem_compras_nova": linha_origem,
                            "valor_final_eh_metadado": True,
                            "measure_aliases": measure_aliases,
                            "dedupe_sources": [source_hash or identity_hash],
                            "peso": extracted_weight,
                        }
                    )
                    db.add(item)
                    db.flush()
                else:
                    item.import_run_id = run.id
                    item.active = is_active
                    item.quality_status = quality_status
                    item.needs_review = needs_review_bool
                    item.review_reason = motivo_rev
                    if cyber_prod and not item.cybersul_product_id:
                        item.cybersul_product_id = cyber_prod.id
                    if cod_cybersul and not item.internal_code:
                        item.internal_code = cod_cybersul[:100]
                    
                    # Atualizar novas colunas
                    item.display_name = display_name_clean[:255]
                    item.base_name = base_name_clean[:255]
                    item.normalized_name = _norm(display_name_clean)[:255]
                    if var_clean:
                        item.variation_label = var_clean[:100]
                    if medida_norm and not item.normalized_measure:
                        item.normalized_measure = medida_norm[:100]
                    item.canonical_measure_key = item.canonical_measure_key or (measure_key[:160] if measure_key else None)
                    item.measure_display = item.measure_display or (measure_display[:160] if measure_display else None)
                    item.measure_kind = item.measure_kind or (measure_kind[:40] if measure_kind else None)
                    merged_aliases = sorted(set((item.measure_aliases_json or []) + measure_aliases))
                    item.measure_aliases_json = merged_aliases
                    item.canonical_identity_key = item.canonical_identity_key or canonical_identity_key
                    item.canonical_category = canonical_key
                    item.canonical_category_display = canonical_display
                    item.visibility_scope = visibility_scope
                    item.review_type = review_type
                    item.is_operational = is_operational
                    item.is_searchable_common = is_searchable_common
                    item.is_tree_visible_common = is_tree_visible_common
                    item.is_parser_junk = is_junk
                    item.metadata_json = {
                        **(item.metadata_json or {}),
                        "source_hash": source_hash,
                        "status_importacao": status_imp,
                        "origem_preco": origem_preco,
                        "linha_origem_compras_nova": linha_origem,
                        "valor_final_eh_metadado": True,
                        "measure_aliases": merged_aliases,
                        "dedupe_sources": sorted(set((item.metadata_json or {}).get("dedupe_sources", []) + [source_hash or identity_hash])),
                        "peso": extracted_weight or (item.metadata_json or {}).get("peso")
                    }
                    db.flush()

                # Adicionar especificação primária se houver
                if especificacao:
                    spec = db.query(StockCatalogItemSpec).filter(
                        StockCatalogItemSpec.item_id == item.id,
                        StockCatalogItemSpec.spec_key == "medida"
                    ).first()
                    if not spec:
                        spec = StockCatalogItemSpec(
                            item_id=item.id,
                            spec_key="medida",
                            spec_value=especificacao,
                            normalized_value=_norm(especificacao)
                        )
                        db.add(spec)
                        db.flush()

                # Adicionar especificação de peso se houver
                if extracted_weight:
                    spec_peso = db.query(StockCatalogItemSpec).filter(
                        StockCatalogItemSpec.item_id == item.id,
                        StockCatalogItemSpec.spec_key == "peso"
                    ).first()
                    if not spec_peso:
                        spec_peso = StockCatalogItemSpec(
                            item_id=item.id,
                            spec_key="peso",
                            spec_value=extracted_weight,
                            normalized_value=_norm(extracted_weight)
                        )
                        db.add(spec_peso)
                        db.flush()

                # Adicionar Vínculo/Link se houver Cybersul e Compras Nova pareados
                if cyber_prod and status_imp in ("CYBERSUL_ENRIQUECIDO_COMPRAS", "INATIVO_CYBERSUL"):
                    link = db.query(StockCatalogLink).filter(
                        StockCatalogLink.stock_catalog_item_id == item.id,
                        StockCatalogLink.cybersul_product_id == cyber_prod.id
                    ).first()
                    if not link:
                        link = StockCatalogLink(
                            stock_catalog_item_id=item.id,
                            cybersul_product_id=cyber_prod.id,
                            match_type="EXACT" if conf_vinculo and float(conf_vinculo) >= 95 else "FUZZY",
                            confidence=float(conf_vinculo)/100.0 if conf_vinculo else 1.0,
                            match_reason=f"Vinculação mestre importada com status {status_imp}.",
                            needs_review=needs_review_bool
                        )
                        db.add(link)
                        db.flush()

                # Criar entrada na fila de revisão se precisar de revisão técnica ativa
                if needs_review_bool:
                    total_review += 1
                    rev_entry = db.query(StockCatalogReviewQueue).filter(
                        StockCatalogReviewQueue.item_id == item.id,
                        StockCatalogReviewQueue.status == "PENDING"
                    ).first()
                    if not rev_entry:
                        rev_entry = StockCatalogReviewQueue(
                            item_id=item.id,
                            import_run_id=run.id,
                            review_type=review_type if review_type != "none" else "AMBIGUOUS_VARIATION",
                            severity="HIGH" if quality_status == "QUARANTINED" else "MEDIUM",
                            title=f"Revisar item: {item.display_name}"[:255],
                            description=motivo_rev or "Revisão manual solicitada.",
                            status="PENDING"
                        )
                        db.add(rev_entry)
                        db.flush()

                # --- D. Alimentação do Índice de Busca ---
                search_text = (
                    f"{item.display_name} {item.base_name} {item.variation_label or ''} "
                    f"{item.specification_text or ''} {item.internal_code or ''} "
                    f"{forn_principal or ''} {sheet_title} "
                    f"{family_title} {' '.join(measure_aliases)}"
                )
                normalized_search_text = _norm(search_text)
                tokens = list(set([t for t in normalized_search_text.split() if len(t) > 1]))

                search_index = db.query(StockCatalogSearchIndex).filter(
                    StockCatalogSearchIndex.item_id == item.id
                ).first()

                if not search_index:
                    search_index = StockCatalogSearchIndex(
                        item_id=item.id,
                        search_text=search_text,
                        normalized_search_text=normalized_search_text,
                        tokens_json=tokens,
                        supplier_names=[forn_principal[:255]] if forn_principal else [],
                        cybersul_code=cod_cybersul[:100] if cod_cybersul else None,
                        source_sheet=sheet_title[:100],
                        family_path=f"{sheet_title} > {family_title}"[:1000],
                        last_price=preco_val,
                        last_supplier=forn_principal[:255] if forn_principal else None,
                        updated_at=datetime.now(timezone.utc)
                    )
                    db.add(search_index)
                else:
                    search_index.search_text = search_text
                    search_index.normalized_search_text = normalized_search_text
                    search_index.tokens_json = tokens
                    search_index.last_price = preco_val
                    search_index.cybersul_code = cod_cybersul[:100] if cod_cybersul else None
                    supplier_names = search_index.supplier_names or []
                    if forn_principal and forn_principal[:255] not in supplier_names:
                        supplier_names.append(forn_principal[:255])
                    search_index.supplier_names = supplier_names
                    if forn_principal:
                        search_index.last_supplier = forn_principal[:255]
                    search_index.updated_at = datetime.now(timezone.utc)
                
                db.flush()

                items_cache[identity_hash] = item
                if canonical_identity_key:
                    items_cache[canonical_identity_key] = item
                if cod_cybersul:
                    items_cache[cod_cybersul] = item
                for key_source in (
                    f"{display_name_clean}|{variacao or ''}|{sheet_title}",
                    f"{nome_limpo or display_name_clean}|{variacao or ''}|{sheet_title}",
                    f"{display_name_clean}|{sheet_title}",
                    f"{nome_limpo or display_name_clean}|{sheet_title}",
                ):
                    key_norm = _norm(key_source)
                    if key_norm:
                        items_cache[key_norm] = item

            # --- PARTE 3: IMPORTAR OFERTAS E PREÇOS ---
            print("Importando ofertas e preços...")
            ws_ofertas = wb["OFERTAS_E_PRECOS"]
            ofertas_rows = list(ws_ofertas.iter_rows(values_only=True))
            o_header = [str(x).strip().lower() if x is not None else "" for x in ofertas_rows[0]]
            col_o_map = {name: idx for idx, name in enumerate(o_header)}

            total_offers = 0
            for r_idx in range(1, len(ofertas_rows)):
                row = ofertas_rows[r_idx]
                if not row or len(row) < 5:
                    continue

                chave_sugerida = _text(row[col_o_map.get("chave_sugerida_portal", 1)])
                cod_cyber = _text(row[col_o_map.get("codigo_cybersul", 2)])
                nome_oferta = _text(row[col_o_map.get("nome_limpo_portal", 3)])
                variacao_oferta = _text(row[col_o_map.get("variacao", 4)])
                forn_name = _text(row[col_o_map.get("fornecedor", 5)])
                price_val = _money(row[col_o_map.get("preco", 6)])
                final_val = _money(row[col_o_map.get("valor_final", 7)])
                final_meta = _text(row[col_o_map.get("valor_final_eh_metadado", 8)])
                unit_val = _text(row[col_o_map.get("unidade", 10)]) or "un"
                origem = _text(row[col_o_map.get("origem", 11)])
                aba_orig = _text(row[col_o_map.get("aba_origem", 12)])
                lin_orig = row[col_o_map.get("linha_origem", 13)]
                data_ref = _text(row[col_o_map.get("data_referencia", 14)])
                cont_email = _email(row[col_o_map.get("contato_email", 15)])
                cont_phone = _phone(row[col_o_map.get("contato_telefone", 16)])
                is_curr = row[col_o_map.get("atual", 17)]
                prec_rev = _text(row[col_o_map.get("precisa_revisao", 18)])
                observacao = _text(row[col_o_map.get("observacao", 19)])

                if not forn_name:
                    continue

                norm_forn = _norm(forn_name)
                supplier = forn_map.get(norm_forn)
                if not supplier:
                    supplier = db.query(StockCatalogSupplier).filter(
                        StockCatalogSupplier.normalized_name == norm_forn
                    ).first()
                if not supplier:
                    supplier = StockCatalogSupplier(
                        name=forn_name[:255],
                        normalized_name=norm_forn[:255],
                        email=cont_email[:255] if cont_email else None,
                        phone=cont_phone[:50] if cont_phone else None,
                        source="UNIFIED_IMPORT_OFFERS",
                        active=True
                    )
                    db.add(supplier)
                    db.flush()
                    forn_map[norm_forn] = supplier
                else:
                    if cont_email and not supplier.email:
                        supplier.email = cont_email[:255]
                    if cont_phone and not supplier.phone:
                        supplier.phone = cont_phone[:50]

                item = None
                if cod_cyber:
                    item = items_cache.get(cod_cyber)
                if not item:
                    item = db.query(StockCatalogItem).filter(
                        StockCatalogItem.internal_code == cod_cyber
                    ).first() if cod_cyber else None
                if not item:
                    for key_source in (
                        f"{nome_oferta}|{variacao_oferta}|{aba_orig}",
                        f"{nome_oferta}|{aba_orig}",
                        chave_sugerida,
                    ):
                        key_norm = _norm(key_source)
                        if key_norm and key_norm in items_cache:
                            item = items_cache[key_norm]
                            break
                if not item and nome_oferta:
                    offer_measure = normalize_measure(variacao_oferta, nome_oferta)
                    offer_key = _norm(
                        "|".join(
                            [
                                get_canonical_category(aba_orig, nome_oferta) or "",
                                normalize_category_title(aba_orig) or "",
                                "",
                                normalize_category_title(nome_oferta) or nome_oferta,
                                offer_measure.canonical_key or _norm(variacao_oferta or nome_oferta),
                            ]
                        )
                    )[:500]
                    item = db.query(StockCatalogItem).filter(
                        StockCatalogItem.canonical_identity_key == offer_key
                    ).first()
                if not item and nome_oferta:
                    item = db.query(StockCatalogItem).filter(
                        StockCatalogItem.normalized_name == _norm(nome_oferta)[:255],
                        StockCatalogItem.source_sheet == aba_orig
                    ).first()
                
                if not item:
                    continue

                total_offers += 1

                # Inativar ofertas anteriores deste item do mesmo fornecedor
                db.query(StockCatalogOffer).filter(
                    StockCatalogOffer.item_id == item.id,
                    StockCatalogOffer.supplier_id == supplier.id
                ).update({"is_current": False})

                is_offer_curr = _yes(is_curr)
                
                # Ignorar ofertas com preço zerado ou ausente como ativa no catálogo
                if price_val is None or price_val <= 0:
                    is_offer_curr = False

                offer = StockCatalogOffer(
                    item_id=item.id,
                    supplier_id=supplier.id,
                    import_run_id=run.id,
                    source_sheet=aba_orig[:100] if aba_orig else "OFERTAS_E_PRECOS",
                    source_row=lin_orig if isinstance(lin_orig, int) else 0,
                    price=price_val,
                    price_raw=str(price_val) if price_val else None,
                    final_value=final_val,
                    final_value_raw=str(final_val) if final_val else None,
                    currency="BRL",
                    unit=unit_val[:30] if unit_val else "un",
                    contact_email=cont_email[:255] if cont_email else None,
                    contact_phone=cont_phone[:50] if cont_phone else None,
                    is_current=is_offer_curr,
                    is_consolidated_line=False,
                    confidence=1.0,
                    needs_review=_yes(prec_rev),
                    metadata_json={
                        "oferta_id": _text(row[col_o_map.get("oferta_id", 0)]),
                        "chave_sugerida_portal": chave_sugerida,
                        "valor_final_eh_metadado": _yes(final_meta),
                        "origem": origem,
                        "data_referencia": data_ref,
                        "observacao": observacao
                    }
                )
                db.add(offer)
                db.flush()

                search_index = db.query(StockCatalogSearchIndex).filter(
                    StockCatalogSearchIndex.item_id == item.id
                ).first()
                if search_index:
                    extra_search = " ".join(
                        x for x in [
                            search_index.search_text,
                            supplier.name,
                            cont_email or "",
                            cont_phone or "",
                            nome_oferta or "",
                            variacao_oferta or "",
                            origem or "",
                            observacao or ""
                        ] if x
                    )
                    search_index.search_text = extra_search
                    search_index.normalized_search_text = _norm(extra_search)
                    search_index.tokens_json = list(set([t for t in search_index.normalized_search_text.split() if len(t) > 1]))
                    names = search_index.supplier_names or []
                    if supplier.name not in names:
                        names.append(supplier.name)
                    search_index.supplier_names = names
                    if is_offer_curr and price_val and price_val > 0:
                        search_index.last_price = price_val
                        search_index.last_supplier = supplier.name
                    search_index.updated_at = datetime.now(timezone.utc)

            # --- PARTE 3.1: IMPORTAR TERMOS DE BUSCA E ALIASES ---
            if "TERMOS_BUSCA_ALIAS" in wb.sheetnames:
                print("Importando termos de busca e aliases...")
                ws_alias = wb["TERMOS_BUSCA_ALIAS"]
                alias_rows = list(ws_alias.iter_rows(values_only=True))
                if alias_rows:
                    a_header = [str(x).strip().lower() if x is not None else "" for x in alias_rows[0]]
                    col_a_map = {name: idx for idx, name in enumerate(a_header)}
                    for row in alias_rows[1:]:
                        if not row:
                            continue
                        cod_cyber = _text(row[col_a_map.get("codigo_cybersul", 1)])
                        nome_alias = _text(row[col_a_map.get("nome_limpo_portal", 2)])
                        termo = _text(row[col_a_map.get("termo_busca", 3)])
                        origem_alias = _text(row[col_a_map.get("origem", 5)])
                        if not termo:
                            continue
                        item = items_cache.get(cod_cyber) if cod_cyber else None
                        if not item and nome_alias:
                            item = db.query(StockCatalogItem).filter(
                                StockCatalogItem.normalized_name == _norm(nome_alias)[:255]
                            ).first()
                        if not item:
                            continue
                        search_index = db.query(StockCatalogSearchIndex).filter(
                            StockCatalogSearchIndex.item_id == item.id
                        ).first()
                        if not search_index:
                            continue
                        extra_search = f"{search_index.search_text} {termo} {origem_alias or ''}"
                        search_index.search_text = extra_search
                        search_index.normalized_search_text = _norm(extra_search)
                        search_index.tokens_json = list(set([t for t in search_index.normalized_search_text.split() if len(t) > 1]))
                        search_index.updated_at = datetime.now(timezone.utc)

            # --- PARTE 4: IMPORTAR ALERTAS ---
            print("Importando alertas e divergências...")
            ws_alertas = wb["DIVERGENCIAS_E_ALERTAS"]
            alertas_rows = list(ws_alertas.iter_rows(values_only=True))
            al_header = [str(x).strip().lower() if x is not None else "" for x in alertas_rows[0]]
            col_al_map = {name: idx for idx, name in enumerate(al_header)}

            for r_idx in range(1, len(alertas_rows)):
                row = alertas_rows[r_idx]
                if not row or len(row) < 4:
                    continue
                
                tipo_al = _text(row[col_al_map.get("tipo_alerta", 0)])
                sev_al = _text(row[col_al_map.get("severidade", 1)])
                cod_cyber = _text(row[col_al_map.get("codigo_cybersul", 2)])
                desc_al = _text(row[col_al_map.get("descricao", 3)])
                prob_det = _text(row[col_al_map.get("problema_detectado", 7)])
                
                item = None
                if cod_cyber:
                    item = items_cache.get(cod_cyber)
                if not item and cod_cyber:
                    item = db.query(StockCatalogItem).filter(
                        StockCatalogItem.internal_code == cod_cyber
                    ).first()

                if item:
                    alert = db.query(StockCatalogAlert).filter(
                        StockCatalogAlert.item_id == item.id,
                        StockCatalogAlert.alert_type == tipo_al,
                        StockCatalogAlert.status == "ACTIVE"
                    ).first()
                    
                    if not alert:
                        alert = StockCatalogAlert(
                            item_id=item.id,
                            alert_type=tipo_al,
                            severity="HIGH" if sev_al == "alta" else "MEDIUM",
                            title=desc_al[:255],
                            description=prob_det,
                            status="ACTIVE"
                        )
                        db.add(alert)
                        db.flush()

            run.status = "SUCCESS"
            run.finished_at = datetime.now(timezone.utc)
            run.total_rows = total_rows
            run.total_items = total_items
            run.total_offers = total_offers
            run.total_review = total_review
            
            db.query(StockCatalogItem).filter(
                StockCatalogItem.import_run_id != run.id,
                StockCatalogItem.active == True
            ).update({"active": False}, synchronize_session=False)
            
            db.commit()
            print("Importação mestre unificada concluída com sucesso!")

        except Exception as e:
            db.rollback()
            run.status = "FAILED"
            run.finished_at = datetime.now(timezone.utc)
            run.error_message = str(e)
            db.commit()
            print(f"Erro na importação mestre unificada: {str(e)}")
            raise e

        return run
