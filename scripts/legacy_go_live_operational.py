#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Ativacao operacional local de dados reais do Portal Vesper.

Este script roda no ambiente de desenvolvimento do backend, le fontes legadas em
modo somente leitura, cria staging real, promove dados seguros para tabelas
oficiais e grava exportacoes para o chefe. Ele nao executa apps antigos, macros,
workflows externos, nem altera arquivos legados.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import os
import re
import sqlite3
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "backend"
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))
if str(ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(ROOT / "scripts"))

import openpyxl
from sqlalchemy.orm import Session

from app.core.database import SessionLocal
from app.core.events import emit_event
from app.models.event_log import EventLog
from app.models.kanban import File, KanbanActivity, KanbanBoard, KanbanCard, KanbanColumn
from app.models.legacy_import import (
    LegacyEntityLink,
    LegacyFileIndex,
    LegacyImportBatch,
    LegacyImportRow,
    LegacyOperationalRecord,
)
from app.models.master_data import Person, ProductItem, Supplier
from app.models.notification import Notification, NotificationDelivery
from app.models.purchase import PurchasePriceHistory, PurchasePriceReference
from app.models.user import User
from app.models.it import ITAccessRecord, ITAsset, ITCredential, ITSystem
from app.modules.it.credentials import encrypt_secret
from app.modules.legacy_imports.service import LegacyImportService

from legacy_import_ti_routines import import_ti_routines
from legacy_import_projeto import import_projeto
from legacy_import_propostas import import_propostas


def configured_path(env_name: str) -> Path:
    value = os.getenv(env_name, "").strip()
    return Path(value) if value else Path("__not_configured__")


COMPRAS_XLSX = configured_path("PORTAL_LEGACY_COMPRAS_XLSX")
TI_XLSM = configured_path("PORTAL_LEGACY_TI_XLSM")
PRODUCAO_BASE = configured_path("PORTAL_LEGACY_PRODUCAO_DIR")
PROJETO_DB = configured_path("PORTAL_LEGACY_PROJETO_DB")
PROPOSTAS_BASE = configured_path("PORTAL_LEGACY_PROPOSTAS_DIR")
EXPORT_DIR = configured_path("PORTAL_LEGACY_EXPORT_DIR")

SENSITIVE_KEYS = {"password", "senha", "token", "secret", "key", "chave", "hash", "credential", "pwd"}
MANDATORY_SYSTEMS = [
    "Portal Vesper",
    "Kanban",
    "Help Desk",
    "Abacus",
    "Cybersul",
    "Skymail/E-mail",
    "NAS",
    "AnyDesk",
    "Impacta/VOIP",
    "Office",
    "ESET",
    "Fusion",
    "Compras",
    "TI",
    "Aprovações",
    "Automação IA",
    "Administração",
    "Chat",
    "Knowledge",
    "Propostas",
    "Produção",
    "Projetos",
    "Estoque",
]


@dataclass
class BatchResult:
    batch_id: str
    rows: int
    source_app: str
    module_target: str
    status: str


@dataclass
class GoLiveStats:
    batches: list[BatchResult] = field(default_factory=list)
    rows_by_source: Counter = field(default_factory=Counter)
    official_suppliers_created: int = 0
    official_suppliers_linked: int = 0
    official_items_created: int = 0
    official_items_linked: int = 0
    prices_activated: int = 0
    price_references_created: int = 0
    price_references_updated: int = 0
    access_records_created: int = 0
    access_records_linked: int = 0
    vault_secrets_imported: int = 0
    vault_secrets_updated: int = 0
    assets_created: int = 0
    assets_updated: int = 0
    production_ops: int = 0
    project_records: int = 0
    project_cards_created: int = 0
    proposal_files: int = 0
    knowledge_files_indexed: int = 0
    exports_written: list[str] = field(default_factory=list)
    events_created: int = 0
    notifications_created: int = 0
    errors: list[str] = field(default_factory=list)
    empty_sources: list[str] = field(default_factory=list)


class CaptureClient:
    def __init__(self) -> None:
        self.rows: list[dict[str, Any]] = []
        self.batch_meta: dict[str, Any] = {}

    def create_batch(self, **kwargs: Any) -> str:
        self.batch_meta = kwargs
        return "capture"

    def upload_rows(self, batch_id: str, rows: list[dict[str, Any]]) -> int:
        self.rows.extend(rows)
        return len(rows)


def utcnow() -> datetime:
    return datetime.now(timezone.utc)


def mask_path(path: str | Path) -> str:
    text = str(path)
    parts = Path(text).parts
    if len(parts) <= 2:
        return text
    return os.path.join(parts[0], "...", parts[-1])


def safe_slug(value: str, max_len: int = 80) -> str:
    text = re.sub(r"[^A-Za-z0-9]+", "-", value.strip().upper()).strip("-")
    return text[:max_len] or "REAL"


def short_hash(value: str, length: int = 12) -> str:
    return hashlib.sha1(value.encode("utf-8", errors="ignore")).hexdigest()[:length]


def parse_decimal(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value)).quantize(Decimal("0.01"))
    text = str(value).replace("R$", "").strip()
    if not text:
        return None
    if "," in text and "." in text:
        if text.find(".") < text.find(","):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return Decimal(text).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        return None


def sanitize(data: Any) -> Any:
    return LegacyImportService.sanitize_data(data)


def get_actor(db: Session) -> User:
    user = (
        db.query(User)
        .join(User.role, isouter=True)
        .filter(User.is_active == True, User.username.in_([name.strip() for name in os.getenv("PORTAL_GO_LIVE_USERS", "vesper_admin").split(",") if name.strip()]))
        .order_by(User.id)
        .first()
    )
    if not user:
        user = db.query(User).filter(User.is_active == True).order_by(User.id).first()
    if not user:
        raise RuntimeError("Nenhum usuario ativo encontrado para registrar a ativacao real.")
    return user


def create_or_update_batch(
    db: Session,
    *,
    source_app: str,
    source_name: str,
    source_path_masked: str,
    module_target: str,
    rows: list[dict[str, Any]],
    actor_user_id: int,
    notes: str,
) -> LegacyImportBatch:
    batch = (
        db.query(LegacyImportBatch)
        .filter(LegacyImportBatch.source_app == source_app, LegacyImportBatch.source_name == source_name)
        .first()
    )
    if not batch:
        batch = LegacyImportBatch(
            source_app=source_app,
            source_name=source_name,
            source_path_masked=source_path_masked,
            module_target=module_target,
            status="DISCOVERED",
            notes=notes,
            created_by_user_id=actor_user_id,
        )
        db.add(batch)
        db.flush()
        emit_event(
            db,
            "legacy.import.batch.created",
            "LegacyImportBatch",
            str(batch.id),
            "legacy_imports",
            {
                "batch_id": str(batch.id),
                "source_app": source_app,
                "source_name": source_name,
                "module_target": module_target,
                "actor_user_id": actor_user_id,
                "summary": f"Lote real {source_name} criado para {module_target}.",
            },
            actor_user_id=actor_user_id,
        )

    existing_keys = {
        (row.source_table_or_sheet, row.source_row_id, row.entity_target)
        for row in db.query(LegacyImportRow).filter(LegacyImportRow.batch_id == batch.id).all()
    }
    inserted = 0
    for row in rows:
        key = (row.get("source_table_or_sheet"), str(row.get("source_row_id")), row["entity_target"])
        if key in existing_keys:
            continue
        raw = sanitize(row.get("raw_data_json") or {})
        normalized = sanitize(row.get("normalized_data_json") or raw)
        db.add(
            LegacyImportRow(
                batch_id=batch.id,
                source_app=source_app,
                source_table_or_sheet=key[0],
                source_row_id=key[1],
                module_target=module_target,
                entity_target=row["entity_target"],
                raw_data_json=raw,
                normalized_data_json=normalized,
                confidence_score=row.get("confidence_score") or 1.0,
                status=row.get("status") or "PENDING_REVIEW",
            )
        )
        existing_keys.add(key)
        inserted += 1

    batch.total_rows = db.query(LegacyImportRow).filter(LegacyImportRow.batch_id == batch.id).count()
    batch.valid_rows = batch.total_rows
    batch.duplicate_rows = 0
    batch.error_rows = 0
    batch.status = "EMPTY" if batch.total_rows == 0 else "EXTRACTED"
    batch.finished_at = utcnow()
    db.flush()
    emit_event(
        db,
        "legacy.import.batch.completed",
        "LegacyImportBatch",
        str(batch.id),
        "legacy_imports",
        {
            "batch_id": str(batch.id),
            "source_app": source_app,
            "source_name": source_name,
            "module_target": module_target,
            "total_rows": batch.total_rows,
            "valid_rows": batch.valid_rows,
            "duplicate_rows": 0,
            "error_rows": 0,
            "actor_user_id": actor_user_id,
            "summary": f"Lote real {source_name} extraido com {batch.total_rows} linhas.",
        },
        actor_user_id=actor_user_id,
    )
    if inserted:
        db.flush()
    return batch


def link_entity(
    db: Session,
    *,
    row: LegacyImportRow | None,
    source_app: str,
    entity_target: str,
    official_entity_type: str,
    official_entity_id: str,
    action: str,
    actor_user_id: int,
) -> bool:
    exists = (
        db.query(LegacyEntityLink)
        .filter(
            LegacyEntityLink.source_app == source_app,
            LegacyEntityLink.entity_target == entity_target,
            LegacyEntityLink.official_entity_type == official_entity_type,
            LegacyEntityLink.official_entity_id == str(official_entity_id),
        )
        .first()
    )
    if exists:
        return False
    db.add(
        LegacyEntityLink(
            legacy_row_id=row.id if row else None,
            source_app=source_app,
            entity_target=entity_target,
            official_entity_type=official_entity_type,
            official_entity_id=str(official_entity_id),
            action=action,
            created_by_user_id=actor_user_id,
        )
    )
    return True


def extract_compras_xlsx() -> list[dict[str, Any]]:
    if not COMPRAS_XLSX.exists():
        return []
    wb = openpyxl.load_workbook(COMPRAS_XLSX, read_only=True, data_only=True)
    rows: list[dict[str, Any]] = []
    for ws in wb.worksheets:
        if "indice" in ws.title.lower() or "índice" in ws.title.lower():
            continue
        header = None
        header_row_num = 0
        for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(8, ws.max_row), values_only=True), start=1):
            normalized = [str(c or "").strip().lower() for c in row]
            if any("código" in c or "codigo" in c for c in normalized) and any("preço" in c or "preco" in c for c in normalized):
                header = normalized
                header_row_num = idx
                break
        if not header:
            continue

        def col(*terms: str) -> int | None:
            for i, value in enumerate(header):
                if any(term in value for term in terms):
                    return i
            return None

        col_code = col("código", "codigo")
        col_desc = col("material", "descrição", "descricao")
        col_price = col("valor final", "preço", "preco")
        col_supplier = col("fornecedor")
        col_date = col("atualizado")
        col_company = col("empresa")
        col_unit = col("unidade", "un")

        for row_num, values in enumerate(ws.iter_rows(min_row=header_row_num + 1, values_only=True), start=header_row_num + 1):
            code = str(values[col_code]).strip() if col_code is not None and col_code < len(values) and values[col_code] else ""
            desc = str(values[col_desc]).strip() if col_desc is not None and col_desc < len(values) and values[col_desc] else ""
            supplier = str(values[col_supplier]).strip() if col_supplier is not None and col_supplier < len(values) and values[col_supplier] else ""
            company = str(values[col_company]).strip() if col_company is not None and col_company < len(values) and values[col_company] else ""
            observed = str(values[col_date]).strip() if col_date is not None and col_date < len(values) and values[col_date] else ""
            unit = str(values[col_unit]).strip() if col_unit is not None and col_unit < len(values) and values[col_unit] else "un"
            price = parse_decimal(values[col_price] if col_price is not None and col_price < len(values) else None)
            if not desc or price is None or price <= 0:
                continue
            sku = code or f"IMP-{safe_slug(ws.title, 20)}-{row_num:04d}-{short_hash(desc, 6)}"
            raw = {
                "sheet": ws.title,
                "row_number": row_num,
                "code": code,
                "description": desc,
                "unit_price": float(price),
                "supplier_name": supplier or company or "Fornecedor nao identificado",
                "company": company,
                "updated_at": observed,
                "unit_of_measure": unit if unit and len(unit) <= 30 else "un",
                "source_file": mask_path(COMPRAS_XLSX),
            }
            rows.append(
                {
                    "source_table_or_sheet": ws.title,
                    "source_row_id": f"{ws.title}:{row_num}:{sku}",
                    "entity_target": "PRICE_HISTORY",
                    "raw_data_json": raw,
                    "normalized_data_json": raw,
                    "confidence_score": 0.96,
                    "status": "PENDING_REVIEW",
                }
            )
    wb.close()
    return rows


def get_or_create_supplier(db: Session, name: str, actor_user_id: int) -> tuple[Supplier, bool]:
    clean = (name or "Fornecedor nao identificado").strip()[:255]
    existing = (
        db.query(Supplier)
        .join(Person)
        .filter(Person.name == clean)
        .first()
    )
    if existing:
        return existing, False
    person = Person(type="COMPANY", name=clean, notes="Criado por ativacao real de Compras Nova.")
    db.add(person)
    db.flush()
    supplier = Supplier(
        person_id=person.id,
        supplier_code=f"SUP-{short_hash(clean).upper()}",
        categories=["LEGACY_PURCHASES"],
        status="ACTIVE",
        notes="Fornecedor real ativado a partir de fonte legada.",
    )
    db.add(supplier)
    db.flush()
    return supplier, True


def get_or_create_product(db: Session, sku: str, name: str, category: str) -> tuple[ProductItem, bool]:
    clean_sku = safe_slug(sku, 90)
    product = db.query(ProductItem).filter(ProductItem.sku == clean_sku).first()
    if product:
        return product, False
    product = ProductItem(
        sku=clean_sku,
        name=name[:255],
        description=f"Variação real importada da aba {category}.",
        item_type="RAW_MATERIAL",
        unit_of_measure="un",
        category=category[:100],
        is_active=True,
    )
    db.add(product)
    db.flush()
    return product, True


def promote_compras(db: Session, batch: LegacyImportBatch, actor: User, stats: GoLiveStats) -> None:
    rows = db.query(LegacyImportRow).filter(LegacyImportRow.batch_id == batch.id, LegacyImportRow.entity_target == "PRICE_HISTORY").all()
    for row in rows:
        data = row.normalized_data_json or {}
        price = parse_decimal(data.get("unit_price"))
        if price is None or price <= 0:
            continue
        supplier, supplier_created = get_or_create_supplier(db, data.get("supplier_name") or data.get("company"), actor.id)
        stats.official_suppliers_created += int(supplier_created)
        stats.official_suppliers_linked += int(not supplier_created)
        sku = data.get("code") or row.source_row_id or data.get("description")
        product, item_created = get_or_create_product(db, sku, data.get("description") or sku, data.get("sheet") or "Compras")
        stats.official_items_created += int(item_created)
        stats.official_items_linked += int(not item_created)

        source_id = f"go-live:{short_hash(str(row.id), 16)}"
        history = db.query(PurchasePriceHistory).filter(
            PurchasePriceHistory.source_type == "IMPORTED_XLSX",
            PurchasePriceHistory.source_id == source_id,
        ).first()
        if not history:
            history = PurchasePriceHistory(
                product_item_id=product.id,
                supplier_id=supplier.id,
                unit_price=price,
                currency="BRL",
                unit_of_measure=data.get("unit_of_measure") or product.unit_of_measure,
                observed_at=utcnow(),
                source_type="IMPORTED_XLSX",
                source_id=source_id,
                created_by_user_id=actor.id,
            )
            db.add(history)
            db.flush()
            stats.prices_activated += 1
            row.status = "IMPORTED"
            link_entity(db, row=row, source_app=batch.source_app, entity_target="PRICE_HISTORY", official_entity_type="PurchasePriceHistory", official_entity_id=str(history.id), action="CREATED", actor_user_id=actor.id)

        reference = db.query(PurchasePriceReference).filter(
            PurchasePriceReference.product_item_id == product.id,
            PurchasePriceReference.supplier_id == supplier.id,
            PurchasePriceReference.is_active == True,
        ).first()
        if not reference:
            db.add(
                PurchasePriceReference(
                    product_item_id=product.id,
                    supplier_id=supplier.id,
                    current_unit_price=price,
                    currency="BRL",
                    unit_of_measure=data.get("unit_of_measure") or product.unit_of_measure,
                    source_history_id=history.id,
                    approved_by_user_id=actor.id,
                    notes="Referencia real ativada automaticamente a partir de Compras Nova.xlsx.",
                    is_active=True,
                )
            )
            stats.price_references_created += 1
        elif Decimal(str(reference.current_unit_price)) != price:
            reference.current_unit_price = price
            reference.source_history_id = history.id
            reference.approved_by_user_id = actor.id
            reference.approved_at = utcnow()
            reference.notes = "Referencia real atualizada pela ativacao operacional."
            stats.price_references_updated += 1
        link_entity(db, row=row, source_app=batch.source_app, entity_target="PRODUCT_ITEM", official_entity_type="ProductItem", official_entity_id=str(product.id), action="CREATED" if item_created else "LINKED_EXISTING", actor_user_id=actor.id)
        link_entity(db, row=row, source_app=batch.source_app, entity_target="SUPPLIER", official_entity_type="Supplier", official_entity_id=str(supplier.id), action="CREATED" if supplier_created else "LINKED_EXISTING", actor_user_id=actor.id)


def capture_import(import_func, *args: Any) -> list[dict[str, Any]]:
    client = CaptureClient()
    import_func(client, False, *args) if args else import_func(client, False)
    return client.rows


def ensure_system(db: Session, name: str) -> ITSystem:
    system = db.query(ITSystem).filter(ITSystem.name == name).first()
    if not system:
        system = ITSystem(name=name, description=f"Sistema real ativado: {name}", is_active=True)
        db.add(system)
        db.flush()
    return system


def find_user_by_name(db: Session, name: str | None) -> User | None:
    if not name:
        return None
    needle = name.strip().lower()
    if not needle:
        return None
    return (
        db.query(User)
        .filter(
            (User.username.ilike(f"%{needle}%"))
            | (User.email.ilike(f"{needle}%"))
        )
        .first()
    )


def upsert_access(
    db: Session,
    *,
    system_name: str,
    legacy_user_name: str,
    access_profile: str,
    has_secret: bool,
    origin: str,
) -> tuple[ITAccessRecord, bool]:
    system = ensure_system(db, system_name)
    access = db.query(ITAccessRecord).filter(
        ITAccessRecord.system_name == system_name,
        ITAccessRecord.legacy_user_name == legacy_user_name,
        ITAccessRecord.access_profile == access_profile,
    ).first()
    created = False
    if not access:
        access = ITAccessRecord(
            system_id=system.id,
            system_name=system_name,
            legacy_user_name=legacy_user_name[:150],
            access_profile=access_profile[:100],
            status="ACTIVE",
            origin=origin,
            has_secret=has_secret,
            last_updated_at=utcnow(),
        )
        db.add(access)
        db.flush()
        created = True
    else:
        access.system_id = system.id
        access.status = "ACTIVE"
        access.has_secret = access.has_secret or has_secret
        access.last_updated_at = utcnow()
    return access, created


def import_secret(
    db: Session,
    *,
    title: str,
    system_name: str,
    username: str | None,
    secret: str | None,
    actor: User,
    notes: str,
) -> tuple[ITCredential | None, bool, bool]:
    if not secret or not str(secret).strip() or str(secret).strip() == "******":
        return None, False, False
    clean_secret = str(secret).strip()
    credential = db.query(ITCredential).filter(
        ITCredential.title == title,
        ITCredential.system_name == system_name,
        ITCredential.username == username,
        ITCredential.is_active == True,
    ).first()
    encrypted = encrypt_secret(clean_secret)
    if credential:
        credential.secret_encrypted = encrypted
        credential.updated_by_user_id = actor.id
        credential.updated_at = utcnow()
        return credential, False, True
    credential = ITCredential(
        title=title[:255],
        system_name=system_name,
        username=username,
        secret_encrypted=encrypted,
        secret_hint=f"{len(clean_secret)} caracteres",
        notes=notes,
        visibility_level="IT_MANAGER",
        created_by_user_id=actor.id,
        is_active=True,
    )
    db.add(credential)
    db.flush()
    return credential, True, False


def promote_ti(db: Session, batch: LegacyImportBatch, rows_raw: list[dict[str, Any]], actor: User, stats: GoLiveStats) -> None:
    for system in MANDATORY_SYSTEMS:
        ensure_system(db, system)

    staging_by_source = {
        (row.source_table_or_sheet, row.source_row_id, row.entity_target): row
        for row in db.query(LegacyImportRow).filter(LegacyImportRow.batch_id == batch.id).all()
    }

    for raw_row in rows_raw:
        raw = raw_row.get("raw_data_json") or {}
        target = raw_row.get("entity_target")
        staging_row = staging_by_source.get((raw_row.get("source_table_or_sheet"), str(raw_row.get("source_row_id")), target))
        if target == "IT_ASSET":
            asset_code = raw.get("asset_code") or raw_row.get("source_row_id")
            hostname = raw.get("hostname")
            specs = raw.get("specifications") or {}
            asset = db.query(ITAsset).filter(ITAsset.asset_tag == str(asset_code)).first()
            created = False
            if not asset:
                asset = ITAsset(asset_tag=str(asset_code), name=str(hostname or asset_code or "Ativo TI"), asset_type="COMPUTADOR", status="EM_USO")
                db.add(asset)
                db.flush()
                created = True
            asset.hostname = hostname or asset.hostname
            asset.ip_address = raw.get("ip") or raw.get("ip_address") or asset.ip_address
            asset.ram = specs.get("Memoria Ram") or specs.get("RAM") or asset.ram
            asset.processor = specs.get("Processador") or specs.get("CPU") or asset.processor
            asset.motherboard = specs.get("Placa Mãe") or asset.motherboard
            asset.gpu = specs.get("Placa de Video") or asset.gpu
            asset.network_card = specs.get("Place de Rede") or specs.get("Placa de Rede") or asset.network_card
            asset.monitor = specs.get("Monitor") or asset.monitor
            asset.mouse = specs.get("Mouse") or asset.mouse
            asset.keyboard = specs.get("Teclado") or asset.keyboard
            asset.storage = specs.get("HDDs") or specs.get("Storage") or asset.storage
            asset.custom_fields = {**(asset.custom_fields or {}), "legacy_raw": sanitize(raw)}
            asset.last_collection_source = "LEGACY_XLSM"
            asset.last_collection_at = utcnow()
            stats.assets_created += int(created)
            stats.assets_updated += int(not created)
            if staging_row:
                staging_row.status = "IMPORTED"
                link_entity(db, row=staging_row, source_app=batch.source_app, entity_target="IT_ASSET", official_entity_type="ITAsset", official_entity_id=str(asset.id), action="CREATED" if created else "UPDATED_EXISTING", actor_user_id=actor.id)

            # A ficha individual tambem carrega segredos de acesso.
            secret_map = [
                ("AnyDesk", raw.get("user") or raw.get("assigned_user") or str(asset_code), raw.get("anydesk_id"), raw.get("anydesk_password")),
                ("Cybersul", raw.get("user") or raw.get("assigned_user") or str(asset_code), raw.get("cybersul_login"), raw.get("cybersul_password")),
                ("NAS", raw.get("user") or raw.get("assigned_user") or str(asset_code), raw.get("nas_login"), raw.get("nas_password")),
            ]
            for system_name, label, username, secret in secret_map:
                if username or secret:
                    access, created_access = upsert_access(db, system_name=system_name, legacy_user_name=str(label), access_profile=str(username or "USER"), has_secret=bool(secret), origin="LEGACY_XLSM")
                    stats.access_records_created += int(created_access)
                    stats.access_records_linked += int(not created_access)
                    cred, created_cred, updated_cred = import_secret(db, title=f"{system_name} - {label}", system_name=system_name, username=str(username or label), secret=secret, actor=actor, notes="Importado da ficha de TI real.")
                    stats.vault_secrets_imported += int(created_cred)
                    stats.vault_secrets_updated += int(updated_cred)

        elif target in {"IT_REMOTE_ACCESS", "CREDENTIAL_METADATA", "NAS_ACCESS", "EMAIL_ACCOUNT_ACCESS", "VOIP_ACCOUNT"}:
            if target == "IT_REMOTE_ACCESS":
                system_name, user_label, username, secret = "AnyDesk", raw.get("user_name") or raw.get("alias"), raw.get("anydesk_id") or raw.get("alias"), raw.get("password")
            elif target == "CREDENTIAL_METADATA":
                system_name, user_label, username, secret = raw.get("system") or "Cybersul", raw.get("user_name"), raw.get("login"), raw.get("password")
            elif target == "NAS_ACCESS":
                system_name, user_label, username, secret = "NAS", raw.get("user_name"), raw.get("login"), raw.get("password")
            elif target == "EMAIL_ACCOUNT_ACCESS":
                system_name, user_label, username, secret = "Skymail/E-mail", raw.get("email"), raw.get("email"), raw.get("password")
            else:
                system_name, user_label, username, secret = "Impacta/VOIP", raw.get("user_name"), raw.get("extension"), raw.get("password")
            access, created_access = upsert_access(db, system_name=system_name, legacy_user_name=str(user_label or username or "Nao identificado"), access_profile=str(username or "USER"), has_secret=bool(secret), origin="LEGACY_XLSM")
            stats.access_records_created += int(created_access)
            stats.access_records_linked += int(not created_access)
            cred, created_cred, updated_cred = import_secret(db, title=f"{system_name} - {user_label or username}", system_name=system_name, username=str(username or user_label or ""), secret=secret, actor=actor, notes=f"Importado da planilha TI, origem {target}.")
            stats.vault_secrets_imported += int(created_cred)
            stats.vault_secrets_updated += int(updated_cred)
            if staging_row:
                staging_row.status = "IMPORTED"
                link_entity(db, row=staging_row, source_app=batch.source_app, entity_target=target, official_entity_type="ITAccessRecord", official_entity_id=str(access.id), action="CREATED" if created_access else "LINKED_EXISTING", actor_user_id=actor.id)


def locate_producao_db() -> Path | None:
    if not PRODUCAO_BASE.exists():
        return None
    for path in PRODUCAO_BASE.rglob("*.db"):
        try:
            con = sqlite3.connect(str(path))
            tables = {r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type='table'").fetchall()}
            con.close()
            if "ops" in tables:
                return path
        except sqlite3.Error:
            continue
    return None


def extract_producao_rows() -> list[dict[str, Any]]:
    db_path = locate_producao_db()
    if not db_path:
        return []
    con = sqlite3.connect(str(db_path))
    con.row_factory = sqlite3.Row
    rows_db = con.execute("SELECT * FROM ops").fetchall()
    con.close()
    rows: list[dict[str, Any]] = []
    for r in rows_db:
        data = dict(r)
        title = str(data.get("numero_op") or data.get("id") or "OP")
        normalized = {
            "numero_op": data.get("numero_op") or data.get("id"),
            "title": f"OP {title}",
            "status": data.get("status_geral") or data.get("status") or "Importado",
            "responsible": data.get("responsavel") or data.get("responsible") or "",
            "record_date": data.get("data_inicio") or data.get("created_at"),
            "data": sanitize(data),
        }
        rows.append({"source_table_or_sheet": "ops", "source_row_id": title, "entity_target": "PRODUCTION_OP", "raw_data_json": data, "normalized_data_json": normalized, "confidence_score": 1.0, "status": "PENDING_REVIEW"})
    return rows


def ensure_board(db: Session, *, name: str, slug: str, actor: User) -> KanbanBoard:
    board = db.query(KanbanBoard).filter(KanbanBoard.slug == slug).first()
    if not board:
        board = KanbanBoard(name=name, slug=slug, description="Dados reais ativados de app legado.", module_origin="legacy_go_live", created_by_user_id=actor.id)
        db.add(board)
        db.flush()
    return board


def ensure_column(db: Session, board: KanbanBoard, name: str, position: int) -> KanbanColumn:
    column = db.query(KanbanColumn).filter(KanbanColumn.board_id == board.id, KanbanColumn.name == name).first()
    if not column:
        column = KanbanColumn(board_id=board.id, name=name, position=position, is_done_column=name.lower() in {"feito", "concluido", "concluído", "finalizado"})
        db.add(column)
        db.flush()
    return column


def promote_operational_rows(db: Session, batch: LegacyImportBatch, actor: User, stats: GoLiveStats, *, board_name: str, board_slug: str, module: str) -> None:
    board = ensure_board(db, name=board_name, slug=board_slug, actor=actor)
    columns: dict[str, KanbanColumn] = {}
    rows = db.query(LegacyImportRow).filter(LegacyImportRow.batch_id == batch.id).all()
    for row in rows:
        data = row.normalized_data_json or {}
        title = data.get("title") or data.get("numero_op") or data.get("status") or row.source_row_id or "Registro legado"
        status = data.get("status") or data.get("status_geral") or "Importado"
        if status not in columns:
            columns[status] = ensure_column(db, board, status, len(columns) + 1)
        record = db.query(LegacyOperationalRecord).filter(LegacyOperationalRecord.legacy_row_id == row.id).first()
        if not record:
            record = LegacyOperationalRecord(
                legacy_row_id=row.id,
                source_app=batch.source_app,
                entity_target=row.entity_target,
                title=str(title)[:200],
                status=str(status)[:50],
                responsible=str(data.get("responsible") or "")[:150],
                data_json=data,
            )
            db.add(record)
            db.flush()
        card = None
        for existing_card in db.query(KanbanCard).filter(KanbanCard.board_id == board.id).all():
            if (existing_card.custom_fields or {}).get("legacy_row_id") == str(row.id):
                card = existing_card
                break
        if not card:
            card = KanbanCard(
                board_id=board.id,
                column_id=columns[status].id,
                title=str(title),
                description=f"Registro real importado de {batch.source_app}.",
                position=db.query(KanbanCard).filter(KanbanCard.column_id == columns[status].id).count() + 1,
                priority="MEDIUM",
                status=str(status),
                created_by_user_id=actor.id,
                custom_fields={"legacy_row_id": str(row.id), "source_app": batch.source_app, "module": module},
            )
            db.add(card)
            db.flush()
            stats.project_cards_created += 1
            db.add(KanbanActivity(board_id=board.id, card_id=card.id, actor_user_id=actor.id, action="legacy.real_card.created", metadata_json={"legacy_row_id": str(row.id)}))
        row.status = "IMPORTED"
        link_entity(db, row=row, source_app=batch.source_app, entity_target=row.entity_target, official_entity_type="LegacyOperationalRecord", official_entity_id=str(record.id), action="CREATED", actor_user_id=actor.id)
        link_entity(db, row=row, source_app=batch.source_app, entity_target=row.entity_target, official_entity_type="KanbanCard", official_entity_id=str(card.id), action="CREATED", actor_user_id=actor.id)
        if module == "production":
            stats.production_ops += 1
        else:
            stats.project_records += 1


def promote_files(db: Session, batch: LegacyImportBatch, actor: User, stats: GoLiveStats) -> None:
    rows = db.query(LegacyImportRow).filter(LegacyImportRow.batch_id == batch.id).all()
    for row in rows:
        data = row.normalized_data_json or {}
        file_name = data.get("file_name") or row.source_row_id or "arquivo"
        existing = db.query(LegacyFileIndex).filter(LegacyFileIndex.file_name == file_name, LegacyFileIndex.file_path_masked == data.get("file_path_masked")).first()
        if not existing:
            existing = LegacyFileIndex(
                legacy_row_id=row.id,
                file_name=file_name[:255],
                file_path_masked=(data.get("file_path_masked") or "")[:500],
                file_type=(data.get("file_type") or Path(file_name).suffix.replace(".", "").upper() or "FILE")[:30],
                file_size_bytes=data.get("file_size_bytes"),
                category=data.get("category") or "DOCUMENT",
                suggested_module=data.get("suggested_module") or batch.module_target,
                tags=data.get("tags") or ["legado"],
                metadata_json=data,
            )
            db.add(existing)
            db.flush()
            stats.knowledge_files_indexed += 1
        # Registro tambem na tabela files como referencia externa legada, sem copiar bytes.
        storage_key = data.get("file_path_masked") or file_name
        file_row = db.query(File).filter(File.storage_provider == "legacy_reference", File.storage_key == storage_key).first()
        if not file_row:
            file_row = File(
                original_filename=file_name,
                stored_filename=file_name,
                content_type=f"application/{(data.get('file_type') or 'octet-stream').lower()}",
                size_bytes=data.get("file_size_bytes") or 0,
                storage_provider="legacy_reference",
                storage_bucket=batch.source_app.lower(),
                storage_key=storage_key,
                checksum_sha256=str(data.get("file_hash") or short_hash(storage_key)),
                uploaded_by_user_id=actor.id,
            )
            db.add(file_row)
            db.flush()
        row.status = "IMPORTED"
        link_entity(db, row=row, source_app=batch.source_app, entity_target=row.entity_target, official_entity_type="LegacyFileIndex", official_entity_id=str(existing.id), action="CREATED", actor_user_id=actor.id)
        stats.proposal_files += int(batch.source_app == "PROPOSTASAPP")


def sync_portal_access(db: Session, actor: User, stats: GoLiveStats) -> None:
    module_to_system = {
        "dashboard": "Portal Vesper",
        "kanban": "Kanban",
        "it": "TI",
        "purchases": "Compras",
        "approvals": "Aprovações",
        "admin": "Administração",
        "chat": "Chat",
        "files": "Knowledge",
        "proposals": "Propostas",
        "automations": "Automação IA",
        "stock": "Estoque",
    }
    for system in MANDATORY_SYSTEMS:
        ensure_system(db, system)
    for user in db.query(User).filter(User.is_active == True).all():
        access, created = upsert_access(db, system_name="Portal Vesper", legacy_user_name=user.username, access_profile=user.role.name if user.role else "USER", has_secret=False, origin="PORTAL")
        access.user_id = user.id
        stats.access_records_created += int(created)
        stats.access_records_linked += int(not created)
        for ma in user.module_accesses:
            system_name = module_to_system.get(ma.module.code, ma.module.name)
            access, created = upsert_access(db, system_name=system_name, legacy_user_name=user.username, access_profile=ma.permission_level, has_secret=False, origin="PORTAL")
            access.user_id = user.id
            stats.access_records_created += int(created)
            stats.access_records_linked += int(not created)


def write_exports(db: Session, stats: GoLiveStats) -> None:
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    today = datetime.now().strftime("%Y%m%d-%H%M%S")
    access_path = EXPORT_DIR / f"matriz_acessos_real_{today}.csv"
    prices_path = EXPORT_DIR / f"compras_precos_reais_{today}.csv"
    activated_path = EXPORT_DIR / f"dados_ativados_real_{today}.csv"

    with access_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(["Nome", "E-mail", "PC", "Sistema", "Perfil", "Status", "Origem", "Tem credencial", "Atualizado em"])
        for r in db.query(ITAccessRecord).order_by(ITAccessRecord.system_name, ITAccessRecord.legacy_user_name).all():
            writer.writerow([
                r.user.username if r.user else r.legacy_user_name,
                r.user.email if r.user else "",
                "",
                r.system_name,
                r.access_profile,
                r.status,
                r.origin,
                "Sim" if r.has_secret else "Nao",
                r.last_updated_at.isoformat() if r.last_updated_at else "",
            ])

    with prices_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(["SKU", "Item/Variacao", "Preco atual", "Ultimo preco", "Fornecedor", "Data", "Origem"])
        refs = db.query(PurchasePriceReference).filter(PurchasePriceReference.is_active == True).all()
        for ref in refs:
            writer.writerow([
                ref.product_item.sku,
                ref.product_item.name,
                float(ref.current_unit_price),
                float(ref.source_history.unit_price) if ref.source_history else "",
                ref.supplier.person.name if ref.supplier and ref.supplier.person else "",
                ref.approved_at.isoformat() if ref.approved_at else "",
                "Compras Nova.xlsx",
            ])

    with activated_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(["Metrica", "Valor"])
        rows_total = db.query(LegacyImportRow).count()
        batches_total = db.query(LegacyImportBatch).count()
        suppliers_total = db.query(Supplier).count()
        items_total = db.query(ProductItem).count()
        prices_total = db.query(PurchasePriceHistory).count()
        references_total = db.query(PurchasePriceReference).filter(PurchasePriceReference.is_active == True).count()
        access_total = db.query(ITAccessRecord).count()
        secrets_total = db.query(ITCredential).filter(ITCredential.is_active == True).count()
        assets_total = db.query(ITAsset).count()
        ops_total = db.query(LegacyImportRow).filter(LegacyImportRow.entity_target == "PRODUCTION_OP").count()
        projects_total = db.query(LegacyOperationalRecord).filter(LegacyOperationalRecord.entity_target == "PROJECT_TASK").count()
        files_total = db.query(LegacyFileIndex).count()
        for key, value in [
            ("batches", batches_total),
            ("linhas_importadas", rows_total),
            ("fornecedores_oficiais", suppliers_total),
            ("itens_variacoes_oficiais", items_total),
            ("historicos_preco_ativados", prices_total),
            ("referencias_preco_ativas", references_total),
            ("acessos_reais", access_total),
            ("segredos_ativos_no_cofre", secrets_total),
            ("ativos_ti", assets_total),
            ("ops_producao", ops_total),
            ("projetos", projects_total),
            ("arquivos_indexados", files_total),
        ]:
            writer.writerow([key, value])

    stats.exports_written.extend([str(access_path), str(prices_path), str(activated_path)])


def create_summary_notifications(db: Session, actor: User, stats: GoLiveStats) -> None:
    admins = [
        u for u in db.query(User).filter(User.is_active == True).all()
        if u.role and u.role.name in {"ADMIN", "MESSIAS"}
    ]
    for user in admins:
        notification = Notification(
            user_id=user.id,
            module="legacy_promotion",
            event_type="real_data.activation.completed",
            title="Dados reais ativados",
            message=f"Ativacao operacional concluiu {sum(b.rows for b in stats.batches)} linhas reais e {stats.vault_secrets_imported} segredos no Cofre.",
            severity="SUCCESS",
            source_type="real_data_activation",
            source_id="local-go-live",
            payload_json={"contains_secrets": False},
        )
        db.add(notification)
        db.flush()
        db.add(NotificationDelivery(notification_id=notification.id, user_id=user.id))
        stats.notifications_created += 1


def main() -> None:
    parser = argparse.ArgumentParser(description="Ativa dados reais legados no Portal Vesper local.")
    parser.add_argument("--execute", action="store_true", help="Executa escrita no banco local.")
    parser.add_argument("--limit-compras", type=int, default=0, help="Limita linhas de Compras para teste. 0 = todas.")
    args = parser.parse_args()
    if not args.execute:
        print("Use --execute para ativar dados reais no banco local. Sem esta flag nada e gravado.")
        return

    stats = GoLiveStats()
    with SessionLocal() as db:
        actor = get_actor(db)
        before_events = db.query(EventLog).count()
        before_notifications = db.query(Notification).count()
        try:
            emit_event(db, "real_data.activation.started", "RealDataActivation", "local-go-live", "legacy_promotion", {"summary": "Ativacao real iniciada.", "actor_user_id": actor.id}, actor_user_id=actor.id)

            compras_rows = extract_compras_xlsx()
            if args.limit_compras:
                compras_rows = compras_rows[: args.limit_compras]
            compras_batch = create_or_update_batch(db, source_app="COMPRAS_NOVA_XLSX", source_name="Compras Nova.xlsx", source_path_masked=mask_path(COMPRAS_XLSX), module_target="purchases", rows=compras_rows, actor_user_id=actor.id, notes="Go-live operacional real de compras e precos.")
            stats.batches.append(BatchResult(str(compras_batch.id), compras_batch.total_rows, compras_batch.source_app, compras_batch.module_target, compras_batch.status))
            stats.rows_by_source["compras"] = compras_batch.total_rows
            if compras_batch.total_rows == 0:
                stats.empty_sources.append("Compras Nova.xlsx")
            promote_compras(db, compras_batch, actor, stats)
            db.commit()

            ti_client = CaptureClient()
            import_ti_routines(ti_client, False, str(TI_XLSM))
            ti_batch = create_or_update_batch(db, source_app="TI_ROUTINES_XLSM", source_name="RQ-CRT Controle de Rotinas de TI", source_path_masked=mask_path(TI_XLSM), module_target="it", rows=ti_client.rows, actor_user_id=actor.id, notes="Go-live operacional real da planilha de TI.")
            stats.batches.append(BatchResult(str(ti_batch.id), ti_batch.total_rows, ti_batch.source_app, ti_batch.module_target, ti_batch.status))
            stats.rows_by_source["ti_routines"] = ti_batch.total_rows
            promote_ti(db, ti_batch, ti_client.rows, actor, stats)
            sync_portal_access(db, actor, stats)
            db.commit()

            producao_rows = extract_producao_rows()
            producao_batch = create_or_update_batch(db, source_app="PRODUCAOAPP", source_name="Produção App SQLite", source_path_masked=mask_path(locate_producao_db() or PRODUCAO_BASE), module_target="production", rows=producao_rows, actor_user_id=actor.id, notes="Go-live operacional real de OPs de producao.")
            stats.batches.append(BatchResult(str(producao_batch.id), producao_batch.total_rows, producao_batch.source_app, producao_batch.module_target, producao_batch.status))
            if producao_batch.total_rows == 0:
                stats.empty_sources.append("ProduçãoApp")
            promote_operational_rows(db, producao_batch, actor, stats, board_name="Produção Real Legada", board_slug="producao-real-legada", module="production")
            db.commit()

            projeto_client = CaptureClient()
            import_projeto(projeto_client, False)
            projeto_batch = create_or_update_batch(db, source_app="PROJETOAPP", source_name="Projeto App SQLite", source_path_masked=mask_path(PROJETO_DB), module_target="projects", rows=projeto_client.rows, actor_user_id=actor.id, notes="Go-live operacional real de projetos.")
            stats.batches.append(BatchResult(str(projeto_batch.id), projeto_batch.total_rows, projeto_batch.source_app, projeto_batch.module_target, projeto_batch.status))
            promote_operational_rows(db, projeto_batch, actor, stats, board_name="Projetos Reais Legados", board_slug="projetos-reais-legados", module="projects")
            db.commit()

            propostas_client = CaptureClient()
            import_propostas(propostas_client, False)
            propostas_batch = create_or_update_batch(db, source_app="PROPOSTASAPP", source_name="Propostas App arquivos", source_path_masked=mask_path(PROPOSTAS_BASE), module_target="proposals", rows=propostas_client.rows, actor_user_id=actor.id, notes="Go-live operacional real de propostas e templates.")
            stats.batches.append(BatchResult(str(propostas_batch.id), propostas_batch.total_rows, propostas_batch.source_app, propostas_batch.module_target, propostas_batch.status))
            promote_files(db, propostas_batch, actor, stats)
            db.commit()

            write_exports(db, stats)
            create_summary_notifications(db, actor, stats)
            emit_event(db, "real_data.activation.completed", "RealDataActivation", "local-go-live", "legacy_promotion", {"summary": "Ativacao real concluida.", "actor_user_id": actor.id, "batches": len(stats.batches), "rows": sum(b.rows for b in stats.batches), "contains_secrets": False}, actor_user_id=actor.id)
            db.commit()
        except Exception as exc:
            db.rollback()
            stats.errors.append(str(exc))
            raise
        finally:
            stats.events_created = max(db.query(EventLog).count() - before_events, 0)
            stats.notifications_created += max(db.query(Notification).count() - before_notifications, 0)

    print("GO_LIVE_OPERATIONAL_SUMMARY")
    print(f"batches={len(stats.batches)}")
    print(f"rows={sum(b.rows for b in stats.batches)}")
    print(f"suppliers_created={stats.official_suppliers_created}")
    print(f"items_created={stats.official_items_created}")
    print(f"prices_activated={stats.prices_activated}")
    print(f"access_records_created={stats.access_records_created}")
    print(f"vault_secrets_imported={stats.vault_secrets_imported}")
    print(f"assets_created={stats.assets_created}")
    print(f"production_ops={stats.production_ops}")
    print(f"project_records={stats.project_records}")
    print(f"proposal_files={stats.proposal_files}")
    print(f"knowledge_files_indexed={stats.knowledge_files_indexed}")
    print(f"exports={len(stats.exports_written)}")
    for path in stats.exports_written:
        print(f"export={path}")
    if stats.empty_sources:
        print(f"empty_sources={','.join(stats.empty_sources)}")


if __name__ == "__main__":
    main()
